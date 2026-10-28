# -*- coding: utf-8 -*-
"""Verifica: 'gen-26' (data Excel auto-convertita) = '01/2026' (testo)?
Confronta cosa calcola RevMemo per i due modi di scrivere la scadenza.
"""
import datetime as dt
import os
import shutil
import sys
import tempfile

PACCHETTO = "/home/z/my-project/download/RevMemo"
tmp = tempfile.mkdtemp(prefix="revmemo_gen26_")
for f in ("revmemo.py", "template_email.html", "logo.png"):
    shutil.copy2(os.path.join(PACCHETTO, f), os.path.join(tmp, f))
sys.path.insert(0, tmp)
import revmemo  # noqa: E402
from openpyxl import Workbook  # noqa: E402

wb = Workbook()
ws = wb.active
ws.title = "Clienti"
# stessa struttura del programma: intestazioni alla riga 4
intest = ["", "Nome Cliente", "Email", "Telefono", "Targa",
          "Scadenza Revisione", "Revisione Fatta", "Attivo"]
for col, nome in enumerate(intest, start=1):
    ws.cell(row=4, column=col, value=nome)

# cliente A: come scritto dall'utente -> Excel lo converte in data (gen-26)
ws.cell(row=5, column=2, value="Prova Data Gen")
ws.cell(row=5, column=3, value="prova.data@email.it")
ws.cell(row=5, column=5, value="GA111AA")
ws.cell(row=5, column=6, value=dt.datetime(2026, 1, 1))   # <- gen-26

# cliente B: versione TESTO "01/2026" (celle formattate come testo)
ws.cell(row=6, column=2, value="Prova Testo")
ws.cell(row=6, column=3, value="prova.testo@email.it")
ws.cell(row=6, column=5, value="GB222BB")
ws.cell(row=6, column=6, value="01/2026")

percorso = os.path.join(tmp, "clienti_revisioni.xlsx")
wb.save(percorso)

clienti = revmemo.leggi_clienti(percorso)
print()
for c in clienti:
    print(f"  {c['nome']:<18} valore cella -> scadenza calcolata: "
          f"{c['scadenza'].strftime('%d/%m/%Y')}")

a = next(c for c in clienti if c["targa"] == "GA111AA")
b = next(c for c in clienti if c["targa"] == "GB222BB")
uguale = a["scadenza"] == b["scadenza"] == dt.date(2026, 1, 31)
print()
print("gen-26 (data) e 01/2026 (testo) danno lo STESSO risultato: "
      f"{'SI - 31/01/2026' if uguale else 'NO - PROBLEMA'}")
sys.exit(0 if uguale else 1)
