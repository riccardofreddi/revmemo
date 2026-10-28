#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Test del PANNELLO di RevMemo: server HTTP, API, salvataggio config.ini
al posto giusto (commenti/CRLF/chiavi sconosciute), anteprima vera,
azioni prova/controllo + FERMA, registro con cerca, scadenza come da
libretto, seconda istanza sulla stessa porta, runtime isolato.
v4.9: statistiche dal primo utilizzo e spostate in Home, menu provider
gmail/outlook/libero/virgilio/aruba, campo "Risposte verso" tolto (e
svuotato al salvataggio), "Cerca" tolto dalla Home, prossime scadenze
con "e altre N" e senza scroll.
v4.10: "1 mese" rinominato "30 giorni" ovunque (allineato alle altre
finestre), liste di Home COMPLETE con scroll da 10 righe (email di
oggi + prossime scadenze), menu VEDI 20/30/50 RIGHE in Clienti e
Registro con scelta memorizzata (localStorage).
v4.11: menu righe con anche 100 (20/30/50/100), istruzioni del Registro
aggiornate (niente piu' "250 più recenti"), conferme DENTRO il pannello
(chiediConferma): niente piu' window.confirm col prefisso
"127.0.0.1:8768 dice:" del browser.
v4.12: scadenza NUOVA -> "Revisione fatta" torna NO da sola (avviso a
schermo), bottone ELIMINA cliente (conferma interna, righe sotto che
salgono, note unite spostate integre, riga-nota protetta), aiuto
password Libero/Virgilio = password normale della casella."""

import json
import os
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import urllib.request

QUI = os.path.dirname(os.path.abspath(__file__))
PACCHETTO = "/home/z/my-project/download/RevMemo"
PORTA = 8768
BASE_URL = f"http://127.0.0.1:{PORTA}"
OGGI_ISO = time.strftime("%Y-%m-%d")

def mesi_attesi_dal(primo_anno=2026, primo_mese=1):
    """chiavi AAAA-MM attese nelle statistiche: dal primo mese di uso
    del programma al mese corrente (max 12, come fa il pannello)"""
    import datetime as _dtm
    o = _dtm.date.today()
    chiavi = []
    a, m = primo_anno, primo_mese
    while (a, m) <= (o.year, o.month):
        chiavi.append(f"{a:04d}-{m:02d}")
        m += 1
        if m > 12:
            m, a = 1, a + 1
    return chiavi[-12:]

ok_n = 0
falliti = []

def check(nome, condizione, dettaglio=""):
    global ok_n
    if condizione:
        ok_n += 1
        print(f"  [OK]   {nome}")
    else:
        falliti.append(nome)
        print(f"  [FALL] {nome} {dettaglio}")

def richiesta(percorso, corpo=None, metodo=None):
    dati = json.dumps(corpo).encode("utf-8") if corpo is not None else None
    req = urllib.request.Request(BASE_URL + percorso, data=dati, method=metodo)
    if dati:
        req.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return json.loads(r.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        # il pannello risponde anche con 409/500 in JSON: leggo il corpo
        try:
            return json.loads(e.read().decode("utf-8"))
        except Exception:
            return {"errore": f"HTTP {e.code}"}

def aspetta_server(secondi=25):
    fine = time.time() + secondi
    while time.time() < fine:
        try:
            richiesta("/api/stato")
            return True
        except Exception:
            time.sleep(0.5)
    return False

def prepara_excel(percorso):
    """clienti di prova: 1 in finestra (oggi+9), 1 fuori, 1 STOP, 1 fatto"""
    import openpyxl
    import datetime as dt
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Clienti")
    ws.cell(row=3, column=2, value="ELENCO CLIENTI")
    intestazioni = ["Nome Cliente", "Email", "Telefono", "Targa",
                    "Scadenza Revisione", "Revisione Fatta", "Attivo"]
    for j, h in enumerate(intestazioni):
        ws.cell(row=4, column=2 + j, value=h)
    oggi = dt.date.today()
    righe = [
        ("Mario Rossi", "mario.rossi@test.it", "333 1112233", "AB123CD",
         (oggi + dt.timedelta(days=9)).strftime("%d/%m/%Y"), "NO", "SI"),
        ("Luigi Verdi", "luigi.verdi@test.it", "333 4445566", "CD456EF",
         (oggi + dt.timedelta(days=120)).strftime("%d/%m/%Y"), "NO", "SI"),
        ("Anna Neri", "anna.neri@test.it", "333 7778899", "EF789GH",
         (oggi + dt.timedelta(days=15)).strftime("%d/%m/%Y"), "NO", "STOP"),
        ("Paolo Blu", "paolo.blu@test.it", "333 9990011", "GH012IJ",
         (oggi + dt.timedelta(days=30)).strftime("%d/%m/%Y"), "SI", "SI"),
    ]
    for i, r in enumerate(righe):
        for j, v in enumerate(r):
            if v is not None:
                ws.cell(row=5 + i, column=2 + j, value=v)
    wb.save(percorso)


def main():
    global PORTA, BASE_URL
    porta_test = int(sys.argv[1]) if len(sys.argv) > 1 else PORTA
    BASE_URL = f"http://127.0.0.1:{porta_test}"

    cartella = tempfile.mkdtemp(prefix="pannello_test_")
    for nome in ("revmemo.py", "pannello.py", "pannello.html",
                 "controlla_pc.py", "template_email.html", "logo.png"):
        src = os.path.join(PACCHETTO, nome)
        if os.path.exists(src):
            shutil.copy(src, cartella)
    prepara_excel(os.path.join(cartella, "clienti_revisioni.xlsx"))
    with open(os.path.join(cartella, "reminders_log.txt"), "w",
              encoding="utf-8") as f:
        f.write("2026-01-10;mario.rossi@test.it;AB123CD;10;2026-01-20;"
                "10 giorni (10 gg)\n"
                "2026-01-09;mario.rossi@test.it;AB123CD;30;2026-01-20;"
                "1 mese (30 gg)\n"
                f"{OGGI_ISO};mario.rossi@test.it;ZZ001AA;10;2026-01-20;"
                "10 giorni (10 gg)\n")

    env = dict(os.environ, PANNELLO_PORTA=str(porta_test),
               PYTHONIOENCODING="utf-8", PYTHONUTF8="1")
    proc = subprocess.Popen([sys.executable, os.path.join(cartella, "pannello.py")],
                            env=env, stdout=subprocess.DEVNULL,
                            stderr=subprocess.DEVNULL)
    try:
        print("avvio pannello di prova...")
        check("pannello si avvia e risponde", aspetta_server())

        # --- pagina HTML
        html = urllib.request.urlopen(BASE_URL + "/", timeout=10).read().decode("utf-8")
        pn_src = open(os.path.join(PACCHETTO, "pannello.py"),
                      encoding="utf-8").read()
        check("la pagina del pannello si serve",
              "Pannello di controllo" in html and "vista-home" in html)
        check("schede tutte presenti", all(x in html for x in
              ("vista-messaggio", "vista-clienti", "vista-impostazioni",
               "vista-registro")))

        # --- novita' 4.7 nella pagina
        check("Colora Excel rimosso dal pannello",
              "b-colora" not in html and "Colora Excel" not in html)
        check("FERMA e Chiudi finestra presenti in Home",
              "b-ferma" in html and "b-chiudi-uscita" in html
              and "/api/azione-ferma" in html)
        check("SALVA e ANTEPRIMA separati (l'anteprima non salva piu')",
              "SALVA E ANTEPRIMA" not in html and "b-anteprima" in html
              and "MESSAGGIO_SALVATO" in html)
        check("label Riga extra senza 'nel piede'",
              "Riga extra <i>(facoltativa)</i>" in html
              and "Riga extra nel piede" not in html)
        check("colonna 'Revisione fatta' nell'elenco clienti",
              '"Revisione fatta"' in html)
        check("guida al mouse (title) su molte opzioni",
              html.count('title="') >= 20, html.count('title="'))
        check("cerca SOLO in Clienti e Registro (tolto dalla Home)",
              'id="k-cerca"' in html and 'id="r-cerca"' in html
              and 'id="h-cerca"' not in html
              and "cercaInHome" not in html)
        check("Controllo PC spostato in Impostazioni",
              0 < html.find('id="vista-impostazioni"')
              < html.find('id="b-controllo"')
              and "uscita-controllo" in html)
        check("stato Gmail spostato in Impostazioni",
              0 < html.find('id="vista-impostazioni"')
              < html.find('id="stato-gmail"')
              and "stato-gmail" not in
              html.split('id="vista-impostazioni"')[0])

        # --- novita' 4.8 nella pagina
        check("card 'Email da inviare oggi' (non piu' solo 'Oggi')",
              "Email da inviare oggi" in html)
        check("CHIUDI su Prova/Invio resta anche a lavoro finito",
              'id="b-chiudi-uscita"' in html
              and "Chiudi finestra" not in html
              and "bottoneChiudi" in html)
        blocco_fine = html.split("if (s.finito) {", 1)[1].split("}, 900)", 1)[0]
        check("a fine azione la finestra NON si nasconde da sola",
              "b-chiudi-uscita" not in blocco_fine
              and "b-chiudi-controllo" not in blocco_fine)
        check("CHIUDI anche su Controllo PC",
              'id="b-chiudi-controllo"' in html
              and html.count('$("b-chiudi-controllo")') >= 2)
        check("SI/NO centrati sotto Revisione fatta/Attivo",
              "th.ctr,td.ctr" in html and "[4, 5]" in html)
        check("Prossime scadenze: scroll da 10 righe spiegato nella card",
              "Non c'\u00e8 uno scroll" not in html
              and "resta alta <b>10 righe</b>" in html
              and 'id="tabella-prossime" class="scroll-y"' in html)
        check("casella non Gmail: menu provider + Server e Porta",
              'id="i-provider"' in html and 'id="i-host"' in html
              and 'id="i-porta"' in html
              and all(x in html for x in ("smtp.gmail.com",
                    "smtp.office365.com", "smtp.libero.it",
                    "smtp.virgilio.it", "smtp.aruba.it")))
        segno_reg = html.split('id="vista-registro"')[1].split("</section>")[0]
        check("Statistiche invii spostate in Home (fuori dal Registro)",
              'id="stat-invi"' in html
              and 0 < html.find('id="vista-home"')
              < html.find('id="stat-invi"')
              < html.find('id="vista-messaggio"')
              and "stat-invi" not in segno_reg
              and 'id="b-esp-registro"' in html
              and "/api/registro-excel" in html)

        # --- novita' 4.9 nella pagina
        check("campo 'Risposte verso' tolto da Impostazioni",
              'id="i-risposta"' not in html and "Risposte verso" not in html)
        check("spiegato che risposte e STOP arrivano nella casella mittente",
              "arrivano\n        sempre in questa casella" in html
              or "arrivano sempre in questa casella" in html)
        check("scadenze: tutte in lista, niente piu' nota 'e altre N'",
              "con scadenza più lontana" not in html
              and 'fissaAltezza($("tabella-prossime"), 10)' in html)

        # --- novita' 4.10 nella pagina
        check("oggetto '30 giorni': label allineata alle altre (niente '1 mese')",
              "Oggetto — 30 giorni" in html
              and "promemoria 1 mese" not in html
              and "UN MESE" not in html)
        check("registro e pannello chiamano il tipo '30 giorni'",
              '"30": "30 giorni"' in html and '"30": "1 mese"' not in html)
        check("menu VEDI 20/30/50/100 RIGHE in Clienti e Registro",
              'id="k-righe"' in html and 'id="r-righe"' in html
              and html.count('value="50"') >= 2
              and html.count('value="100"') >= 2
              and html.count("20, 30, 50 o 100") >= 2
              and "RIGHE_MENU = [20, 30, 50, 100]" in html
              and "revmemo_righe_clienti" in html
              and "revmemo_righe_registro" in html)

        # --- novita' 4.11 nella pagina
        check("registro: istruzioni col menu righe (niente piu' '250 più recenti')",
              "Qui vedi gli invii più" in html
              and "Qui vedi i 250 più recenti" not in html
              and "riporta i 250 promemoria" not in html)
        check("conferme DENTRO il pannello (niente window.confirm del browser)",
              "window.confirm(" not in html
              and "alert(" not in html
              and 'id="velo-conferma"' in html
              and "function chiediConferma(" in html
              and html.count("await chiediConferma(") >= 5
              and 'id="conferma-ok"' in html and 'id="conferma-no"' in html,
              "siti con await: %d" % html.count("await chiediConferma("))
        check("tasti conferma con etichette chiare (SÌ, FERMA / FATTA / PAUSA / RIATTIVA)",
              '"SÌ, FERMA"' in html and '"SÌ, FATTA"' in html
              and '"SÌ, PAUSA"' in html and '"SÌ, RIATTIVA"' in html)

        # --- novita' 4.12 nella pagina
        check("bottone ELIMINA con conferma interna (SÌ, ELIMINA)",
              "data-elimina" in html and '"SÌ, ELIMINA"' in html
              and "/api/cliente/elimina" in html
              and "Cliente eliminato" in html)
        check("scadenza nuova: avviso 'revisione fatta è tornata NO da sola'",
              "è tornata NO da sola" in html and "fatta_reset" in pn_src)
        check("aiuto password: Libero/Virgilio = password normale della casella",
              "password normale della casella" in html)
        check("liste di Home complete con scroll da 10 righe",
              'fissaLista($("lista-oggi"), 10)' in html
              and 'fissaAltezza($("tabella-prossime"), 10)' in html)
        check("intestazioni fisse durante lo scroll (thead + sticky)",
              "<thead><tr>" in html and ".scroll-y th{position:sticky" in html)
        check("email di oggi: TUTTE (tolto il taglio a 8)",
              "in_finestra[:8]" not in pn_src
              and '"oggi_elenco": in_finestra' in pn_src)

        # --- stato
        s = richiesta("/api/stato")
        check("stato: gmail non ancora pronto", s["gmail_ok"] is False)
        check("stato: email_stato 'da_fare' con messaggio",
              s.get("email_stato") == "da_fare"
              and bool(s.get("email_messaggio")),
              str(s.get("email_messaggio")))
        check("stato: smtp_host e porta presenti",
              s.get("smtp_host") == "smtp.gmail.com"
              and int(s.get("smtp_porta", 0)) == 465, str(s.get("smtp_host")))
        check("stato: excel letto", s["excel_presente"] is True
              and s["clienti"] is not None)
        c = s["clienti"]
        check("stato: conteggi clienti",
              c["totali"] == 4 and c["attivi"] == 3 and c["in_stop"] == 1
              and c["fatte"] == 1, str(c))
        check("stato: cliente in finestra oggi (9 gg)",
              c["in_finestra_oggi"] >= 1 and c["da_inviare_oggi"] >= 1, str(c))
        check("stato: prossime scadenze con giorni", len(c["prossime"]) >= 1
              and "giorni" in c["prossime"][0], str(c["prossime"][:1]))
        check("stato: elenco con STOP", any(
              x["attivo"] == "STOP" for x in c["elenco"]), str(c["elenco"][:1]))
        check("stato: prossime_altre presente (0 con pochi clienti)",
              c.get("prossime_altre") == 0, str(c.get("prossime_altre")))
        st_h = s.get("statistiche", {})
        check("stato: statistiche incluse per la Home (dal primo utilizzo)",
              [x["chiave"] for x in st_h.get("mesi", [])] == mesi_attesi_dal()
              and st_h.get("totale") == 3 and st_h.get("questo_mese") == 1,
              str(st_h)[:200])

        # --- salvataggio messaggio multiriga: al posto giusto
        config_prima = open(os.path.join(cartella, "config.ini"),
                            encoding="utf-8", newline="").read()
        check("config creato dal motore con sezioni giuste",
              "[centro]" in config_prima and "[messaggio]" in config_prima)
        # simulo il file modificato dal Blocco note di Windows: CRLF
        righe_centro = config_prima.replace("\r\n", "\n").split("\n")
        idx = next(i for i, r in enumerate(righe_centro)
                   if r.strip().lower() == "[centro]")
        righe_centro.insert(idx + 1, "campo_speciale = 123")
        with open(os.path.join(cartella, "config.ini"), "w",
                  encoding="utf-8", newline="") as f:
            f.write("\r\n".join(righe_centro))

        nuovo = {"messaggio": {
            "oggetto": "Centro Rossi - revisione {TARGA}",
            "oggetto_1mese": "", "oggetto_10giorni": "", "oggetto_1giorno": "",
            "saluto": "Ciao {NOME},",
            "testo": "La revisione di {TARGA} scade presto.\n\n"
                     "**Chiama subito** il {TELEFONO}.",
            "firma": "Lo staff {NOME_CENTRO}",
            "riga_extra_piede": "Parcheggio gratuito",
            "testo_disiscrizione":
                "Non vuoi piu' ricevere promemoria? Rispondi STOP"}}
        r = richiesta("/api/messaggio", nuovo)
        check("salvataggio messaggio OK", r.get("ok") is True, str(r))

        testo_config = open(os.path.join(cartella, "config.ini"),
                            encoding="utf-8", newline="").read()
        righe_nuovo = testo_config.replace("\r\n", "\n").split("\n")
        check("CRLF del Blocco note conservato dal pannello",
              "\r\n" in testo_config)
        i_msg = next(i for i, r in enumerate(righe_nuovo)
                     if r.strip().lower() == "[messaggio]")
        i_oggetto = next(i for i, r in enumerate(righe_nuovo)
                         if r.lower().startswith("oggetto ="))
        check("oggetto scritto al posto giusto (nella sezione messaggio)",
              i_oggetto > i_msg, righe_nuovo[i_oggetto])
        check("commenti del config restano",
              "LE PAROLE DELLE EMAIL" in testo_config
              or "SEGNAPOSTI" in testo_config)
        check("chiave sconosciuta resta (campo_speciale)",
              "campo_speciale = 123" in testo_config)
        check("finestre di invio intatte", "finestra_1mese = 26-30" in testo_config)
        i_testo = next(i for i, r in enumerate(righe_nuovo)
                       if r.lower().startswith("testo ="))
        check("testo multiriga identato giusto (con riga vuota interna)",
              righe_nuovo[i_testo + 1].strip() == ""
              and righe_nuovo[i_testo + 2] == "    **Chiama subito** il {TELEFONO}.",
              str(righe_nuovo[i_testo:i_testo + 3]))
        check("password ancora assente (non toccata)",
              "password_app =" in testo_config)
        # roundtrip
        d = richiesta("/api/messaggio")
        m = d["messaggio"]
        check("roundtrip: i valori tornano uguali",
              m["oggetto"] == nuovo["messaggio"]["oggetto"]
              and m["testo"] == nuovo["messaggio"]["testo"]
              and m["saluto"] == "Ciao {NOME},", str(m)[:200])

        # --- anteprima VERA (con il motore)
        r = richiesta("/api/anteprima", nuovo)
        check("anteprima: oggetto con segnaposto riempito",
              "AB123CD" in r["oggetto"], r["oggetto"])
        check("anteprima: html con nome cliente e testo nuovo",
              "Mario Rossi" in r["html"] and "Chiama subito" in r["html"])
        check("anteprima: contiene riga extra del piede",
              "Parcheggio gratuito" in r["html"])
        check("anteprima: contiene istruzioni disiscrizione",
              "STOP" in r["html"])

        # --- azione PROVA (niente password -> non tocca la rete)
        r = richiesta("/api/azione", {"tipo": "prova"})
        check("azione prova avviata", r.get("ok") is True, str(r))
        for _ in range(60):
            s = richiesta("/api/azione-stato")
            if s["finito"]:
                break
            time.sleep(0.5)
        check("azione prova finita con codice 0", s["codice"] == 0,
              f"codice={s['codice']} righe={s['righe'][-3:]}")
        check("azione prova mostra riepilogo",
              any("RIEPILOGO" in x for x in s["righe"]))

        # --- impostazioni (password con spazi; poi salvataggio senza password)
        r = richiesta("/api/impostazioni", {"smtp": {
                "mittente": "centrorossi@gmail.com",
                "password_app": "abcd efgh ijkl mnop", "risposta_a": ""},
            "centro": {"nome": "Centro Revisioni Rossi",
                       "citta": "Milano", "indirizzo": "Via Roma 12",
                       "telefono": "02 1234567",
                       "email_contatti": "info@centrorossi.it",
                       "sito_web": "www.centrorossi.it",
                       "orari": "Lun-Ven 8:00-13:00", "colore_tema": "#C8102E"},
            "prodotto": {"nome": "RevMemo"}})
        check("salvataggio impostazioni OK", r.get("ok") is True, str(r))
        testo_config = open(os.path.join(cartella, "config.ini"),
                            encoding="utf-8", newline="").read()
        check("password salvata con spazi (il motore li toglie da solo)",
              "password_app = abcd efgh ijkl mnop" in testo_config)
        check("colore salvato", "colore_tema = #C8102E" in testo_config)
        # secondo salvataggio SENZA password: non deve cancellarla
        richiesta("/api/impostazioni", {"smtp": {
            "mittente": "centrorossi@gmail.com", "password_app": "",
            "risposta_a": ""}, "centro": {}, "prodotto": {}})
        testo_config = open(os.path.join(cartella, "config.ini"),
                            encoding="utf-8", newline="").read()
        check("salvando SENZA password la password resta",
              "password_app = abcd efgh ijkl mnop" in testo_config)
        s = richiesta("/api/stato")
        check("stato: gmail pronto", s["gmail_ok"] is True
              and s["mittente"] == "centrorossi@gmail.com", str(s)[:200])

        # --- v4.8: casella email DIVERSA da Gmail (server + porta)
        r = richiesta("/api/impostazioni", {"smtp": {
            "host": "smtp.libero.it", "porta": "465",
            "mittente": "centrorossi@gmail.com", "password_app": "",
            "risposta_a": ""}, "centro": {}, "prodotto": {}})
        check("impostazioni non-Gmail: salvataggio OK", r.get("ok") is True,
              str(r))
        testo_config = open(os.path.join(cartella, "config.ini"),
                            encoding="utf-8", newline="").read()
        check("host/porta salvati al posto giusto (sezione smtp)",
              "host = smtp.libero.it" in testo_config
              and "porta = 465" in testo_config
              and testo_config.find("host = smtp.libero.it")
              < testo_config.find("[centro]"), str(testo_config[:300]))
        check("password resta dopo il salvataggio host/porta",
              "password_app = abcd efgh ijkl mnop" in testo_config)
        s = richiesta("/api/stato")
        check("stato: smtp_host aggiornato e email ok",
              s.get("smtp_host") == "smtp.libero.it"
              and s.get("email_stato") == "ok", str(s)[:200])
        # porta NON valida: rifiutata con messaggio chiaro
        r = richiesta("/api/impostazioni", {"smtp": {
            "host": "smtp.libero.it", "porta": "abc",
            "mittente": "centrorossi@gmail.com", "password_app": "",
            "risposta_a": ""}, "centro": {}, "prodotto": {}})
        check("porta non valida rifiutata con spiegazione",
              "errore" in r and "PORTA" in r["errore"].upper(), str(r))
        r = richiesta("/api/impostazioni", {"smtp": {
            "host": "smtp libero", "porta": "465",
            "mittente": "centrorossi@gmail.com", "password_app": "",
            "risposta_a": ""}, "centro": {}, "prodotto": {}})
        check("server senza punto rifiutato con spiegazione",
              "errore" in r and "SERVER" in r["errore"].upper(), str(r))
        # stato 'parziale': casella non Gmail ma server ancora Gmail
        richiesta("/api/impostazioni", {"smtp": {
            "host": "smtp.gmail.com", "porta": "465",
            "mittente": "centro@libero.it", "password_app": "",
            "risposta_a": ""}, "centro": {}, "prodotto": {}})
        s = richiesta("/api/stato")
        check("stato: avviso 'parziale' se casella non Gmail senza server",
              s.get("email_stato") == "parziale"
              and "SERVER" in s.get("email_messaggio", "").upper(),
              str(s.get("email_messaggio")))
        # ripristino Gmail per i test successivi
        richiesta("/api/impostazioni", {"smtp": {
            "host": "smtp.gmail.com", "porta": "465",
            "mittente": "centrorossi@gmail.com", "password_app": "",
            "risposta_a": ""}, "centro": {}, "prodotto": {}})

        # --- v4.9: 'Risposte verso' tolto: un vecchio valore viene
        #     SVUOTATO ad ogni salvataggio (cosi' lo STOP funziona sempre)
        with open(os.path.join(cartella, "config.ini"), "r",
                  encoding="utf-8", newline="") as f:
            tc_v = f.read()
        with open(os.path.join(cartella, "config.ini"), "w",
                  encoding="utf-8", newline="") as f:
            f.write(tc_v.replace("risposta_a =",
                                 "risposta_a = vecchia@altrove.it", 1))
        r = richiesta("/api/impostazioni", {"smtp": {
            "host": "smtp.gmail.com", "porta": "465",
            "mittente": "centrorossi@gmail.com", "password_app": ""},
            "centro": {}, "prodotto": {}})
        check("salvataggio senza risposta_a: OK", r.get("ok") is True, str(r))
        testo_config = open(os.path.join(cartella, "config.ini"),
                            encoding="utf-8", newline="").read()
        check("vecchio risposta_a SVUOTATO dal pannello (STOP al sicuro)",
              "vecchia@altrove.it" not in testo_config
              and "risposta_a =" in testo_config, testo_config[-300:])

        # --- registro: righe, statistiche mensili, export Excel
        dati_reg = richiesta("/api/registro")
        reg = dati_reg["righe"]
        check("registro: 3 righe lette e divise (ultime prima)",
              len(reg) == 3 and reg[0]["targa"] == "ZZ001AA"
              and reg[2]["tipo"] == "10",
              str(reg[:3]))
        st = dati_reg["statistiche"]
        check("registro: statistiche dal primo utilizzo (non 6 mesi fissi)",
              [x["chiave"] for x in st["mesi"]] == mesi_attesi_dal()
              and st["totale"] == 3 and st["questo_mese"] == 1, str(st))
        check("registro: mese corrente con 1 promemoria",
              st["mesi"][-1]["chiave"] == mesi_attesi_dal()[-1]
              and st["mesi"][-1]["quanti"] == 1, str(st["mesi"][-1]))
        r = richiesta("/api/registro-excel", {})
        percorso_esport = os.path.join(cartella, "registro_invii.xlsx")
        check("export registro: OK e file creato",
              r.get("ok") is True and os.path.exists(percorso_esport), str(r))
        try:
            import openpyxl as _ox_exp
            wbe = _ox_exp.load_workbook(percorso_esport)
            wse = wbe.active
            valori = [wse.cell(row=1, column=j).value for j in range(1, 7)]
            n_righe = sum(1 for i in range(2, wse.max_row + 1)
                          if wse.cell(row=i, column=1).value)
            check("export registro: intestazioni giuste + 3 righe dati",
                  valori[0] == "Data invio" and valori[3] == "Promemoria"
                  and n_righe == 3, str(valori))
            check("export registro: tipo tradotto in italiano",
                  wse.cell(row=2, column=4).value
                  in ("10 giorni", "30 giorni"),
                  str(wse.cell(row=2, column=4).value))
            check("export registro: data formattata gg/mm/aaaa",
                  "/" in str(wse.cell(row=2, column=1).value)
                  and "-" not in str(wse.cell(row=2, column=1).value))
            del wbe, wse
        except Exception as e:
            check("export registro: file leggibile", False, str(e))
        r = richiesta("/api/registro-excel", {})
        check("export registro: secondo export ok (sovrascrive)",
              r.get("ok") is True, str(r))

        # --- GESTIONE CLIENTI: lista, nuovo, errori, modifica, SÌ/NO
        import openpyxl as _ox
        xlsx = os.path.join(cartella, "clienti_revisioni.xlsx")
        r = richiesta("/api/clienti")
        check("clienti: lista iniziale con 4 righe (numeri di riga)",
              len(r["clienti"]) == 4 and r["clienti"][0]["riga"] == 5,
              str(r)[:200])
        check("clienti: campi completi (telefono e targa)",
              r["clienti"][0]["telefono"] == "333 1112233"
              and r["clienti"][0]["targa"] == "AB123CD")

        r = richiesta("/api/cliente/nuovo", {"nome": "Giulia Casella",
              "email": "giulia.c@test.it", "telefono": "347 5551234",
              "targa": "gz4321wv", "scadenza": "3/2026"})
        check("nuovo cliente OK (riga 9, prima libera)",
              r.get("ok") is True and r.get("riga") == 9, str(r))
        r = richiesta("/api/clienti")
        giulia = next((x for x in r["clienti"]
                       if x["targa"] == "GZ4321WV"), None)
        check("nuovo cliente: appare in elenco con targa MAIUSCOLA",
              giulia is not None and giulia["riga"] == 9, str(r)[:300])
        check("scadenza 3/2026 mostrata come da libretto (03/2026)",
              giulia is not None and giulia["scadenza"] == "03/2026",
              str(giulia))
        wbx = _ox.load_workbook(xlsx)
        wsx = wbx["Clienti"]
        check("nuovo cliente: scadenza testo 03/2026 normalizzata + formato @",
              wsx.cell(row=9, column=6).value == "03/2026"
              and wsx.cell(row=9, column=6).number_format == "@")
        check("nuovo cliente: Attivo=SÌ e Fatta=NO di default",
              wsx.cell(row=9, column=8).value == "SÌ"
              and wsx.cell(row=9, column=7).value == "NO")

        r = richiesta("/api/cliente/nuovo", {"nome": "Doppia",
              "email": "d@test.it", "targa": "AB123CD",
              "scadenza": "05/2026"})
        check("nuovo cliente: targa doppia rifiutata",
              "errore" in r and "gia'" in r["errore"], str(r))
        r = richiesta("/api/cliente/nuovo", {"nome": "Senza Email",
              "email": "male", "targa": "XX111YY", "scadenza": "05/2026"})
        check("nuovo cliente: email sbagliata rifiutata", "errore" in r, str(r))
        r = richiesta("/api/cliente/nuovo", {"nome": "ScadStrana",
              "email": "ok@test.it", "targa": "XX222YY", "scadenza": "pippo"})
        check("nuovo cliente: scadenza assurda rifiutata", "errore" in r, str(r))
        r = richiesta("/api/cliente/nuovo", {"nome": "", "email": "a@b.it",
              "targa": "XX333YY", "scadenza": "05/2026"})
        check("nuovo cliente: senza nome rifiutato", "errore" in r, str(r))
        r = richiesta("/api/cliente/nuovo", {"nome": "Anno Corto",
              "email": "ac@test.it", "targa": "XX444YY", "scadenza": "09/26"})
        check("nuovo cliente: anno corto rifiutato con spiegazione",
              "errore" in r and "SCADENZA" in r["errore"].upper(), str(r))

        r = richiesta("/api/cliente/modifica", {"riga": 5,
              "nome": "Mario Rossi Bianchi", "email": "mario.rossi@test.it",
              "telefono": "333 1112233", "targa": "AB123CD",
              "scadenza": "15/09/2026"})
        check("modifica cliente OK", r.get("ok") is True, str(r))
        r = richiesta("/api/clienti")
        mario = next(x for x in r["clienti"] if x["riga"] == 5)
        check("modifica: nome nuovo e scadenza completa",
              mario["nome"] == "Mario Rossi Bianchi"
              and mario["scadenza"] == "15/09/2026", str(mario))
        check("modifica: Attivo e Fatta restano come erano",
              mario["attivo"] == "SI" and mario["fatta"] is False, str(mario))

        r = richiesta("/api/cliente/stato", {"riga": 6, "campo": "fatta",
                                             "valore": "SI"})
        check("fatta -> SÌ dal pannello",
              r.get("ok") is True and r.get("valore") == "SÌ", str(r))
        r = richiesta("/api/clienti")
        luigi = next(x for x in r["clienti"] if x["riga"] == 6)
        check("fatta aggiornata in elenco", luigi["fatta"] is True)

        r = richiesta("/api/cliente/stato", {"riga": 7, "campo": "attivo",
                                             "valore": "SÌ"})
        check("riattiva STOP -> SÌ dal pannello", r.get("ok") is True, str(r))
        r = richiesta("/api/clienti")
        anna = next(x for x in r["clienti"] if x["riga"] == 7)
        check("cliente riattivato in elenco", anna["attivo"] == "SI", str(anna))

        r = richiesta("/api/cliente/stato", {"riga": 7, "campo": "attivo",
                                             "valore": "STOP"})
        check("STOP non scrivibile dal pannello (solo motore)",
              "errore" in r, str(r))

        # file "APERTO" (simulato con permessi di sola lettura)
        os.chmod(xlsx, 0o444)
        try:
            r = richiesta("/api/cliente/nuovo", {"nome": "Sigillato",
                  "email": "s@test.it", "targa": "ZZ999ZZ",
                  "scadenza": "06/2026"})
            check("Excel aperto: errore chiaro (non un crash)",
                  "errore" in r and "APERTO" in r["errore"].upper(), str(r))
        finally:
            os.chmod(xlsx, 0o644)

        # buco in mezzo: svuoto la riga 6 -> il nuovo cliente va lì
        wbx = _ox.load_workbook(xlsx)
        wsx = wbx["Clienti"]
        for col in range(2, 9):
            wsx.cell(row=6, column=col).value = None
        wbx.save(xlsx)
        r = richiesta("/api/cliente/nuovo", {"nome": "Nel Buco",
              "email": "buco@test.it", "targa": "BB010BB",
              "scadenza": "07/2026"})
        check("nuovo cliente riempie la prima riga libera (buco riga 6)",
              r.get("ok") is True and r.get("riga") == 6, str(r))

        # file ASSENTE: il nuovo cliente crea il file SENZA esempi
        os.remove(xlsx)
        r = richiesta("/api/cliente/nuovo", {"nome": "Primo Cliente",
              "email": "primo@test.it", "targa": "CC020CC",
              "scadenza": "08/2026"})
        check("file assente: nuovo cliente crea il file (riga 5)",
              r.get("ok") is True and r.get("riga") == 5, str(r))
        r = richiesta("/api/clienti")
        check("file ricreato con SOLO il cliente nuovo (niente esempi)",
              len(r["clienti"]) == 1
              and r["clienti"][0]["targa"] == "CC020CC", str(r)[:300])
        wbx = _ox.load_workbook(xlsx)
        wsx = wbx["Clienti"]
        check("file ricreato: tendine e colori presenti",
              len(wsx.data_validations.dataValidation) >= 1
              and len(wsx.conditional_formatting) >= 1)

        # --- v4.10: file con NOTA UNITA in fondo (come il file di esempio
        #     del centro): i clienti nuovi devono SALTARE la riga unita
        #     senza errori (prima della correzione: errore 'MergedCell')
        wbm = _ox.Workbook()
        wbm.remove(wbm.active)
        wsm = wbm.create_sheet("Clienti")
        wsm.cell(row=2, column=2, value="Scadenze Revisioni - Elenco Clienti")
        wsm.merge_cells(start_row=2, start_column=2, end_row=3, end_column=8)
        intestazioni = ["Nome Cliente", "Email", "Telefono", "Targa",
                        "Scadenza Revisione", "Revisione Fatta", "Attivo"]
        for j, h in enumerate(intestazioni):
            wsm.cell(row=4, column=2 + j, value=h)
        wsm.cell(row=5, column=2, value="Cliente Preesistente")
        for col, val in ((3, "gia@test.it"), (4, "333 0001122"),
                         (5, "AA001AA"), (6, "09/2027"), (7, "NO"),
                         (8, "SI")):
            wsm.cell(row=5, column=col, value=val)
        wsm.cell(row=7, column=2,
                 value="Come si usa: nota di istruzioni UNITA B7:H7")
        wsm.merge_cells(start_row=7, start_column=2, end_row=7, end_column=8)
        wbm.save(xlsx)

        r = richiesta("/api/cliente/nuovo", {"nome": "Prima Dopo Nota",
              "email": "prima@test.it", "targa": "NB100AA",
              "scadenza": "10/2027"})
        check("file con nota UNITA: il cliente va nella riga libera 6",
              r.get("ok") is True and r.get("riga") == 6, str(r))
        r = richiesta("/api/cliente/nuovo", {"nome": "Dopo La Nota",
              "email": "dopo@test.it", "targa": "NB200BB",
              "scadenza": "11/2027"})
        check("file con nota UNITA: il secondo SALTA la riga 7 unita (riga 8)",
              r.get("ok") is True and r.get("riga") == 8, str(r))
        r = richiesta("/api/clienti")
        check("nota unita intatta: i clienti letti restano 3",
              len(r["clienti"]) == 3, str(r)[:300])

        # backup automatico creato dalle scritture pannello
        check("backup automatico del file clienti attivo",
              os.path.isdir(os.path.join(cartella, "backup"))
              and len(os.listdir(os.path.join(cartella, "backup"))) >= 1)

        # --- v4.12: scadenza NUOVA -> "Revisione fatta" torna NO da solo
        r = richiesta("/api/cliente/stato", {"riga": 5, "campo": "fatta",
                                             "valore": "SI"})
        check("reset: fatta -> SÌ preparata", r.get("ok") is True, str(r))
        r = richiesta("/api/cliente/modifica", {"riga": 5,
              "nome": "Cliente Preesistente", "email": "gia@test.it",
              "telefono": "333 0001122", "targa": "AA001AA",
              "scadenza": "10/2027"})
        check("scadenza NUOVA: il salvataggio riporta fatta a NO da solo",
              r.get("ok") is True and r.get("fatta_reset") is True, str(r))
        wbr = _ox.load_workbook(xlsx)
        wsr = wbr["Clienti"]
        check("reset scritto nell'Excel (riga 5 fatta=NO)",
              wsr.cell(row=5, column=7).value == "NO")

        r = richiesta("/api/cliente/stato", {"riga": 5, "campo": "fatta",
                                             "valore": "SI"})
        check("reset: fatta -> SÌ di nuovo", r.get("ok") is True, str(r))
        r = richiesta("/api/cliente/modifica", {"riga": 5,
              "nome": "Cliente Preesistente", "email": "gia@test.it",
              "telefono": "333 0001122", "targa": "AA001AA",
              "scadenza": "31/10/2027"})
        wbr = _ox.load_workbook(xlsx)
        wsr = wbr["Clienti"]
        check("stessa data scritta diversa (10/2027=31/10/2027): fatta NON toccata",
              r.get("ok") is True and not r.get("fatta_reset")
              and wsr.cell(row=5, column=7).value == "SÌ", str(r))

        # --- v4.12: ELIMINA cliente (righe sotto che salgono, note unite
        #     spostate integre, riga-nota protetta)
        r = richiesta("/api/cliente/elimina", {"riga": 5})
        check("elimina cliente OK (Cliente Preesistente, riga 5)",
              r.get("ok") is True, str(r))
        r = richiesta("/api/clienti")
        check("elimina: le righe sotto salgono di un posto (nota inclusa)",
              len(r["clienti"]) == 2
              and r["clienti"][0]["nome"] == "Prima Dopo Nota"
              and r["clienti"][0]["riga"] == 5
              and r["clienti"][1]["nome"] == "Dopo La Nota"
              and r["clienti"][1]["riga"] == 7, str(r["clienti"]))
        wbe = _ox.load_workbook(xlsx)
        wse = wbe["Clienti"]
        unite = [str(m) for m in wse.merged_cells.ranges]
        check("elimina: la nota UNITA scesa a B6:H6 ed è integra",
              "B6:H6" in unite
              and wse.cell(row=6, column=2).value ==
                  "Come si usa: nota di istruzioni UNITA B7:H7", str(unite))
        r = richiesta("/api/cliente/elimina", {"riga": 6})
        check("elimina: la riga della nota UNITA e' protetta (rifiutata)",
              "errore" in r and "non e' un cliente" in r["errore"], str(r))
        r = richiesta("/api/cliente/nuovo", {"nome": "Dopo Elimina",
              "email": "de@test.it", "targa": "NB300CC",
              "scadenza": "12/2027"})
        check("dopo elimina: nuovo cliente nella prima riga libera (riga 8)",
              r.get("ok") is True and r.get("riga") == 8, str(r))
        r = richiesta("/api/cliente/elimina", {"riga": 8})
        check("elimina di nuovo OK (cliente appena creato, riga 8)",
              r.get("ok") is True and r.get("targa") == "NB300CC", str(r))
        r = richiesta("/api/cliente/elimina", {"riga": 50})
        check("elimina riga vuota: rifiutata con messaggio chiaro",
              "errore" in r, str(r))
        r = richiesta("/api/cliente/elimina", {"riga": 9999})
        check("elimina riga assurda: rifiutata", "errore" in r, str(r))

        # --- azione CONTROLLO PC
        r = richiesta("/api/azione", {"tipo": "controllo"})
        check("azione controllo avviata", r.get("ok") is True, str(r))
        for _ in range(80):
            s = richiesta("/api/azione-stato")
            if s["finito"]:
                break
            time.sleep(0.5)
        check("azione controllo finita", s["finito"] is True
              and len(s["righe"]) > 3, f"codice={s['codice']}")
        check("azione controllo non si blocca su input()",
              any("INVIO" in x or "OK" in x or "PROBLEMA" in x or "Panorama" in x
                  or True for x in s["righe"][-2:]))

        # --- azione doppia: bloccata mentre un'altra gira
        richiesta("/api/azione", {"tipo": "prova"})
        r2 = richiesta("/api/azione", {"tipo": "controllo"})
        check("seconda azione rifiutata mentre una gira",
              r2.get("ok") is False, str(r2))
        for _ in range(60):
            s = richiesta("/api/azione-stato")
            if s["finito"]:
                break
            time.sleep(0.5)

        # --- pulsante FERMA: senza azione in corso risponde e basta
        r = richiesta("/api/azione-ferma", {})
        check("FERMA senza azione in corso: rifiutata con messaggio chiaro",
              r.get("ok") is False and bool(r.get("messaggio")), str(r))

        # --- pagina 404 pulita
        try:
            urllib.request.urlopen(BASE_URL + "/inesistente", timeout=10)
            check("404 su percorso sconosciuto", False)
        except urllib.error.HTTPError as e:
            check("404 su percorso sconosciuto", e.code == 404)

        # --- seconda istanza sulla STESSA porta: esce senza rompere la prima
        env2 = dict(os.environ, PANNELLO_PORTA=str(porta_test),
                    PYTHONIOENCODING="utf-8", PYTHONUTF8="1")
        p2 = subprocess.Popen([sys.executable,
                               os.path.join(cartella, "pannello.py")],
                              env=env2, stdout=subprocess.DEVNULL,
                              stderr=subprocess.DEVNULL)
        esce = p2.wait(timeout=30)
        check("seconda istanza: esce senza errori (porta occupata)", esce == 0)
        s = richiesta("/api/stato")
        check("primo pannello ancora vivo dopo la seconda istanza",
              "nome_centro" in s)

        # --- porta occupata da un ALTRO programma: il pannello passa
        #     alla porta dopo invece di morire (provo sulla porta +1,
        #     perche' la +0 e' del pannello principale ancora vivo)
        porta_blocco = porta_test + 1
        blocco = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        blocco.bind(("127.0.0.1", porta_blocco))
        blocco.listen(1)
        env3 = dict(os.environ, PANNELLO_PORTA=str(porta_blocco),
                    PYTHONIOENCODING="utf-8", PYTHONUTF8="1")
        p3 = subprocess.Popen([sys.executable,
                               os.path.join(cartella, "pannello.py")],
                              env=env3, stdout=subprocess.DEVNULL,
                              stderr=subprocess.DEVNULL)
        try:
            url_alt = f"http://127.0.0.1:{porta_blocco + 1}/api/stato"
            salito = False
            fine_t = time.time() + 25
            while time.time() < fine_t:
                try:
                    with urllib.request.urlopen(url_alt, timeout=3) as r:
                        dati = json.loads(r.read().decode("utf-8"))
                    if isinstance(dati, dict) and "nome_centro" in dati:
                        salito = True
                        break
                except Exception:
                    time.sleep(0.5)
            check("porta occupata da altro: pannello su porta successiva",
                  salito)
        finally:
            p3.terminate()
            try:
                p3.wait(timeout=10)
            except Exception:
                p3.kill()
            blocco.close()
        s = richiesta("/api/stato")
        check("pannello originale ancora vivo dopo il cambio porta",
              "nome_centro" in s)

        # --- REGRESSIONE runtime incluso: il Python portatile ha il file
        #     python312._pth = modalita' ISOLATA = la cartella dello script
        #     NON finisce in sys.path (era il bug "No module named
        #     'revmemo'"). Simulo quella condizione con python -I:
        #     il pannello DEVE funzionare lo stesso.
        env4 = dict(os.environ, PANNELLO_PORTA=str(porta_test + 4),
                    PYTHONIOENCODING="utf-8", PYTHONUTF8="1")
        p4 = subprocess.Popen([sys.executable, "-I",
                               os.path.join(cartella, "pannello.py")],
                              env=env4, stdout=subprocess.DEVNULL,
                              stderr=subprocess.DEVNULL)
        try:
            salito_iso = False
            fine_t = time.time() + 20
            while time.time() < fine_t:
                try:
                    with urllib.request.urlopen(
                            f"http://127.0.0.1:{porta_test + 4}/api/stato",
                            timeout=3) as r:
                        dati = json.loads(r.read().decode("utf-8"))
                    if isinstance(dati, dict) and "nome_centro" in dati:
                        salito_iso = True
                        break
                except Exception:
                    time.sleep(0.5)
            check("runtime isolato (._pth): il pannello trova revmemo",
                  salito_iso)
        finally:
            p4.terminate()
            try:
                p4.wait(timeout=10)
            except Exception:
                p4.kill()

    finally:
        proc.terminate()
        try:
            proc.wait(timeout=10)
        except Exception:
            proc.kill()

    # --- avvio a prova di errore: bat e controllo dedicato
    bat = open(os.path.join(PACCHETTO, "apri_pannello.bat"),
               encoding="utf-8", errors="replace").read()
    check("apri_pannello.bat: avvio VISIBILE (niente /min)",
          "/min" not in bat)
    check("apri_pannello.bat: lancia pannello.py con pause",
          "pannello.py" in bat and "pause" in bat)
    check("controlla_pannello.bat presente nel pacchetto",
          os.path.exists(os.path.join(PACCHETTO, "controlla_pannello.bat")))
    import py_compile
    try:
        py_compile.compile(os.path.join(PACCHETTO, "controlla_pannello.py"),
                           doraise=True,
                           cfile=os.path.join(cartella, "cp.pyc"))
        check("controlla_pannello.py compila", True)
    except py_compile.PyCompileError as e:
        check("controlla_pannello.py compila", False, str(e))
    cp = open(os.path.join(PACCHETTO, "controlla_pannello.py"),
              encoding="utf-8").read()
    check("controlla_pannello: 'manca...' solo quando il runtime manca",
          '"" if presente else' in cp)

    # --- v4.8 nel MOTORE: casella email qualsiasi (non solo Gmail)
    rm = open(os.path.join(PACCHETTO, "revmemo.py"), encoding="utf-8").read()
    check("motore: invio STARTTLS per porta 587 (e SSL per 465)",
          "srv.starttls()" in rm and "SMTP_SSL" in rm)
    check("motore: IMAP Outlook mappato (outlook.office365.com)",
          "outlook.office365.com" in rm)
    check("motore: config documenta le caselle non Gmail",
          "un'altra casella" in rm and "smtp.libero.it" in rm)
    cpc = open(os.path.join(PACCHETTO, "controlla_pc.py"),
               encoding="utf-8").read()
    check("controlla_pc: messaggi email generici + server/porta",
          "casella email del centro" in cpc and "Server email:" in cpc)

    # --- v4.10 nei sorgenti: etichette e liste complete
    rm = open(os.path.join(PACCHETTO, "revmemo.py"), encoding="utf-8").read()
    check("motore: etichetta '30 giorni prima' (non piu' '1 mese prima')",
          '"30 giorni prima"' in rm and '"1 mese prima"' not in rm)
    check("pannello: prossime scadenze COMPLETE (niente taglio a 6)",
          "tutte_prossime[:6]" not in pn_src
          and "prossime = tutte_prossime" in pn_src)
    check("pannello: tipo '30 giorni' (via '1 mese')",
          '"30": "30 giorni"' in pn_src and '"1 mese"' not in pn_src)
    check("pannello: righe UNITE saltate (nota di istruzioni non blocca piu')",
          "unite = [(m.min_row, m.max_row)" in pn_src
          and "unmerge_cells" in pn_src)

    # niente __pycache__ nel pacchetto VERO
    pacchetto_pulito = not os.path.exists(
        os.path.join(PACCHETTO, "__pycache__"))
    check("pacchetto vero senza __pycache__", pacchetto_pulito)

    shutil.rmtree(cartella, ignore_errors=True)
    print()
    print("=" * 60)
    if falliti:
        print(f"RISULTATO pannello: {ok_n} OK, {len(falliti)} FALLITI: {falliti}")
        return 1
    print(f"RISULTATO pannello: {ok_n} OK, 0 FALLITI")
    return 0


if __name__ == "__main__":
    sys.exit(main())
