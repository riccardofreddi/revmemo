#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
=====================================================================
 RevMemo - Promemoria revisioni automatici        versione 3 (PRO)
=====================================================================
Cosa fa:
 - Legge i clienti dal file Excel (clienti_revisioni.xlsx)
 - Calcola la scadenza per MESE come sul libretto (es. 09/2026 =
   ultimo giorno di settembre 2026)
 - Invia email professionali in HTML (logo, colori, contatti del
   centro) a: 1 mese, 10 giorni e 1 giorno dalla scadenza
 - Manda al titolare un riepilogo giornaliero di cio' che e' partito
 - Registra ogni invio in reminders_log.txt (zero doppioni: ogni
   promemoria parte UNA sola volta per ogni scadenza del cliente)
 - Backup automatico del file clienti prima di ogni esecuzione

Regole dei clienti:
 - Attivo = NO            -> nessuna email
 - Revisione Fatta = SI'  -> nessuna email (finche' non aggiorni la scadenza)
 - Finestre di invio (config.ini): 26-30 gg, 8-10 gg, 1-2 gg

Tutta la configurazione sta in config.ini: il codice NON si tocca.

Uso:
    python revmemo.py             invio reale
    python revmemo.py --dry-run   PROVA: mostra senza inviare
    python revmemo.py --file x.xlsx   usa un file clienti alternativo
=====================================================================
"""

import argparse
import calendar
import configparser
import datetime as dt
import os
import re
import shutil
import smtplib
import sys
import time
import zipfile
from email.message import EmailMessage
from email.utils import formataddr

# Rende sicura la stampa a schermo su qualsiasi console Windows
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))

# Se c'e' il runtime incluso e NON lo stiamo usando, ripartiamo con quello:
# cosi' il Python installato sul PC non conta piu' (zero problemi di versione).
RUNTIME_PY = os.path.join(BASE, "runtime", "python.exe")
if (os.name == "nt"
        and os.path.isfile(RUNTIME_PY)
        and os.path.abspath(sys.executable).lower() != os.path.abspath(RUNTIME_PY).lower()
        and os.environ.get("REVMEMO_NEL_RUNTIME") != "1"):
    os.environ["REVMEMO_NEL_RUNTIME"] = "1"
    try:
        os.execv(RUNTIME_PY, [RUNTIME_PY, os.path.abspath(__file__)] + sys.argv[1:])
    except Exception:
        pass  # se il rilancio fallisce, prosegue con il Python corrente

try:
    from openpyxl import load_workbook
except ImportError as e:
    print("Problema con la libreria openpyxl:", e)
    print("Soluzione: fai partire il programma con la cartella runtime inclusa")
    print("nel pacchetto (doppio click su avvia_prova.bat o avvia_invio.bat).")
    print("Se il runtime non c'e', riscarica il pacchetto completo o esegui:")
    print("   pip install openpyxl")
    sys.exit(1)

# =====================================================================
# 1) CONFIGURAZIONE (da config.ini)
# =====================================================================

PERCORSO_CONFIG = os.path.join(BASE, "config.ini")

CONFIG_DEFAULT = """\
; =====================================================================
;  RevMemo - CONFIGURAZIONE
;  Questo e' l'unico file da modificare. Il codice non si tocca mai.
; =====================================================================

[smtp]
; Casella Gmail DALLA CUI partono i promemoria
host = smtp.gmail.com
porta = 465
mittente = latuaemail@gmail.com
; Password per app di Google (16 caratteri) - vedi ISTRUZIONI_CENTRO.md punto 3
password_app =
; Risposte dei clienti (facoltativo): quando un cliente risponde a un
; promemoria, la risposta arriva nella casella mittente (quella sopra).
; Se vuoi che arrivino su un'altra casella, scrivila qui.
risposta_a =

[centro]
; Questi dati appaiono nelle email che ricevono i CLIENTI
nome = Centro Revisioni Esempio
citta = Milano
telefono = 02 1234567
email_contatti = info@centrorevisioni.it
; Colore principale delle email (codice esadecimale, es. blu #1F3B66, rosso #C8102E)
colore_tema = #1F3B66

[prodotto]
; Nome del prodotto mostrato in piccolo nel piede delle email.
; Cambiabile liberamente (white label).
nome = RevMemo

[promemoria]
; Finestre di invio: (da giorni-al giorni) di anticipo sulla scadenza
finestra_1mese = 26-30
finestra_10giorni = 8-10
finestra_1giorno = 1-2

[invii]
; Pausa tra un'email e l'altra (secondi) - 1.5 consigliato
pausa_secondi = 1.5
; Riepilogo giornaliero al titolare: SI / NO
notifica_titolare = SI
; Email del titolare (se vuota, usa la stessa del mittente)
email_titolare =
; Nomi dei file (non serve toccarli)
file_excel = clienti_revisioni.xlsx
file_log = reminders_log.txt

[sms]
; Modulo SMS predisposto: quando il centro decide di attivare gli SMS
; (gateway tipo Aruba/Skebby) questa sezione verra' collegata.
; Per ora lasciare NO.
attivo = NO
"""


def _si(v, predefinito=False):
    if v is None:
        return predefinito
    return str(v).strip().upper() in ("SI", "SÌ", "S", "TRUE", "VERO", "1", "YES", "Y")


def carica_config():
    if not os.path.exists(PERCORSO_CONFIG):
        try:
            with open(PERCORSO_CONFIG, "w", encoding="utf-8") as f:
                f.write(CONFIG_DEFAULT)
            print("config.ini non c'era: l'ho ricreato adesso (valori di partenza).")
            print("ATTENZIONE: se avevi gia' messo mittente o password, reinseriscili.\n")
        except Exception as e:
            print(f"ERRORE: manca config.ini e non riesco a crearlo: {e}")
            sys.exit(1)
    cfg = configparser.ConfigParser(interpolation=None)
    cfg.read(PERCORSO_CONFIG, encoding="utf-8")

    c = {}
    smtp = cfg["smtp"] if cfg.has_section("smtp") else {}
    c["smtp_host"] = smtp.get("host", "smtp.gmail.com")
    c["smtp_porta"] = int(smtp.get("porta", "465"))
    c["mittente"] = smtp.get("mittente", "").strip()
    # Tolgo TUTTI gli spazi: Google mostra la chiave a gruppi di 4
    # ("abcd efgh ijkl mnop") e se viene incollata con gli spazi il login fallirebbe.
    c["password"] = "".join(smtp.get("password_app", "").split())
    c["risposta_a"] = smtp.get("risposta_a", "").strip()

    off = cfg["centro"] if cfg.has_section("centro") else {}
    c["nome_centro"] = off.get("nome", "Centro Revisioni")
    c["citta"] = off.get("citta", "")
    c["telefono"] = off.get("telefono", "")
    c["email_contatti"] = off.get("email_contatti", "")
    colore = off.get("colore_tema", "#1F3B66").strip()
    if not colore.startswith("#"):
        colore = "#" + colore
    c["colore"] = colore

    prod = cfg["prodotto"] if cfg.has_section("prodotto") else {}
    c["prodotto"] = prod.get("nome", "RevMemo")

    prom = cfg["promemoria"] if cfg.has_section("promemoria") else {}

    def finestra(val, predef):
        try:
            a, b = str(val).split("-")
            return (int(a), int(b))
        except Exception:
            try:
                n = int(val)
                return (n, n)
            except Exception:
                return predef

    c["finestre"] = {
        30: finestra(prom.get("finestra_1mese", "26-30"), (26, 30)),
        10: finestra(prom.get("finestra_10giorni", "8-10"), (8, 10)),
        1: finestra(prom.get("finestra_1giorno", "1-2"), (1, 2)),
    }
    c["etichette"] = {30: "1 mese prima", 10: "10 giorni prima", 1: "1 giorno prima"}

    inv = cfg["invii"] if cfg.has_section("invii") else {}
    c["pausa"] = float(inv.get("pausa_secondi", "1.5"))
    c["notifica_titolare"] = _si(inv.get("notifica_titolare", "SI"), True)
    c["email_titolare"] = inv.get("email_titolare", "").strip()

    c["file_excel"] = inv.get("file_excel", "clienti_revisioni.xlsx")
    c["file_log"] = inv.get("file_log", "reminders_log.txt")
    return c


CFG = carica_config()
SMTP_HOST = CFG["smtp_host"]
SMTP_PORT = CFG["smtp_porta"]
MITTENTE_EMAIL = CFG["mittente"]
MITTENTE_PASSWORD = CFG["password"]
RISPOSTA_A = CFG["risposta_a"]
NOME_CENTRO = CFG["nome_centro"]
PRODOTTO = CFG["prodotto"]
FILE_EXCEL = CFG["file_excel"]
FILE_LOG = CFG["file_log"]

# =====================================================================
# 2) LETTURA DEL FILE EXCEL
# =====================================================================

COLONNE = {
    "nome":     "Nome Cliente",
    "email":    "Email",
    "telefono": "Telefono",          # opzionale (per SMS futuri)
    "targa":    "Targa",
    "scadenza": "Scadenza Revisione",
    "fatta":    "Revisione Fatta",
    "attivo":   "Attivo",
}
RIGA_INTESTAZIONI = 4
FOGLIO = "Clienti"


def ultimo_giorno_mese(anno, mese):
    return dt.date(anno, mese, calendar.monthrange(anno, mese)[1])


def converti_scadenza(val):
    """(data_scadenza, da_mese) — accetta 09/2026, 9/2026, date Excel
    auto-convertite (giorno 1 = mese del libretto) e date complete."""
    if isinstance(val, dt.datetime):
        val = val.date()
    if isinstance(val, dt.date):
        if val.day == 1:
            return ultimo_giorno_mese(val.year, val.month), True
        return val, False
    if isinstance(val, str):
        v = val.strip()
        for f in ("%m/%Y", "%m-%Y"):
            try:
                d = dt.datetime.strptime(v, f)
                return ultimo_giorno_mese(d.year, d.month), True
            except ValueError:
                continue
        for f in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
            try:
                return dt.datetime.strptime(v, f).date(), False
            except ValueError:
                continue
    return None, False


def leggi_clienti(percorso_excel):
    if not os.path.exists(percorso_excel):
        print(f"ERRORE: non trovo il file Excel: {percorso_excel}")
        sys.exit(1)
    try:
        wb = load_workbook(percorso_excel, data_only=True)
    except PermissionError:
        print("ERRORE: il file Excel non si puo' aprire perche' risulta APERTO")
        print("in Excel (o bloccato da OneDrive mentre sincronizza).")
        print("Soluzione: salva (Ctrl+S) e chiudi Excel, aspetta qualche secondo")
        print("e rilancia il programma. Se continua: riavvia il PC.")
        sys.exit(1)
    except zipfile.BadZipFile:
        print("ERRORE: il file Excel non e' leggibile (non e' un file Excel valido).")
        print("Aprilo con Excel: se si apre, salvalo di nuovo (Ctrl+S) e riprova.")
        sys.exit(1)
    ws = wb[FOGLIO] if FOGLIO in wb.sheetnames else wb.active

    indice = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=RIGA_INTESTAZIONI, column=col).value
        if val is not None:
            indice[str(val).strip().lower()] = col

    colonna_di, mancanti = {}, []
    for chiave, nome_col in COLONNE.items():
        c = indice.get(nome_col.strip().lower())
        if c is None:
            if chiave == "telefono":      # unica colonna facoltativa
                continue
            mancanti.append(nome_col)
        else:
            colonna_di[chiave] = c
    if mancanti:
        print("ERRORE: nel foglio Excel mancano le colonne: " + ", ".join(mancanti))
        print("Le intestazioni alla riga 4 devono essere: "
              + " | ".join(COLONNE.values()))
        sys.exit(1)

    clienti = []
    for row in range(RIGA_INTESTAZIONI + 1, ws.max_row + 1):
        def cella(chiave):
            if chiave not in colonna_di:
                return None
            return ws.cell(row=row, column=colonna_di[chiave]).value

        nome, email, targa = cella("nome"), cella("email"), cella("targa")
        scadenza, fatta, attivo = cella("scadenza"), cella("fatta"), cella("attivo")
        telefono = cella("telefono")

        # riga vuota o nota a fondo tabella -> salta
        if all(v is None or str(v).strip() == "" for v in (email, targa, scadenza)):
            continue

        data_scad, da_mese = converti_scadenza(scadenza)
        clienti.append({
            "riga": row,
            "nome": str(nome).strip() if nome else "",
            "email": str(email).strip() if email else "",
            "telefono": str(telefono).strip() if telefono else "",
            "targa": str(targa).strip().upper() if targa else "",
            "scadenza": data_scad,
            "da_mese": da_mese,
            "fatta": _si(fatta),
            "attivo": _si(attivo),
        })
    return clienti

# =====================================================================
# 3) REGISTRO INVII (anti-doppioni)
# =====================================================================

def percorso_log():
    return os.path.join(BASE, FILE_LOG)


_DATA_ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def gia_inviato(email, targa, tipo, scadenza):
    """True se quel promemoria e' GIA' partito per questa scadenza.
    La chiave e' email;targa;tipo;scadenza: ogni promemoria parte UNA
    sola volta per ciclo (se il cliente rinnova la revisione, la scadenza
    cambia e i promemoria ripartono da soli). Le righe del vecchio
    formato (senza scadenza) valgono comunque come 'gia' inviato':
    nessun rischio di doppioni dopo l'aggiornamento."""
    percorso = percorso_log()
    if not os.path.exists(percorso):
        return False
    email_l, targa_u = email.lower().strip(), targa.upper().strip()
    with open(percorso, "r", encoding="utf-8") as f:
        for line in f:
            parti = line.split(";")
            if len(parti) < 5:
                continue
            if (parti[1].strip().lower() == email_l
                    and parti[2].strip().upper() == targa_u
                    and parti[3].strip() == str(tipo)):
                quarto = parti[4].strip()
                # formato nuovo: campo 4 = scadenza AAAA-MM-GG
                # formato vecchio: campo 4 = descrizione -> comunque inviato
                if quarto == scadenza or not _DATA_ISO.match(quarto):
                    return True
    return False


def registra_invio(email, targa, tipo, giorni, descrizione, scadenza):
    riga = (f"{dt.date.today().isoformat()};{email.lower()};{targa.upper()};"
            f"{tipo};{scadenza};{descrizione} ({giorni} gg)")
    with open(percorso_log(), "a", encoding="utf-8") as f:
        f.write(riga + "\n")

# =====================================================================
# 4) TEMPLATE EMAIL (HTML + testo)
# =====================================================================

PERCORSO_TEMPLATE = os.path.join(BASE, "template_email.html")
PERCORSO_LOGO = os.path.join(BASE, "logo.png")


def leggi_template():
    if os.path.exists(PERCORSO_TEMPLATE):
        with open(PERCORSO_TEMPLATE, "r", encoding="utf-8") as f:
            return f.read()
    return None


_LOGO_CACHE = None


def leggi_logo():
    """Legge logo.png una volta sola; avvisa se e' illeggibile o vuoto
    (cosi' mai email mute senza spiegazione, es. file OneDrive non sincronizzato)."""
    global _LOGO_CACHE
    if _LOGO_CACHE is None and os.path.exists(PERCORSO_LOGO):
        try:
            with open(PERCORSO_LOGO, "rb") as f:
                dati = f.read()
            if len(dati) < 100:
                print("AVVISO: logo.png e' vuoto o danneggiato: le email usciranno SENZA logo.")
                print("        Ricopialo dal pacchetto o mettici il logo del centro.")
            else:
                _LOGO_CACHE = dati
        except Exception as e:
            print(f"AVVISO: non riesco a leggere logo.png ({e}): email senza logo.")
    return _LOGO_CACHE


def componi_email(cliente, giorni):
    data_scad = cliente["scadenza"].strftime("%d/%m/%Y")
    targa = cliente["targa"]
    if giorni > 1:
        tra = f"tra {giorni} giorni"
    else:
        tra = "tra 1 giorno"
    oggetto = f"[{NOME_CENTRO}] Promemoria revisione {targa} - scade {tra}"

    testo = (
        f"Gentile Cliente,\n\n"
        f"le ricordiamo che la revisione del veicolo targato {targa} "
        f"scade entro il {data_scad}.\n"
        f"Giorni rimanenti: {giorni}.\n\n"
        f"Per prenotare la revisione risponda a questa email "
        f"oppure chiami il numero {CFG['telefono']}.\n\n"
        f"{NOME_CENTRO}\n"
        f"---\n"
        f"Non desidera piu' ricevere promemoria? Risponda a questa email "
        f"con oggetto: STOP\n"
        f"Promemoria inviato con {PRODOTTO}\n"
    )

    template = leggi_template()
    html = None
    if template:
        giorni_label = f"{giorni} GIORNI" if giorni > 1 else "1 GIORNO"
        logo_presente = leggi_logo() is not None
        logo_img = (f'<img src="cid:logo" height="38" alt="logo {NOME_CENTRO}" '
                    'style="display:block;border:0;">') if logo_presente else ""
        citta_footer = f" &nbsp;|&nbsp; {CFG['citta']}" if CFG["citta"] else ""
        email_footer = (f" &nbsp;|&nbsp; {CFG['email_contatti']}"
                        if CFG["email_contatti"] else "")
        valori = {
            "{NOME_CENTRO}": NOME_CENTRO,
            "{CITTA}": CFG["citta"],
            "{TELEFONO}": CFG["telefono"],
            "{EMAIL_CENTRO}": CFG["email_contatti"],
            "{COLORE}": CFG["colore"],
            "{TARGA}": targa,
            "{SCADENZA}": data_scad,
            "{GIORNI}": str(giorni),
            "{GIORNI_LABEL}": giorni_label,
            "{TRAFRA}": tra,
            "{PRODOTTO}": PRODOTTO,
            "{ANNO}": str(dt.date.today().year),
            "{LOGO_IMG}": logo_img,
            "{CITTA_FOOTER}": citta_footer,
            "{EMAIL_FOOTER}": email_footer,
        }
        html = template
        for k, v in valori.items():
            html = html.replace(k, v)
    return oggetto, testo, html

# =====================================================================
# 5) INVIO (una sola connessione, con retry e pausa)
# =====================================================================

class Sender:
    """Gestisce la connessione SMTP: si connette al primo invio e resta
    aperta; se cade si riconnette; ogni email ha 3 tentativi."""

    def __init__(self):
        self.server = None

    def _connetti(self):
        if self.server is None:
            self.server = smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=30)
            self.server.login(MITTENTE_EMAIL, MITTENTE_PASSWORD)

    def _chiudi(self):
        try:
            if self.server is not None:
                self.server.quit()
        except Exception:
            pass
        self.server = None

    def invia(self, destinatario, oggetto, testo, html=None, logo=None):
        ultimo_errore = None
        for tentativo in range(1, 4):
            try:
                self._connetti()
                msg = EmailMessage()
                msg["From"] = formataddr((NOME_CENTRO, MITTENTE_EMAIL))
                msg["To"] = destinatario
                msg["Subject"] = oggetto
                if RISPOSTA_A:
                    msg["Reply-To"] = RISPOSTA_A
                msg.set_content(testo)
                if html:
                    msg.add_alternative(html, subtype="html")
                    if logo:
                        # Content-ID DEVE avere le parentesi angolari (<logo>):
                        # e' lo standard RFC 2392 per src="cid:logo". Senza
                        # parentesi Gmail e altri client NON mostrano il logo.
                        msg.get_payload()[1].add_related(
                            logo, maintype="image", subtype="png", cid="<logo>")
                self.server.send_message(msg)
                return None
            except (smtplib.SMTPException, OSError) as e:
                ultimo_errore = e
                self._chiudi()
                if tentativo < 3:
                    time.sleep(3 * tentativo)
        return ultimo_errore

    def chiudi(self):
        self._chiudi()

# =====================================================================
# 6) BACKUP AUTOMATICO DEL FILE CLIENTI
# =====================================================================

def backup_excel(percorso_excel):
    try:
        cartella = os.path.join(BASE, "backup")
        os.makedirs(cartella, exist_ok=True)
        nome = f"clienti_{dt.date.today().isoformat()}_{dt.datetime.now().strftime('%H%M')}.xlsx"
        shutil.copy2(percorso_excel, os.path.join(cartella, nome))
        # tieni solo i 14 piu' recenti
        files = sorted(
            (os.path.join(cartella, f) for f in os.listdir(cartella)),
            key=os.path.getmtime, reverse=True)
        for vecchio in files[14:]:
            os.remove(vecchio)
        return True
    except Exception as e:
        print(f"  (backup non riuscito: {e} - continuo comunque)")
        return False

# =====================================================================
# 7) RIEPILOGO GIORNALIERO AL TITOLARE
# =====================================================================

def invia_riepilogo(sender, inviati, errori, esclusi, scadute, dry_run):
    if not CFG["notifica_titolare"]:
        return
    if not inviati and not errori:
        return                      # niente da segnalare: zero email al titolare
    oggi = dt.date.today().strftime("%d/%m/%Y")
    destinatario = CFG["email_titolare"] or MITTENTE_EMAIL

    righe_html = ""
    for i in inviati:
        righe_html += (f"<tr><td style='padding:6px 10px;border-bottom:1px solid #eee;'>{i['targa']}</td>"
                       f"<td style='padding:6px 10px;border-bottom:1px solid #eee;'>{i['email']}</td>"
                       f"<td style='padding:6px 10px;border-bottom:1px solid #eee;'>{i['desc']}</td></tr>")
    righe_err = ""
    for e in errori:
        righe_err += (f"<tr><td style='padding:6px 10px;border-bottom:1px solid #eee;'>{e['targa']}</td>"
                      f"<td style='padding:6px 10px;border-bottom:1px solid #eee;color:#C0392B;'>{e['errore']}</td></tr>")

    html = f"""<div style="font-family:Arial,sans-serif;color:#333;max-width:600px;">
  <h2 style="color:{CFG['colore']};margin-bottom:4px;">Riepilogo promemoria - {oggi}</h2>
  <p style="margin-top:0;">Generato automaticamente da {PRODOTTO}</p>
  <p><b>{len(inviati)}</b> promemoria inviati &nbsp;|&nbsp;
     <b>{len(errori)}</b> errori &nbsp;|&nbsp;
     <b>{esclusi}</b> clienti esclusi &nbsp;|&nbsp;
     <b>{scadute}</b> revisioni gia' scadute</p>
  {f"<h3>Invii di oggi</h3><table style='border-collapse:collapse;width:100%;'><tr style='background:{CFG['colore']};color:#fff;'><th style='padding:6px 10px;text-align:left;'>Targa</th><th style='padding:6px 10px;text-align:left;'>Cliente</th><th style='padding:6px 10px;text-align:left;'>Promemoria</th></tr>{righe_html}</table>" if righe_html else ""}
  {f"<h3>Errori da controllare</h3><table style='border-collapse:collapse;width:100%;'><tr style='background:#C0392B;color:#fff;'><th style='padding:6px 10px;text-align:left;'>Targa</th><th style='padding:6px 10px;text-align:left;'>Errore</th></tr>{righe_err}</table>" if righe_err else ""}
  <p style="font-size:12px;color:#888;">Se un cliente risponde STOP: metti Attivo = NO nel file Excel.</p>
</div>"""
    testo = (f"Riepilogo promemoria del {oggi}\n"
             f"Inviati: {len(inviati)} | Errori: {len(errori)} | "
             f"Esclusi: {esclusi} | Scadute: {scadute}")
    oggetto = (f"[{PRODOTTO}] Riepilogo del {oggi} - "
               f"{len(inviati)} promemoria inviati")
    if dry_run:
        print("\n  [RIEPILOGO TITOLARE - solo PROVA]")
        print(f"  Oggetto: {oggetto}")
        print(f"  A: {destinatario}")
        print(f"  Inviati: {len(inviati)}, Errori: {len(errori)}, Esclusi: {esclusi}, Scadute: {scadute}")
        return
    errore = sender.invia(destinatario, oggetto, testo, html)
    if errore:
        print(f"  ATTENZIONE: riepilogo al titolare non inviato ({errore})")

# =====================================================================
# 8) PROGRAMMA PRINCIPALE
# =====================================================================

def trova_tipo(giorni):
    for tipo, (da, a) in CFG["finestre"].items():
        if da <= giorni <= a:
            return tipo
    return None


def crea_excel_se_manca(percorso_excel):
    """Se il file clienti non c'e' ma c'e' l'esempio nel pacchetto,
    lo copia: cosi' aggiornare il programma non cancella mai i clienti
    (lo zip NON contiene il file clienti del centro, solo l'esempio)."""
    if os.path.exists(percorso_excel):
        return
    esempio = os.path.join(BASE, "clienti_ESEMPIO.xlsx")
    if not os.path.isfile(esempio):
        return
    try:
        shutil.copy2(esempio, percorso_excel)
        print(f"File clienti non trovato: l'ho creato dall'esempio "
              f"({os.path.basename(percorso_excel)}).\n")
    except Exception as e:
        print(f"AVVISO: non riesco a creare il file clienti dall'esempio ({e}).\n")


def main():
    parser = argparse.ArgumentParser(description="RevMemo - promemoria revisioni")
    parser.add_argument("--dry-run", action="store_true",
                        help="modalita' prova: mostra senza inviare")
    parser.add_argument("--file", default=None,
                        help="file clienti alternativo")
    args = parser.parse_args()

    percorso_excel = args.file or os.path.join(BASE, FILE_EXCEL)
    crea_excel_se_manca(percorso_excel)

    print("=" * 62)
    print(f"  {PRODOTTO} - Promemoria revisioni",
          "(MODALITA' PROVA - nessuna email partira')" if args.dry_run else "")
    print("=" * 62)
    print(f"Oggi: {dt.date.today().strftime('%d/%m/%Y')}")
    print(f"Centro: {NOME_CENTRO}   |   Mittente: {MITTENTE_EMAIL or '(non impostato)'}")
    print(f"File clienti: {percorso_excel}")
    if args.dry_run and not MITTENTE_PASSWORD:
        print("Nota: password non ancora inserita in config.ini - la prova funziona lo stesso")
    print()

    if not args.dry_run:
        if not MITTENTE_EMAIL or not MITTENTE_PASSWORD:
            print("ERRORE: compila config.ini sezione [smtp] prima dell'invio vero:")
            print("  mittente = la casella da cui inviare")
            print("  password_app = la Password per app di Google (16 caratteri)")
            print("Guida: ISTRUZIONI_CENTRO.md punto 3")
            sys.exit(1)
        backup_excel(percorso_excel)

    clienti = leggi_clienti(percorso_excel)
    print(f"Trovati {len(clienti)} clienti nel file.\n")

    sender = Sender()
    inviati, errori = [], []
    esclusi = scadute = saltati_log = 0
    ogggi = dt.date.today()

    for c in clienti:
        targa = c["targa"] or "(targa mancante)"
        riga_lbl = f"riga {c['riga']:>2} | {(c['nome'] or '(senza nome)')[:18]:<18} | {targa:<8} | "

        if not c["attivo"]:
            print(f"  - {riga_lbl}cliente NON attivo -> saltato")
            esclusi += 1
            continue
        if c["fatta"]:
            print(f"  - {riga_lbl}revisione GIA' FATTA -> nessuna email")
            esclusi += 1
            continue
        if not c["email"] or "@" not in c["email"]:
            print(f"  - {riga_lbl}email mancante o non valida -> saltato")
            esclusi += 1
            continue
        if c["scadenza"] is None:
            print(f"  - {riga_lbl}scadenza non leggibile (scrivi es. 09/2026) -> saltato")
            esclusi += 1
            continue

        giorni = (c["scadenza"] - ogggi).days
        if giorni < 0:
            print(f"  - {riga_lbl}revisione SCADUTA da {-giorni} gg -> nessun invio")
            scadute += 1
            esclusi += 1
            continue
        if giorni == 0:
            print(f"  - {riga_lbl}scade OGGI: ultimo giorno utile, nessun promemoria")
            esclusi += 1
            continue

        tipo = trova_tipo(giorni)
        if tipo is None:
            print(f"  - {riga_lbl}mancano {giorni:>3} gg -> non e' momento di inviare")
            esclusi += 1
            continue

        descrizione = CFG["etichette"][tipo]
        oggetto, testo, html = componi_email(c, giorni)

        scadenza_iso = c["scadenza"].isoformat()
        if gia_inviato(c["email"], c["targa"], tipo, scadenza_iso):
            print(f"  - {riga_lbl}promemoria '{descrizione}' GIA' INVIATO -> saltato")
            saltati_log += 1
            continue

        if args.dry_run:
            print(f"  > {riga_lbl}PROVA invio '{descrizione}' a {c['email']}")
            print(f"      Oggetto: {oggetto}")
            print(f"      Email HTML con logo: {'SI' if html else 'no (template mancante)'}")
            inviati.append({"targa": targa, "email": c["email"], "desc": descrizione})
            continue

        errore = sender.invia(c["email"], oggetto, testo, html, leggi_logo())
        if errore is None:
            registra_invio(c["email"], c["targa"], tipo, giorni, descrizione, scadenza_iso)
            print(f"  [OK] {riga_lbl}email inviata a {c['email']} ({descrizione})")
            inviati.append({"targa": targa, "email": c["email"], "desc": descrizione})
        else:
            print(f"  [X] {riga_lbl}ERRORE dopo 3 tentativi: {errore}")
            errori.append({"targa": targa, "errore": str(errore)})
        time.sleep(CFG["pausa"])

    if not args.dry_run:
        sender.chiudi()

    print()
    print("-" * 62)
    print(f"RIEPILOGO: {len(inviati)} inviate, {saltati_log} doppioni saltati, "
          f"{esclusi} esclusi, {len(errori)} errori.")
    if args.dry_run:
        invia_riepilogo(sender, inviati, errori, esclusi, scadute, dry_run=True)
        print("\nTutto corretto? Per l'invio vero: doppio click su avvia_invio.bat")
    else:
        invia_riepilogo(sender, inviati, errori, esclusi, scadute, dry_run=False)
        if inviati:
            print(f"Registro aggiornato: {percorso_log()}")
    print("-" * 62)
    return 0


if __name__ == "__main__":
    sys.exit(main())
