#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Controllo del pacchetto RevMemo (versione 2 - auto-riparante).
- Se c'e' la cartella runtime inclusa, TUTTI i test girano sul runtime
  (cosi' il Python installato sul PC non conta piu').
- Se manca config.ini, lo ricrea da solo.
- Riconosce i problemi tipici e spiega cosa fare.
Avvio: doppio click su controlla_pc.bat
"""

import configparser
import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

CARTELLA = os.path.dirname(os.path.abspath(__file__))
RUNTIME_PY = os.path.join(CARTELLA, "runtime", "python.exe")
OK, MALE, INFO, FATTO = "[OK]  ", "[!!]  ", "[i]   ", "[+*]  "
problemi = []
nel_runtime = False


def riga(sim, testo, dettaglio="", problema=False):
    global problemi
    if problema:
        problemi.append(testo)
    print(f"  {sim}{testo}" + (f"\n        {dettaglio}" if dettaglio else ""))


# ---------------------------------------------------------------
# 0) Se c'e' il runtime incluso e non lo stiamo usando, ripartiamo
#    con quello: i test valgono per il Python "giusto".
# ---------------------------------------------------------------
if (os.name == "nt"
        and os.path.isfile(RUNTIME_PY)
        and os.path.abspath(sys.executable).lower() != os.path.abspath(RUNTIME_PY).lower()
        and os.environ.get("REVMEMO_NEL_RUNTIME") != "1"):
    os.environ["REVMEMO_NEL_RUNTIME"] = "1"
    try:
        os.execv(RUNTIME_PY, [RUNTIME_PY, os.path.abspath(__file__)])
    except Exception:
        pass  # se il rilancio fallisce, prosegue con il Python corrente

nel_runtime = os.path.abspath(sys.executable).lower().startswith(
    os.path.abspath(os.path.join(CARTELLA, "runtime")).lower())
runtime_presente = os.path.isdir(os.path.join(CARTELLA, "runtime"))

print("=" * 62)
print("  CONTROLLO DEL PACCHETTO - RevMemo")
print("=" * 62)

# 1) Python usato
if nel_runtime:
    riga(OK, f"Girando sul runtime incluso - Python {sys.version.split()[0]}")
    riga(OK, "Niente da installare sul PC: tutto e' dentro il pacchetto")
else:
    riga(OK, f"Programma avviato - Python {sys.version.split()[0]}")

# 2) runtime incluso
if runtime_presente:
    if not nel_runtime:
        riga(INFO, "Cartella runtime presente (ma non usata in questo avvio)")
else:
    riga(MALE, "Manca la cartella runtime (il Python incluso nel pacchetto)",
         "Il pacchetto sul PC e' incompleto: copia/estraici TUTTA la cartella\n"
         "        RevMemo dal file zip scaricato, poi rifai il controllo", True)

# 3) file .py con nomi "riservati" (causano errori di import)
NOMI_VIETATI = {
    "decimal", "numbers", "math", "re", "operator", "contextvars", "abc",
    "functools", "inspect", "warnings", "datetime", "calendar", "time",
    "os", "sys", "shutil", "argparse", "configparser", "smtplib", "email",
    "socket", "ssl", "openpyxl", "string", "types", "io", "enum",
}
ombra = [f for f in os.listdir(CARTELLA)
         if f.lower().endswith(".py")
         and f[:-3].lower() in NOMI_VIETATI]
if ombra:
    riga(MALE, "File Python con nome 'riservato': " + ", ".join(ombra),
         "Questi file mandano in confusione il programma (errori di import\n"
         "        tipo 'cannot import name Decimal'). Rinominali o spostali fuori.", True)
else:
    riga(OK, "Nessun file con nome che disturba Python")

# 4) openpyxl
openpyxl_ok = False
errore_openpyxl = ""
try:
    import openpyxl  # noqa: F401
    openpyxl_ok = True
except ImportError as e:
    errore_openpyxl = str(e)
except Exception as e:
    errore_openpyxl = str(e)

if openpyxl_ok:
    riga(OK, "Libreria openpyxl presente e funzionante")
else:
    if "No module named" in errore_openpyxl and "openpyxl" in errore_openpyxl:
        if nel_runtime:
            riga(MALE, "openpyxl manca anche nel runtime: riscarica il pacchetto", "", True)
        elif runtime_presente:
            riga(INFO, "openpyxl non installata nel Python del PC (ma il runtime la include)")
        else:
            riga(MALE, "Manca openpyxl",
                 "Soluzione migliore: copia la cartella runtime dal pacchetto\n"
                 "        In alternativa: Prompt dei comandi -> pip install openpyxl", True)
    else:
        dettaglio = ("Questo problema e' tipico del Python 'nuovo' (3.14) del PC:\n"
                     "        la soluzione e' usare il runtime incluso nel pacchetto.\n"
                     "        Copia la cartella runtime accanto a questo file e rifai il controllo.")
        riga(MALE, f"openpyxl non si avvia con questo Python: {errore_openpyxl}",
             dettaglio, True)

# 5) config.ini (se manca, lo creo io)
CONFIG_DEFAULT = """\
; =====================================================================
;  RevMemo - CONFIGURAZIONE
;  Questo e' l'unico file da modificare. Il codice non si tocca mai.
; =====================================================================

[smtp]
; Casella Gmail DALLA CUI partono i promemoria
host = smtp.gmail.com
porta = 465
mittente = latuaemail@gmail.com
; Password per app di Google (16 caratteri) - vedi ISTRUZIONI_CENTRO.md punto 3
password_app =
; Risposte dei clienti (facoltativo): quando un cliente risponde a un
; promemoria, la risposta arriva nella casella mittente (quella sopra).
; Se vuoi che arrivino su un'altra casella, scrivila qui.
risposta_a =

[centro]
; Questi dati appaiono nelle email che ricevono i CLIENTI
nome = Centro Revisioni Esempio
citta = Milano
telefono = 02 1234567
email_contatti = info@centrorevisioni.it
; Colore principale delle email (codice esadecimale, es. blu #1F3B66, rosso #C8102E)
colore_tema = #1F3B66

[prodotto]
; Nome del prodotto mostrato in piccolo nel piede delle email.
; Cambiabile liberamente (white label).
nome = RevMemo

[promemoria]
; Finestre di invio: (da giorni-al giorni) di anticipo sulla scadenza
finestra_1mese = 26-30
finestra_10giorni = 8-10
finestra_1giorno = 1-2

[invii]
; Pausa tra un'email e l'altra (secondi) - 1.5 consigliato
pausa_secondi = 1.5
; Riepilogo giornaliero al titolare: SI / NO
notifica_titolare = SI
; Email del titolare (se vuota, usa la stessa del mittente)
email_titolare =
; Nomi dei file (non serve toccarli)
file_excel = clienti_revisioni.xlsx
file_log = reminders_log.txt

[sms]
; Modulo SMS predisposto: quando il centro decide di attivare gli SMS
; (gateway tipo Aruba/Skebby) questa sezione verra' collegata.
; Per ora lasciare NO.
attivo = NO
"""
percorso_config = os.path.join(CARTELLA, "config.ini")
if not os.path.exists(percorso_config):
    try:
        with open(percorso_config, "w", encoding="utf-8") as f:
            f.write(CONFIG_DEFAULT)
        riga(FATTO, "config.ini mancante: l'ho ricreato adesso (valori di partenza)",
             problema=True)
        print("        ATTENZIONE: se avevi gia' messo mittente o password, reinseriscili.")
    except Exception as e:
        riga(MALE, f"Manca config.ini e non riesco a crearlo: {e}", "", True)

cfg = configparser.ConfigParser(interpolation=None)
if os.path.exists(percorso_config):
    cfg.read(percorso_config, encoding="utf-8")
    riga(OK, "config.ini trovato")
    mittente = cfg.get("smtp", "mittente", fallback="").strip()
    password = cfg.get("smtp", "password_app", fallback="").strip()
    nome_centro = cfg.get("centro", "nome", fallback="")
    prodotto = cfg.get("prodotto", "nome", fallback="RevMemo")
    if mittente and mittente != "latuaemail@gmail.com":
        riga(OK, f"Mittente impostato: {mittente}")
    else:
        riga(MALE, "Mittente non impostato",
             "config.ini sezione [smtp]: metti la casella Gmail del centro", True)
    if password:
        riga(OK, "Password per app inserita")
    else:
        riga(MALE, "Manca la Password per app di Google",
             "config.ini sezione [smtp], riga password_app\n"
             "        Come crearla: ISTRUZIONI_CENTRO.md punto 3", True)
    riga(OK, f"Centro configurato: {nome_centro} (prodotto: {prodotto})")
else:
    mittente = password = ""

# 6) template e logo
if os.path.exists(os.path.join(CARTELLA, "template_email.html")):
    riga(OK, "Email HTML: template_email.html presente")
else:
    riga(MALE, "Manca template_email.html", "Le email partono comunque in formato testo", True)
if os.path.exists(os.path.join(CARTELLA, "logo.png")):
    dim_logo = os.path.getsize(os.path.join(CARTELLA, "logo.png"))
    if dim_logo < 100:
        riga(MALE, "logo.png e' vuoto o danneggiato",
             "Le email uscirebbero SENZA logo: ricopialo dal pacchetto\n"
             "        o mettici il logo del centro (PNG).", True)
    else:
        riga(OK, "Logo presente (logo.png) - sostituiscilo con quello del centro")
else:
    riga(INFO, "Nessun logo.png: le email usciranno senza logo")

# 6-bis) registro anti-doppioni (reminders_log.txt)
nome_log = "reminders_log.txt"
try:
    _cfg_tmp = configparser.ConfigParser(interpolation=None)
    _cfg_tmp.read(percorso_config, encoding="utf-8")
    nome_log = _cfg_tmp.get("invii", "file_log", fallback="reminders_log.txt")
except Exception:
    pass
percorso_log_chk = os.path.join(CARTELLA, nome_log)
if os.path.exists(percorso_log_chk):
    try:
        with open(percorso_log_chk, "r", encoding="utf-8") as f:
            n_invii = sum(1 for l in f if l.strip())
        riga(OK, f"Registro anti-doppioni presente: {nome_log} ({n_invii} invii registrati)",
             "Se e' una prova e vuoi RE-inviare le stesse email: cancella\n"
             "        quel file, si ricrea da solo al prossimo invio.")
    except Exception:
        riga(OK, f"Registro anti-doppioni presente: {nome_log}")
else:
    riga(OK, f"Registro anti-doppioni assente ({nome_log}): si crea da solo al primo invio")

# 7) file clienti + intestazioni (solo se openpyxl funziona)
percorso_excel = os.path.join(CARTELLA, "clienti_revisioni.xlsx")
if not os.path.exists(percorso_excel):
    esempio_excel = os.path.join(CARTELLA, "clienti_ESEMPIO.xlsx")
    if os.path.isfile(esempio_excel):
        try:
            import shutil as _shutil
            _shutil.copy2(esempio_excel, percorso_excel)
            riga(FATTO, "clienti_revisioni.xlsx mancante: l'ho creato dall'esempio",
                 "Adesso ci sono i clienti di prova: sostituiscili con i tuoi.", True)
        except Exception as e:
            riga(MALE, f"Manca clienti_revisioni.xlsx e non riesco a crearlo: {e}", "", True)
    else:
        riga(MALE, "Manca clienti_revisioni.xlsx", "Mettilo accanto a revmemo.py", True)
elif openpyxl_ok:
    riga(OK, "File clienti trovato: clienti_revisioni.xlsx")
    try:
        from openpyxl import load_workbook
        wb = load_workbook(percorso_excel, read_only=True)
        ws = wb["Clienti"] if "Clienti" in wb.sheetnames else wb.active
        intestazioni_attese = ["nome cliente", "email", "targa",
                               "scadenza revisione", "revisione fatta", "attivo"]
        trovate = []
        for col in range(1, 12):
            v = ws.cell(row=4, column=col).value
            if v is not None:
                trovate.append(str(v).strip().lower())
        mancano = [h for h in intestazioni_attese if h not in trovate]
        if mancano:
            riga(MALE, f"Nell'Excel mancano le colonne: {', '.join(mancano)}",
                 "La riga 4 deve contenere le intestazioni originali", True)
        else:
            riga(OK, "Colonne dell'Excel corrette (riga 4)")
        wb.close()
    except Exception as e:
        dettaglio = ""
        if "Permission denied" in str(e) or "Errno 13" in str(e):
            dettaglio = ("Il file Excel risulta APERTO in Excel (o bloccato da OneDrive):\n"
                         "        chiudi Excel, aspetta qualche secondo e rifai il controllo.")
        riga(MALE, f"Non riesco a leggere l'Excel: {e}", dettaglio, True)
else:
    riga(MALE, "Non posso verificare l'Excel (openpyxl non funziona: vedi sopra)", "", True)

# 8) riepilogo e consigli
print()
print("-" * 62)
if not runtime_presente:
    print("  COSA FARE ADESSO (pacchetto incompleto):")
    print("   1. Apri il file RevMemo_v3.zip che avevi scaricato")
    print("   2. Tasto destro -> 'Estrai tutto'")
    print("   3. La cartella RevMemo deve contenere ANCHE la cartella")
    print("      'runtime': copiala qui accanto a revmemo.py")
    print("   4. Rifai il controllo (controlla_pc.bat)")
    print()
if problemi:
    print(f"  >> {len(problemi)} cosa/e da sistemare prima dell'invio vero.")
    print("  >> La modalita' PROVA (avvia_prova.bat) funziona gia' da subito.")
else:
    print("  >> TUTTO PRONTO! Fai la prova con avvia_prova.bat,")
    print("  >> poi avvia_invio.bat per le email vere.")
input("\n  Premi INVIO per chiudere...")
