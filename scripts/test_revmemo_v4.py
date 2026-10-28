# -*- coding: utf-8 -*-
"""
Test RevMemo v4 — verifica completa senza toccare il pacchetto reale:
 1. parola STOP riconosciuta in tutte le forme (stop/Stop/STOP/frasi)
 2. host IMAP derivato da SMTP
 3. Excel esempio creato DA CODICE (8 clienti, colonne giuste, SÌ/NO leggibili)
 4. segna_stop_nel_excel: Attivo = STOP + log + idempotenza
 5. estrai_da_e_testo: testo semplice, HTML, nessuno STOP
 6. controlla_risposte con GMAIL FINTA (IMAP simulato): dry-run e reale
 7. main() completo in --dry-run
 8. controlla_pc.py gira e crea config.ini
 9. il pacchetto vero resta PRISTINO (nessun file generato dentro)
"""
import email as em
import email.message
from email.message import EmailMessage
import io
import os
import shutil
import subprocess
import sys
import tempfile
from contextlib import redirect_stdout

PACCHETTO = "/home/z/my-project/download/RevMemo"

ok_count = 0
fail_count = 0


def check(nome, condizione, extra=""):
    global ok_count, fail_count
    if condizione:
        ok_count += 1
        print(f"  [OK]   {nome}")
    else:
        fail_count += 1
        print(f"  [FALLITO] {nome} {extra}")


# ------------------------------------------------------------------
# Preparo una copia di lavoro in una cartella temporanea
# ------------------------------------------------------------------
tmp = tempfile.mkdtemp(prefix="revmemo_test_")
for f in ("revmemo.py", "template_email.html", "logo.png"):
    shutil.copy2(os.path.join(PACCHETTO, f), os.path.join(tmp, f))

CONFIG_TEST = """\
[smtp]
host = smtp.gmail.com
porta = 465
mittente = centro.prova.revmemo@gmail.com
password_app = aaaa bbbb cccc dddd
risposta_a =

[centro]
nome = Centro Revisioni Test
citta = Milano
telefono = 02 1234567
email_contatti = info@test.it
colore_tema = #1F3B66

[prodotto]
nome = RevMemo

[promemoria]
finestra_1mese = 26-30
finestra_10giorni = 8-10
finestra_1giorno = 1-2

[invii]
pausa_secondi = 0
notifica_titolare = SI
email_titolare =
file_excel = clienti_revisioni.xlsx
file_log = reminders_log.txt

[sms]
attivo = NO
"""
with open(os.path.join(tmp, "config.ini"), "w", encoding="utf-8") as f:
    f.write(CONFIG_TEST)

sys.path.insert(0, tmp)
import revmemo  # noqa: E402

print("== 1) Parola STOP ==")
check("STOP maiuscolo", revmemo.parola_stop_presente("STOP"))
check("Stop misto", revmemo.parola_stop_presente("Stop"))
check("stop minuscolo", revmemo.parola_stop_presente("stop"))
check("con punto esclamativo", revmemo.parola_stop_presente("STOP!"))
check("dentro una frase", revmemo.parola_stop_presente("Per favore STOP grazie"))
check("via email: STOP", revmemo.parola_stop_presente("potete fermarvi? STOP\ngrazie"))
check("'stoppa' NON conta", not revmemo.parola_stop_presente("stoppa"))
check("'stoppare' NON conta", not revmemo.parola_stop_presente("voglio stoppare"))
check("testo vuoto", not revmemo.parola_stop_presente(""))

print("== 2) Host IMAP ==")
check("smtp -> imap", revmemo.host_imap_da("smtp.gmail.com") == "imap.gmail.com")
check("vuoto -> gmail", revmemo.host_imap_da("") == "imap.gmail.com")

print("== 3) Excel creato da codice ==")
percorso_excel = os.path.join(tmp, "clienti_revisioni.xlsx")
revmemo.crea_excel_se_manca(percorso_excel)
check("file creato", os.path.exists(percorso_excel))
clienti = revmemo.leggi_clienti(percorso_excel)
check("8 clienti di esempio", len(clienti) == 8, f"(trovati {len(clienti)})")
check("tutte le scadenze valide", all(c["scadenza"] is not None for c in clienti))
check("colonne riconosciute", revmemo.individua_colonne(
    __import__("openpyxl").load_workbook(percorso_excel)["Clienti"])[0].get("attivo") == 8)
check("SÌ letto come attivo", clienti[0]["attivo"] is True)
check("NO letto come non attivo", clienti[6]["attivo"] is False)
check("Fatta=SÌ blocca", clienti[5]["fatta"] is True)

print("== 4) segna_stop_nel_excel ==")
toccate = revmemo.segna_stop_nel_excel(percorso_excel, "mario.rossi@email.it")
check("1 riga aggiornata", len(toccate) == 1, str(toccate))
check("riga 5", toccate and toccate[0][0] == 5, str(toccate))
from openpyxl import load_workbook  # noqa: E402
ws_chk = load_workbook(percorso_excel)["Clienti"]
check("cella Attivo = STOP", ws_chk.cell(row=5, column=8).value == "STOP")
check("log con riga STOP", os.path.exists(os.path.join(tmp, "reminders_log.txt"))
      and "STOP" in open(os.path.join(tmp, "reminders_log.txt")).read())
check("seconda volta: niente (gia' pausa)", revmemo.segna_stop_nel_excel(
    percorso_excel, "mario.rossi@email.it") == [])
check("email sconosciuta: niente", revmemo.segna_stop_nel_excel(
    percorso_excel, "sconosciuto@nessuno.it") == [])
clienti = revmemo.leggi_clienti(percorso_excel)
check("mario ora non attivo con testo STOP",
      clienti[0]["attivo"] is False and clienti[0]["attivo_testo"] == "STOP")

print("== 5) estrai_da_e_testo ==")


def msg_raw(da, oggetto, corpo):
    m = EmailMessage()
    m["From"] = da
    m["To"] = "centro@gmail.com"
    m["Subject"] = oggetto
    m.set_content(corpo)
    return m.as_bytes()


mitt, ogg, corpo = revmemo.estrai_da_e_testo(
    msg_raw("Mario Rossi <MARIO.ROSSI@email.it>", "Re: promemoria", "STOP per favore\n"))
check("mittente estratto in minuscolo", mitt == "mario.rossi@email.it", mitt)
check("STOP nel testo semplice", revmemo.parola_stop_presente(ogg + "\n" + corpo))
RAW_HTML = (b"From: Solo HTML <solohtml@email.it>\r\nTo: c@b.it\r\n"
            b"Subject: Re: promemoria\r\nMIME-Version: 1.0\r\n"
            b"Content-Type: text/html; charset=utf-8\r\n\r\n"
            b"<html><body><p>Ciao, <b>STOP</b> per favore</p></body></html>\r\n")
mitt, ogg, corpo = revmemo.estrai_da_e_testo(RAW_HTML)
check("STOP dentro HTML", revmemo.parola_stop_presente(ogg + "\n" + corpo))
mitt, ogg, corpo = revmemo.estrai_da_e_testo(
    msg_raw("Anna <anna.bianchi@email.it>", "Re: promemoria", "Arrivo lunedi'\n"))
check("risposta normale: nessuno STOP", not revmemo.parola_stop_presente(ogg + "\n" + corpo))

print("== 6) controlla_risposte con Gmail finta ==")
MESSAGGI = {
    1: msg_raw("Luigi Verdi <luigi.verdi@email.it>", "Re: promemoria",
               "Per favore STOP, ho venduto l'auto\n"),
    2: msg_raw("Anna Bianchi <anna.bianchi@email.it>", "Re: promemoria",
               "Arrivo lunedi' con la macchina\n"),
    3: msg_raw("Esterno <esterno@libero.it>", "Ciao", "STOP subito\n"),
}


class FakeIMAP:
    istanza = None

    def __init__(self, host, port, timeout=None):
        self.host, self.port = host, port
        self.marked = []
        self.fetch_full = []
        self.login_ok = False
        FakeIMAP.istanza = self

    def login(self, u, p):
        self.login_ok = (u == "centro.prova.revmemo@gmail.com")
        return ("OK", [b""])

    def select(self, box):
        return ("OK", [b"3"])

    def search(self, charset, *crit):
        mesi = ("Jan", "Feb", "Mar", "Apr", "May", "Jun",
                "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
        assert crit[0] == "SINCE"
        assert crit[1].split("-")[1] in mesi, f"mese non inglese: {crit[1]}"
        return ("OK", [b"1 2 3"])

    def fetch(self, num, spec):
        # come Gmail vera: righe di intestazione con fine riga \r\n
        raw = MESSAGGI[int(num)].replace(b"\n", b"\r\n")
        if "HEADER" in spec:
            cut = raw.find(b"\r\n\r\n")
            hdr = raw[:cut] if cut != -1 else raw
            return ("OK", [(num + b" (RFC822.HEADER {%d}" % len(hdr), hdr), b")"])
        self.fetch_full.append(num)
        return ("OK", [(num + b" (RFC822 {%d}" % len(raw), raw), b")"])

    def store(self, num, flags, flag):
        self.marked.append(num)
        return ("OK", [b""])

    def logout(self):
        return ("BYE", [b""])


imaplib_originale = revmemo.imaplib.IMAP4_SSL
revmemo.imaplib.IMAP4_SSL = FakeIMAP

# 6a) dry-run: NON tocca niente
buf = io.StringIO()
with redirect_stdout(buf):
    eventi = revmemo.controlla_risposte(percorso_excel, dry_run=True)
check("dry-run: 1 evento prova", len(eventi) == 1, str(eventi))
check("dry-run: evento di luigi", eventi and eventi[0]["email"] == "luigi.verdi@email.it")
ws_chk = load_workbook(percorso_excel)["Clienti"]
check("dry-run: Excel NON modificato", ws_chk.cell(row=6, column=8).value == "SÌ")
check("dry-run: nessuna email segnata letta", FakeIMAP.istanza.marked == [])
check("email esterno mai letta interamente", b"3" not in FakeIMAP.istanza.fetch_full)

# 6b) invio vero: registra lo STOP
buf = io.StringIO()
with redirect_stdout(buf):
    eventi = revmemo.controlla_risposte(percorso_excel, dry_run=False)
uscita = buf.getvalue()
check("reale: 1 evento", len(eventi) == 1, str(eventi))
check("reale: riga 6 di luigi", eventi and eventi[0]["righe"] == [(6, "Luigi Verdi", "EF456GH")],
      str(eventi))
ws_chk = load_workbook(percorso_excel)["Clienti"]
check("reale: cella H6 = STOP", ws_chk.cell(row=6, column=8).value == "STOP")
check("reale: email segnata come letta", FakeIMAP.istanza.marked == [b"1"])
check("schermo mostra lo STOP", "[STOP] luigi.verdi@email.it" in uscita)
check("risposta normale NON segnata", FakeIMAP.istanza.marked == [b"1"])

# 6c) rilancio: luigi e' gia' in pausa -> zero nuovi eventi
buf = io.StringIO()
with redirect_stdout(buf):
    eventi = revmemo.controlla_risposte(percorso_excel, dry_run=False)
check("rilancio: zero nuovi eventi", eventi == [], str(eventi))

revmemo.imaplib.IMAP4_SSL = imaplib_originale

print("== 7) main() completo in PROVA ==")
sys.argv = ["revmemo.py", "--dry-run"]
buf = io.StringIO()
with redirect_stdout(buf):
    codice = revmemo.main()
uscita = buf.getvalue()
check("main ritorna 0", codice == 0)
check("dice MODALITA' PROVA", "MODALITA' PROVA" in uscita)
check("clienti con STOP esclusi", "ha chiesto STOP" in uscita)
check("Fatta=SÌ bloccata", "GIA' FATTA" in uscita)
check("riepilogo finale", "RIEPILOGO:" in uscita)

print("== 8) controlla_pc.py ==")
tmp2 = tempfile.mkdtemp(prefix="revmemo_ctrl_")
shutil.copy2(os.path.join(PACCHETTO, "controlla_pc.py"), os.path.join(tmp2, "controlla_pc.py"))
proc = subprocess.run([sys.executable, "controlla_pc.py"], cwd=tmp2, input="\n",
                      capture_output=True, text=True, timeout=120)
check("controlla_pc esce senza errori", proc.returncode == 0, proc.stderr[-300:])
check("controlla_pc stampa il titolo", "CONTROLLO DEL PACCHETTO" in proc.stdout)
check("controlla_pc crea config.ini", os.path.exists(os.path.join(tmp2, "config.ini")))

print("== 9) il pacchetto vero resta pulito ==")
for f in ("config.ini", "clienti_revisioni.xlsx", "reminders_log.txt", "__pycache__"):
    check(f"pacchetto senza {f}", not os.path.exists(os.path.join(PACCHETTO, f)))

print()
print("=" * 50)
print(f"RISULTATO: {ok_count} OK, {fail_count} FALLITI")
print("=" * 50)
sys.exit(1 if fail_count else 0)
