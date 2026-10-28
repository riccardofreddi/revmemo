# -*- coding: utf-8 -*-
"""Crea un Excel di prova temporaneo con clienti che cadono nelle finestre
di invio (26-30 gg, 8-10 gg, 1-2 gg) per verificare i trigger. Da cancellare."""
import datetime as dt
from openpyxl import Workbook

OGGI = dt.date.today()
OUT = "/home/z/my-project/download/promemoria-revisioni/_test_trigger.xlsx"

HEADERS = ["Nome Cliente", "Email", "Targa", "Scadenza Revisione", "Revisione Fatta", "Attivo"]
RIGHE = [
    # dentro le 3 finestre -> devono partire
    ("A - prova 1 mese", "a@test.it", "AA111AA", OGGI + dt.timedelta(days=30), "NO", "SÌ"),
    ("B - prova 10 gg",  "b@test.it", "BB222BB", OGGI + dt.timedelta(days=9),  "NO", "SÌ"),
    ("C - prova 1 gg",   "c@test.it", "CC333CC", OGGI + dt.timedelta(days=2),  "NO", "SÌ"),
    # dentro la finestra ma revisione fatta -> deve essere bloccato
    ("D - fatta ma vicino", "d@test.it", "DD444DD", OGGI + dt.timedelta(days=28), "SÌ", "SÌ"),
    # fuori finestre -> niente
    ("E - 25 gg", "e@test.it", "EE555EE", OGGI + dt.timedelta(days=25), "NO", "SÌ"),
    ("F - 7 gg",  "f@test.it", "FF666FF", OGGI + dt.timedelta(days=7),  "NO", "SÌ"),
    # scadenza "scade oggi" (giorni=0)
    ("G - scade oggi", "g@test.it", "GG777GG", OGGI, "NO", "SÌ"),
]

wb = Workbook()
ws = wb.active
ws.title = "Clienti"
for col, h in enumerate(HEADERS, 2):
    ws.cell(row=4, column=col, value=h)
for i, riga in enumerate(RIGHE):
    r = 5 + i
    for col, v in enumerate(riga, 2):
        cell = ws.cell(row=r, column=col, value=v)
        if col == 5 and isinstance(v, dt.date):
            cell.number_format = "DD/MM/YYYY"
wb.save(OUT)
print("Creato", OUT)
