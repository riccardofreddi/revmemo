# -*- coding: utf-8 -*-
"""Test RevMemo v4.3.1 — "solo le CELLE colorate; STOP non e' un'opzione
da nessuna parte" (correzione della v4.3 che colorava la riga intera e
metteva STOP nel menu a tendina).

Verifica:
 1. converti_scadenza: le celle-data di Excel (es. "gen-26") valgono
    MESE/ANNO con fine mese (31/01/2026, non 01/01/2026)
 2. crea_excel_se_manca: SOLO le celle di Revisione Fatta e Attivo si
    colorano (SÌ verde, NO rosso, STOP rosso scuro); NESSUNA tinta di
    riga; tendina SÌ,NO su entrambe (STOP mai nel menu); colonna
    Scadenza in formato TESTO; intestazioni bloccate
 3. migliora_excel su un file "vecchio stile" E su uno toccato dalla
    v4.3 sbagliata (STOP nel menu): sistema il menu, aggiunge i colori,
    riscrive le scadenze diventate data, NON tocca i dati, e rilanciato
    due volte non crea doppioni
 4. leggi_clienti sul file migliorato: tutto leggibile come prima
 5. CLI --colora-excel: esce 0 e stampa "Fatto"
"""
import io
import os
import shutil
import subprocess
import sys
import tempfile
from contextlib import redirect_stdout

import datetime as dt

PACCHETTO = "/home/z/my-project/download/RevMemo"

ok_count = 0
fail_count = 0


def check(nome, condizione, extra=""):
    global ok_count, fail_count
    if condizione:
        ok_count += 1
        print(f"  [OK]   {nome}")
    else:
        fail_count += 1
        print(f"  [FALLITO] {nome} {extra}")


tmp = tempfile.mkdtemp(prefix="revmemo_v43_")
for f in ("revmemo.py", "template_email.html", "logo.png"):
    shutil.copy2(os.path.join(PACCHETTO, f), os.path.join(tmp, f))

CONFIG_TEST = """\
[smtp]
host = smtp.gmail.com
porta = 465
mittente = centro.prova.revmemo@gmail.com
password_app = aaaabbbbccccdddd
risposta_a =

[centro]
nome = Centro Revisioni Test
citta = Milano
telefono = 02 1234567
email_contatti = info@test.it
colore_tema = #1F3B66

[prodotto]
nome = RevMemo

[promemoria]
finestra_1mese = 26-30
finestra_10giorni = 8-10
finestra_1giorno = 1-2

[invii]
pausa_secondi = 0
notifica_titolare = SI
email_titolare = titolare@test.it
file_excel = clienti_revisioni.xlsx
file_log = reminders_log.txt

[sms]
attivo = NO
"""
with open(os.path.join(tmp, "config.ini"), "w", encoding="utf-8") as f:
    f.write(CONFIG_TEST)

sys.path.insert(0, tmp)
import revmemo  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

print("\n== 1) converti_scadenza: celle-data di Excel (gen-26) ==")
check("data Excel 01/01/2026 -> 31/01/2026 fine mese",
      revmemo.converti_scadenza(dt.datetime(2026, 1, 1)) == (dt.date(2026, 1, 31), True))
check("testo '01/2026' -> 31/01/2026",
      revmemo.converti_scadenza("01/2026") == (dt.date(2026, 1, 31), True))
check("testo '9/2026' -> 30/09/2026",
      revmemo.converti_scadenza("9/2026") == (dt.date(2026, 9, 30), True))
check("testo '15/03/2026' -> data esatta",
      revmemo.converti_scadenza("15/03/2026") == (dt.date(2026, 3, 15), False))
check("data Excel 15/03/2026 (giorno != 1) -> data esatta",
      revmemo.converti_scadenza(dt.datetime(2026, 3, 15)) == (dt.date(2026, 3, 15), False))
check("cella vuota -> None",
      revmemo.converti_scadenza("") == (None, False))


def regole_di(ws):
    trovate = []
    for rng in ws.conditional_formatting:
        for r in rng.rules:
            trovate.append((str(rng.sqref), r))
    return trovate


print("\n== 2) crea_excel_se_manca: solo celle colorate, menu senza STOP ==")
percorso_excel = os.path.join(tmp, "clienti_revisioni.xlsx")
buf = io.StringIO()
with redirect_stdout(buf):
    revmemo.crea_excel_se_manca(percorso_excel)
wb = load_workbook(percorso_excel)
ws = wb["Clienti"]
check("intestazioni bloccate (A5)", ws.freeze_panes == "A5")
dvs = ws.data_validations.dataValidation
check("tendina presente", len(dvs) >= 1)
check("tendina SOLO SÌ,NO (STOP mai nel menu)",
      all("STOP" not in (d.formula1 or "") and '"SÌ,NO"' in (d.formula1 or "")
          for d in dvs))
regole = regole_di(ws)
check("5 regole di colore (2 su Fatta + 3 su Attivo)",
      len(regole) == 5, f"(trovate {len(regole)})")
stop_rules = [r for _, r in regole if r.formula and 'UPPER(H5)="STOP"' in r.formula[0]]
check("regola STOP presente (solo colore della cella Attivo)", len(stop_rules) == 1)
if stop_rules:
    fill_txt = str(stop_rules[0].dxf.fill) if stop_rules[0].dxf else ""
    check("STOP rosso scuro (C00000)", "C00000" in fill_txt, fill_txt)
si_rules = [r for _, r in regole if r.formula and "OR" in r.formula[0]
            and 'UPPER(H5)="SI"' in r.formula[0]]
check("SÌ accetta anche SI senza accento", len(si_rules) >= 1)
no_fatta = [r for _, r in regole if r.formula and 'UPPER(G5)="NO"' in r.formula[0]]
check("Revisione Fatta: NO rosso anche sulla sua colonna", len(no_fatta) == 1)
check("solo celle G e H colorate (niente tinte di riga)",
      all(s.startswith(("G5:", "H5:")) for s, _ in regole),
      str([s for s, _ in regole]))
check("colonna Scadenza in formato testo (@)",
      ws.cell(row=5, column=6).number_format == "@")
val_scad = str(ws.cell(row=5, column=6).value)
check("scadenza di esempio resta MESE/ANNO", val_scad.count("/") == 1, val_scad)
check("intestazione blu 1F3B66",
      str(ws.cell(row=4, column=2).fill.fgColor.rgb).endswith("1F3B66"))

print("\n== 3) migliora_excel su file 'vecchio stile' ==")
vecchio = os.path.join(tmp, "vecchio.xlsx")
wb2 = Workbook()
ws2 = wb2.active
ws2.title = "Clienti"
ws2["B1"] = "ELENCO CLIENTI"
for col, titolo in zip(range(2, 9), ("Nome Cliente", "Email", "Telefono", "Targa",
                                     "Scadenza Revisione", "Revisione Fatta", "Attivo")):
    ws2.cell(row=4, column=col, value=titolo)
dati = [
    ("Mario Rossi", "mario.rossi@email.it", "347 1122334", "AB123CD",
     dt.datetime(2026, 1, 1), "NO", "SI"),
    ("Luigi Verdi", "luigi.verdi@email.it", None, "EF456GH",
     "03/2026", "NO", "STOP"),
    ("Anna Bianchi", "anna.bianchi@email.it", None, "JK789LM",
     dt.datetime(2026, 5, 15), "NO", "SÌ"),
]
for i, riga in enumerate(dati, start=5):
    for col, v in enumerate(riga, start=2):
        if v is not None:
            ws2.cell(row=i, column=col, value=v)
# menu con STOP come nella v4.3 sbagliata di stamattina: migliora_excel
# deve riportarlo a SÌ,NO
dv_vecchia = DataValidation(type="list", formula1='"SÌ,NO,STOP"', allow_blank=True,
                            showErrorMessage=True,
                            error="Scrivi SÌ oppure NO (usa il menu a tendina).")
ws2.add_data_validation(dv_vecchia)
dv_vecchia.add("H5:H300")
ws2.conditional_formatting.add("H5:H300", CellIsRule(
    operator="equal", formula=['"NO"'], fill=PatternFill("solid", fgColor="FFC7CE")))
wb2.save(vecchio)

_wb = load_workbook(vecchio)
_ws = _wb["Clienti"]
# confronto TUTTE le colonne tranne la Scadenza (col. 6): quella DEVE
# cambiare solo dove Excel aveva trasformato il testo in data
prima = tuple(tuple(_ws.cell(row=r, column=c).value for c in range(2, 9) if c != 6)
              for r in range(4, 8))
_wb.close()

buf2 = io.StringIO()
with redirect_stdout(buf2):
    esito = revmemo.migliora_excel(vecchio)
check("migliora_excel esito True", esito is True)
wb3 = load_workbook(vecchio)
ws3 = wb3["Clienti"]
check("scadenza data Excel -> testo '01/2026'",
      ws3.cell(row=5, column=6).value == "01/2026")
check("scadenza convertita in formato testo",
      ws3.cell(row=5, column=6).number_format == "@")
check("scadenza testo '03/2026' intatta",
      ws3.cell(row=6, column=6).value == "03/2026")
check("data completa 15/05/2026 non toccata",
      ws3.cell(row=7, column=6).value == dt.datetime(2026, 5, 15))
dvs3 = ws3.data_validations.dataValidation
check("tendina riportata a SÌ,NO (STOP tolto dal menu)",
      all('"SÌ,NO"' in (d.formula1 or "") and "STOP" not in (d.formula1 or "")
          for d in dvs3))
check("blocco scrittura tendina disattivato",
      all(not d.showErrorMessage for d in dvs3))
regole3 = regole_di(ws3)
check("5 regole di colore (vecchie rimosse, niente doppioni)",
      len(regole3) == 5, f"(trovate {len(regole3)})")
check("regola STOP presente anche qui (colore cella)",
      any(r.formula and 'UPPER(H5)="STOP"' in r.formula[0] for _, r in regole3))
check("niente tinte di riga ne' qui",
      all(s.startswith(("G5:", "H5:")) for s, _ in regole3))
check("intestazioni bloccate (A5)", ws3.freeze_panes == "A5")
check("larghezza colonna Email impostata",
      ws3.column_dimensions["C"].width == 28)
dopo = tuple(tuple(ws3.cell(row=r, column=c).value for c in range(2, 9) if c != 6)
             for r in range(4, 8))
check("I DATI non toccati (righe 4-7 identiche, tranne le scadenze da convertire)",
      dopo == prima)
wb3.close()

with redirect_stdout(io.StringIO()):
    esito2 = revmemo.migliora_excel(vecchio)
wb4 = load_workbook(vecchio)
regole4 = regole_di(wb4["Clienti"])
check("rilanciato 2 volte: sempre 5 regole (niente doppioni)",
      esito2 is True and len(regole4) == 5, f"(trovate {len(regole4)})")
wb4.close()

print("\n== 4) leggi_clienti sul file migliorato ==")
clienti = revmemo.leggi_clienti(vecchio)
check("3 clienti letti", len(clienti) == 3, f"(letti {len(clienti)})")
if len(clienti) == 3:
    mario, luigi, anna = clienti
    check("Mario: scadenza 31/01/2026 da mese",
          mario["scadenza"] == dt.date(2026, 1, 31) and mario["da_mese"])
    check("Mario: 'SI' senza accento = attivo", mario["attivo"] is True)
    check("Luigi: STOP = non attivo",
          luigi["attivo"] is False and luigi["attivo_testo"] == "STOP")
    check("Luigi: scadenza 31/03/2026",
          luigi["scadenza"] == dt.date(2026, 3, 31) and luigi["da_mese"])
    check("Anna: data completa 15/05/2026 rispettata",
          anna["scadenza"] == dt.date(2026, 5, 15) and not anna["da_mese"])

print("\n== 5) casi limite di migliora_excel ==")
with redirect_stdout(io.StringIO()):
    check("file inesistente -> False",
          revmemo.migliora_excel(os.path.join(tmp, "che_non_c'e.xlsx")) is False)
rotto = os.path.join(tmp, "rotto.xlsx")
_wbx = Workbook()
_wbx.active["A1"] = "non e' un foglio clienti"
_wbx.save(rotto)
with redirect_stdout(io.StringIO()):
    check("file senza colonne giuste -> False",
          revmemo.migliora_excel(rotto) is False)

print("\n== 6) CLI --colora-excel (sottoprocesso) ==")
p = subprocess.run([sys.executable, os.path.join(tmp, "revmemo.py"), "--colora-excel"],
                   capture_output=True, text=True, cwd=tmp, timeout=180)
check("CLI esce con codice 0", p.returncode == 0, p.stderr[-200:])
check("CLI stampa 'Fatto'", "Fatto" in p.stdout)

print(f"\nRISULTATO: {ok_count} OK, {fail_count} FALLITI")
sys.exit(1 if fail_count else 0)
