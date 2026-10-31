#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Prove della v4.14: versione unica, blocco invio con controllo del
processo VIVO, serratura condivisa del file clienti (pannello contro
motore), salvataggio atomico, STOP scritto sotto serratura."""
import importlib.util
import os
import subprocess
import sys
import tempfile
import time

PACCHETTO = os.environ.get("REVMEMO_PKG", "/home/z/my-project/download/RevMemo")
PANNELLO = os.path.join(PACCHETTO, "pannello.py")

spec = importlib.util.spec_from_file_location("pannello", PANNELLO)
p = importlib.util.module_from_spec(spec)
spec.loader.exec_module(p)
r = p.motore()                       # il motore importato dal pannello

cart = tempfile.mkdtemp(prefix="revmemo_v414_")
r.BASE = cart                        # lock/log/stop_gestiti nella cartella di prova
p.BASE = cart
PERCORSO_EXCEL_PROVA = os.path.join(cart, "clienti_revisioni.xlsx")
ok, ko = 0, []


def check(nome, cond):
    global ok
    if cond:
        ok += 1
        print(f"  [OK]   {nome}")
    else:
        ko.append(nome)
        print(f"  [FALL] {nome}")


def pid_morto():
    figlio = subprocess.Popen([sys.executable, "-c", "pass"])
    figlio.wait()
    return figlio.pid


def lock_vecchio(percorso, minuti=20):
    vecchio = time.time() - minuti * 60
    os.utime(percorso, (vecchio, vecchio))


print("== A) VERSIONE unica ==")
versione = getattr(r, "VERSIONE", "")
check("VERSIONE presente nel motore", bool(versione))
testo = open(os.path.join(PACCHETTO, "revmemo.py"), encoding="utf-8").read()
check(f"intestazione allineata alla versione ({versione})",
      f"versione {versione} (PRO)" in testo)
check("non esiste piu' la scritta 4.4 (PRO)", "versione 4.4 (PRO)" not in testo)
check("il banner mostra la versione", "v{VERSIONE}" in testo)

print("== B) _pid_vivo / _leggi_pid_lock ==")
check("pid di se stesso: vivo", r._pid_vivo(os.getpid()) is True)
check("pid 0: non vivo", r._pid_vivo(0) is False)
check("pid None: non vivo", r._pid_vivo(None) is False)
check("pid testo strano: non vivo", r._pid_vivo("abc") is False)
morto = pid_morto()
check("pid appena morto: non vivo", r._pid_vivo(morto) is False)

blocco = os.path.join(cart, "revmemo.lock")
open(blocco, "w").write("12345")
check("legge il pid dal blocco", r._leggi_pid_lock(blocco) == 12345)
open(blocco, "w").write("  ")
check("blocco vuoto: pid None", r._leggi_pid_lock(blocco) is None)
os.remove(blocco)
check("blocco mancante: pid None", r._leggi_pid_lock(blocco) is None)

print("== C) BloccoInvio con controllo del vivo ==")
b1 = r.BloccoInvio()
check("primo blocco: preso", b1.prendi() is True)
check("dentro c'e' il nostro pid", r._leggi_pid_lock(blocco) == os.getpid())
b2 = r.BloccoInvio()
check("secondo blocco mentre e' vivo: rifiutato", b2.prendi() is False)
lock_vecchio(blocco)
b3 = r.BloccoInvio()
check("blocco VECCHIO ma processo VIVO: NON rubato (v4.14)", b3.prendi() is False)
b1.rilascia()
check("rilasciato: il file non c'e' piu'", not os.path.exists(blocco))

open(blocco, "w").write(str(morto))
b4 = r.BloccoInvio()
check("blocco di un processo MORTO: ripreso SUBITO (v4.14)",
      b4.prendi() is True and r._leggi_pid_lock(blocco) == os.getpid())
b4.rilascia()

open(blocco, "w").write("   ")
lock_vecchio(blocco)
b5 = r.BloccoInvio()
check("blocco illeggibile e vecchio: ripreso", b5.prendi() is True)
b5.rilascia()

open(blocco, "w").write("   ")
b6 = r.BloccoInvio()
check("blocco illeggibile ma recente: rispetto (v4.14)", b6.prendi() is False)
os.remove(blocco)

print("== D) SerraturaExcel (pannello contro motore) ==")
serra = os.path.join(cart, "excel.lock")
s1 = r.SerraturaExcel(attesa=1)
check("serratura: presa", s1.prendi() is True)
check("nel file c'e' il nostro pid", r._leggi_pid_lock(serra) == os.getpid())
s2 = r.SerraturaExcel(attesa=1)
t0 = time.time()
check("seconda mentre occupata: rifiutata (dopo ~1 secondo)",
      s2.prendi() is False and 0.8 <= time.time() - t0 < 6)
s1.rilascia()
check("rilasciata: file sparito", not os.path.exists(serra))
s3 = r.SerraturaExcel(attesa=1)
check("dopo il rilascio: presa di nuovo", s3.prendi() is True)
s3.rilascia()

open(serra, "w").write(str(morto))
s4 = r.SerraturaExcel(attesa=1)
check("serratura di un processo MORTO: ripresa SUBITO (v4.14)", s4.prendi() is True)
s4.rilascia()

open(serra, "w").write(str(os.getpid()))
s5 = r.SerraturaExcel(attesa=1)
check("serratura di un processo VIVO: non rubata (v4.14)", s5.prendi() is False)
os.remove(serra)

open(serra, "w").write("   ")
lock_vecchio(serra)
s6 = r.SerraturaExcel(attesa=1)
check("serratura illeggibile e vecchia: ripresa", s6.prendi() is True)
s6.rilascia()

print("== E) salva_excel_sicuro (salvataggio atomico) ==")
from openpyxl import Workbook, load_workbook
wbp = Workbook()
wbp.active["A1"] = "uno"
r.salva_excel_sicuro(wbp, PERCORSO_EXCEL_PROVA)
check("file salvato e leggibile",
      os.path.exists(PERCORSO_EXCEL_PROVA)
      and load_workbook(PERCORSO_EXCEL_PROVA).active["A1"].value == "uno")
check("nessun file provvisorio lasciato in giro",
      not os.path.exists(PERCORSO_EXCEL_PROVA + ".nuovo"))
wbp.active["A1"] = "due"
r.salva_excel_sicuro(wbp, PERCORSO_EXCEL_PROVA)
check("secondo salvataggio aggiornato (sostituito tutto)",
      load_workbook(PERCORSO_EXCEL_PROVA).active["A1"].value == "due")

errore_arrivato = []
wbp2 = Workbook()
try:
    os.chmod(PERCORSO_EXCEL_PROVA, 0o444)
    try:
        r.salva_excel_sicuro(wbp2, PERCORSO_EXCEL_PROVA)
    except PermissionError:
        errore_arrivato.append("permessi")
finally:
    os.chmod(PERCORSO_EXCEL_PROVA, 0o644)
check("file protetto -> PermissionError come prima", "permessi" in errore_arrivato)
check("niente provvisorio dopo l'errore",
      not os.path.exists(PERCORSO_EXCEL_PROVA + ".nuovo"))

print("== F) STOP scritto sotto serratura (funzione vera) ==")
PERCORSO_STOP = os.path.join(cart, "clienti_stop_prova.xlsx")
r.crea_excel_se_manca(PERCORSO_STOP)
toccate = r.segna_stop_nel_excel(PERCORSO_STOP, "mario.rossi@email.it")
check("STOP scritto su 1 riga di esempio", len(toccate) == 1)
check("serratura rilasciata dopo lo STOP", not os.path.exists(serra))
wbcheck = load_workbook(PERCORSO_STOP)
wscheck = wbcheck[r.FOGLIO]
colonna_di, _ = r.individua_colonne(wscheck)
valore = wscheck.cell(row=toccate[0][0], column=colonna_di["attivo"]).value
check("nella cella Attivo c'e' scritto STOP", str(valore).strip().upper() == "STOP")

print("== G) pannello agganciato alla serratura condivisa ==")
check("pannello usa SerraturaExcel del motore",
      "SerraturaExcel" in open(PANNELLO, encoding="utf-8").read())
with p._serratura_file_clienti() as occupato:
    libero = occupato
check("pannello prende la serratura quando e' libera", libero == "")
check("niente excel.lock avanzato nella cartella di prova",
      not os.path.exists(serra))

print(f"\nRISULTATO v4.14: {ok} OK, {len(ko)} FALLITI {ko if ko else ''}")
sys.exit(1 if ko else 0)
