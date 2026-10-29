#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Test delle correzioni v4.13 del MOTORE e del PANNELLO: ogni controllo
prova il comportamento vero (non cerca stringhe nel codice).
Uso: REVMEMO_PKG=/percorso/RevMemo python3 test_correzioni_v413.py"""
import datetime as dt
import json
import os
import shutil
import smtplib
import socket
import ssl
import subprocess
import sys
import tempfile
import time
import urllib.request

PACCHETTO = os.environ.get("REVMEMO_PKG", "/home/z/my-project/download/RevMemo")
ok_n, falliti = 0, []


def check(nome, cond, dettaglio=""):
    global ok_n
    if cond:
        ok_n += 1
        print(f"  [OK]   {nome}")
    else:
        falliti.append(nome)
        print(f"  [FALL] {nome} {dettaglio}")


# ambiente isolato: copia del pacchetto con config di esempio
cart = tempfile.mkdtemp(prefix="v413_")
for n in ("revmemo.py", "config_ESEMPIO.ini", "template_email.html", "logo.png"):
    if os.path.exists(os.path.join(PACCHETTO, n)):
        shutil.copy(os.path.join(PACCHETTO, n), cart)
shutil.copy(os.path.join(cart, "config_ESEMPIO.ini"), os.path.join(cart, "config.ini"))
sys.path.insert(0, cart)
import io
from contextlib import redirect_stdout
with redirect_stdout(io.StringIO()):
    import revmemo as r
r.BASE = cart

print("== 1) STOP falso: citazioni di Outlook ==")
outlook = ("Vorrei prenotare martedi' mattina, grazie.\n\n"
           "Da: Centro <centro@gmail.com>\nInviato: giovedi 29 ottobre 2026 08:00\n"
           "A: mario@email.it\nOggetto: Promemoria revisione\n\n"
           "Gentile Mario, la revisione scade il 08/11/2026.\n---\n"
           "Non vuoi piu' ricevere promemoria? Rispondi a questa email con oggetto STOP\n")
check("Outlook: risposta normale NON e' uno STOP",
      not r.parola_stop_presente(r.rimuovi_citazioni(outlook)))
check("Outlook: STOP scritto dal cliente SOPRA la citazione resta valido",
      r.parola_stop_presente(r.rimuovi_citazioni("STOP\n\n" + outlook.split("\n\n", 1)[1])))
check("piede incollato senza intestazione: non e' uno STOP",
      not r.parola_stop_presente(r.rimuovi_citazioni(
          "ok grazie\nNon vuoi piu' ricevere promemoria? Rispondi a questa email con oggetto STOP")))
check("Gmail ('ha scritto:') funziona ancora",
      not r.parola_stop_presente(r.rimuovi_citazioni(
          "Va bene\nIl giorno gio ha scritto:\n> oggetto STOP")))
check("riga 'Da:' normale in un messaggio non lo taglia",
      "prenotare" in r.rimuovi_citazioni("Da: lunedi voglio prenotare\nciao"))

print("== 2) configurazione robusta ==")
cfg = open(os.path.join(cart, "config.ini"), encoding="utf-8").read()
for valore, atteso in (("1,5", 1.5), ("abc", 1.5), ("-3", 0.0)):
    d = tempfile.mkdtemp()
    shutil.copy(os.path.join(cart, "revmemo.py"), d)
    shutil.copy(os.path.join(cart, "template_email.html"), d)
    open(os.path.join(d, "config.ini"), "w", encoding="utf-8").write(
        cfg.replace("pausa_secondi = 1.5", f"pausa_secondi = {valore}"))
    out = subprocess.run([sys.executable, "-c",
                          "import revmemo;print(revmemo.CFG['pausa'])"],
                         cwd=d, capture_output=True, text=True)
    ultima = (out.stdout.strip().splitlines() or [""])[-1]
    check(f"pausa_secondi = {valore} non fa crollare (usa {atteso})",
          out.returncode == 0 and ultima == str(atteso), out.stderr[-120:])
d = tempfile.mkdtemp()
shutil.copy(os.path.join(cart, "revmemo.py"), d)
open(os.path.join(d, "config.ini"), "w", encoding="utf-8").write(
    cfg.replace("finestra_1mese = 26-30", "finestra_1mese = 30-26"))
out = subprocess.run([sys.executable, "-c",
                      "import revmemo;print(revmemo.CFG['finestre'][30])"],
                     cwd=d, capture_output=True, text=True)
check("finestra scritta al contrario (30-26) vale 26-30",
      "(26, 30)" in out.stdout, out.stdout[-80:] + out.stderr[-80:])

print("== 3) email e cella Attivo ==")
for e in ("mario@email.it", "a.b-c@x.co.uk"):
    check(f"email valida: {e}", r.email_valida(e))
for e in ("a@b", "a b@x.it", "a@x.it, b@y.it", "a@x.it;b@y.it", "a\n@x.it", "", None, "a@@x.it"):
    check(f"email NON valida: {e!r}", not r.email_valida(e))
import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Clienti"
for j, h in enumerate(["Nome Cliente", "Email", "Telefono", "Targa",
                       "Scadenza Revisione", "Revisione Fatta", "Attivo"]):
    ws.cell(row=4, column=2 + j, value=h)
for i, att in enumerate([None, "SÌ", "NO", "STOP"]):
    for j, v in enumerate([f"C{i}", f"c{i}@t.it", "", f"AA00{i}BB", "12/2030", "NO", att]):
        ws.cell(row=5 + i, column=2 + j, value=v)
x = os.path.join(cart, "c.xlsx")
wb.save(x)
cl = r.leggi_clienti(x)
check("Attivo VUOTO = cliente attivo", cl[0]["attivo"] is True)
check("Attivo SÌ = attivo", cl[1]["attivo"] is True)
check("Attivo NO = escluso", cl[2]["attivo"] is False)
check("Attivo STOP = escluso", cl[3]["attivo"] is False and cl[3]["attivo_testo"] == "STOP")

print("== 4) STOP: riattivato a mano non viene rimesso in STOP ==")
RAW = (b"From: Mario <c0@t.it>\r\nTo: centro@gmail.com\r\nSubject: Re: Promemoria\r\n"
       b"Message-ID: <uno@t.it>\r\nIn-Reply-To: <x@g>\r\n\r\nSTOP\r\n")
chiamate = []


class FakeIMAP:
    def __init__(self, host, port, timeout=None, ssl_context=None, **kw):
        chiamate.append(("ctx", ssl_context))

    def login(self, *a): return "OK", []
    def select(self, *a, **k): return "OK", [b"1"]
    def search(self, *a): return "OK", [b"1"]

    def fetch(self, num, spec):
        chiamate.append(("fetch", spec))
        if "HEADER" in spec:
            h = RAW.split(b"\r\n\r\n")[0] + b"\r\n\r\n"
            return "OK", [(b"1 (RFC822.HEADER {%d}" % len(h), h), b")"]
        return "OK", [(b"1 (BODY[] {%d}" % len(RAW), RAW), b")"]

    def store(self, *a): return "OK", []
    def logout(self): return "BYE", []


r.imaplib.IMAP4_SSL = FakeIMAP
r.MITTENTE_EMAIL, r.MITTENTE_PASSWORD = "centro@gmail.com", "x"
r.FILE_LOG = "log_prova.txt"


def attivo(riga=5):
    return openpyxl.load_workbook(x)["Clienti"].cell(row=riga, column=8).value


with redirect_stdout(io.StringIO()):
    r.controlla_risposte(x)
check("1a lettura: il cliente va in STOP", attivo() == "STOP")
wb = openpyxl.load_workbook(x)
wb["Clienti"].cell(row=5, column=8).value = "SÌ"
wb.save(x)
with redirect_stdout(io.StringIO()):
    r.controlla_risposte(x)
check("riattivato dal titolare: la stessa email NON lo rimette in STOP", attivo() == "SÌ", str(attivo()))
check("IMAP: contesto TLS con verifica certificato",
      any(c[0] == "ctx" and c[1] is not None and c[1].verify_mode == ssl.CERT_REQUIRED
          and c[1].check_hostname for c in chiamate))
check("IMAP: lettura con BODY.PEEK (non segna come letto)",
      ("fetch", "(BODY.PEEK[])") in chiamate and ("fetch", "(RFC822)") not in chiamate)

print("== 5) invio: errori e sicurezza ==")
r.MITTENTE_EMAIL, r.MITTENTE_PASSWORD = "centro@gmail.com", "x"


class SMTPFinto:
    modalita = "ok"
    tentativi = 0
    inviati = []
    msg = None
    ultimi_kw = {}

    def __init__(self, *a, **k):
        SMTPFinto.tentativi += 1
        SMTPFinto.ultimi_kw = k
        self.k = k

    def login(self, *a):
        if SMTPFinto.modalita == "auth":
            raise smtplib.SMTPAuthenticationError(535, b"bad")

    def send_message(self, m):
        SMTPFinto.msg = m
        if SMTPFinto.modalita == "5xx":
            raise smtplib.SMTPResponseException(550, b"no")
        SMTPFinto.inviati.append(m["To"])

    def quit(self): pass
    def close(self): pass


r.smtplib.SMTP_SSL = SMTPFinto
_SLEEP_VERO = time.sleep
r.time.sleep = lambda s: None
r.SMTP_PORT = 465
s = r.Sender()
SMTPFinto.modalita = "ok"
check("invio normale riesce", s.invia("c@t.it", "Og", "testo", None, None, con_stop=True) is None)
m = SMTPFinto.msg
check("header List-Unsubscribe con mailto STOP", "subject=STOP" in str(m["List-Unsubscribe"]))
check("header Date e Message-ID presenti", bool(m["Date"]) and bool(m["Message-ID"]))
ctx = (SMTPFinto.ultimi_kw or {}).get("context")
check("SMTP_SSL con contesto che VERIFICA il certificato",
      ctx is not None and ctx.verify_mode == ssl.CERT_REQUIRED
      and ctx.check_hostname,
      f"context={ctx}")
SMTPFinto.modalita, SMTPFinto.tentativi = "auth", 0
s2 = r.Sender()
e = s2.invia("c@t.it", "Og", "t")
check("password sbagliata: 1 solo tentativo (non 3)", SMTPFinto.tentativi == 1, str(SMTPFinto.tentativi))
check("password sbagliata: errore fatale, invii fermati", s2.errore_fatale is not None)
tent = SMTPFinto.tentativi
s2.invia("d@t.it", "Og", "t")
check("dopo errore fatale non riprova per gli altri clienti", SMTPFinto.tentativi == tent)
SMTPFinto.modalita, SMTPFinto.tentativi = "5xx", 0
s3 = r.Sender()
s3.invia("c@t.it", "Og", "t")
check("errore 5xx definitivo: nessun ritentativo", SMTPFinto.tentativi == 1, str(SMTPFinto.tentativi))
SMTPFinto.modalita = "ok"
s4 = r.Sender()
e = s4.invia("c@t.it", "Promemoria\nAB123", "t")
check("oggetto con a capo: errore gestito, nessun crash", isinstance(e, ValueError), repr(e))
check("dopo l'a capo il cliente successivo parte", r.Sender().invia("d@t.it", "Og", "t") is None)

print("== 6) backup: i giorni precedenti non spariscono ==")
cb = os.path.join(cart, "backup")
os.makedirs(cb, exist_ok=True)
for f in os.listdir(cb):
    os.remove(os.path.join(cb, f))
for g in range(14, 0, -1):
    p = os.path.join(cb, f"clienti_giorno{g:02d}_0800.xlsx")
    open(p, "wb").write(b"x")
    t = time.time() - g * 86400
    os.utime(p, (t, t))
for i in range(25):
    p = os.path.join(cb, f"clienti_oggi_{i:04d}.xlsx")
    open(p, "wb").write(b"x")
    os.utime(p, (time.time() - 25 + i, time.time() - 25 + i))
r.pulisci_backup(cb)
resti = os.listdir(cb)
vecchi = [f for f in resti if f.startswith("clienti_giorno")]
check("dopo 25 salvataggi in un giorno restano i backup dei 14 giorni prima", len(vecchi) == 14, str(len(vecchi)))
check("totale backup contenuto (<= 20 + 30)", len(resti) <= 50, str(len(resti)))

print("== 7) blocco anti-doppio-invio ==")
b1, b2 = r.BloccoInvio(), r.BloccoInvio()
check("primo invio prende il blocco", b1.prendi())
check("secondo invio contemporaneo e' rifiutato", not b2.prendi())
b1.rilascia()
check("dopo il rilascio il blocco si puo' riprendere", b2.prendi())
b2.rilascia()
b3 = r.BloccoInvio()
open(r.percorso_lock(), "w").write("999")
vecchio = time.time() - r.LOCK_SCADUTO_SECONDI - 60
os.utime(r.percorso_lock(), (vecchio, vecchio))
check("blocco abbandonato (nessun battito da 15 min) viene ripreso", b3.prendi())
b3.rilascia()

print("== 8) riepilogo al titolare: HTML sicuro ==")
spedite = {}


class SenderFinto:
    def invia(self, dest, og, testo, html=None, logo=None, **k):
        spedite["html"] = html
        return None


r.CFG["notifica_titolare"] = True
with redirect_stdout(io.StringIO()):
    r.invia_riepilogo(SenderFinto(), [{"targa": "AB1<b>", "email": "a&b@t.it", "desc": "x"}],
                      [{"targa": "T", "errore": "<script>x</script>"}], 0, 0, False)
h = spedite.get("html", "")
check("caratteri speciali resi sicuri nel riepilogo", "<script>" not in h and "AB1&lt;b&gt;" in h)

time.sleep = _SLEEP_VERO      # nella sezione 5 era finto: qui servono attese vere
print("== 9) pannello: sicurezza e messaggi (server vero) ==")
import re
pan = tempfile.mkdtemp(prefix="v413_pan_")
for n in ("revmemo.py", "pannello.py", "pannello.html", "controlla_pc.py",
          "template_email.html", "logo.png", "config_ESEMPIO.ini"):
    if os.path.exists(os.path.join(PACCHETTO, n)):
        shutil.copy(os.path.join(PACCHETTO, n), pan)
shutil.copy(os.path.join(pan, "config_ESEMPIO.ini"), os.path.join(pan, "config.ini"))
cfg_pan = open(os.path.join(pan, "config.ini"), encoding="utf-8").read()
open(os.path.join(pan, "config.ini"), "w", encoding="utf-8").write(
    cfg_pan.replace("password_app =", "password_app = SEGRETA123456789", 1))
sock = socket.socket()
sock.bind(("127.0.0.1", 0))
PORTA = sock.getsockname()[1]
sock.close()
proc = subprocess.Popen([sys.executable, os.path.join(pan, "pannello.py")],
                        env=dict(os.environ, PANNELLO_PORTA=str(PORTA), PYTHONUTF8="1",
                                 PYTHONDONTWRITEBYTECODE="1"),
                        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
URL = f"http://127.0.0.1:{PORTA}"


def chiama(percorso, corpo=None, intest=None):
    dati = json.dumps(corpo).encode() if corpo is not None else None
    rq = urllib.request.Request(URL + percorso, data=dati)
    for k, v in (intest or {}).items():
        rq.add_header(k, v)
    try:
        with urllib.request.urlopen(rq, timeout=20) as resp:
            return resp.status, resp.read().decode("utf-8")
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode("utf-8")


try:
    for _ in range(60):
        try:
            if chiama("/api/ping")[0] == 200:
                break
        except Exception:
            time.sleep(0.5)
    st, pagina = chiama("/")
    tok = (re.search(r'name="revmemo-token" content="([^"]+)"', pagina) or [None, ""])[1]
    check("la pagina contiene la chiave di sessione", len(tok) >= 20)
    buono = {"X-RevMemo-Token": tok}
    check("POST senza chiave (altro sito) -> 403",
          chiama("/api/impostazioni", {"smtp": {"host": "smtp.evil.example"}},
                 {"Origin": "http://evil.example", "Content-Type": "text/plain"})[0] == 403)
    check("il server SMTP NON e' stato cambiato",
          "evil.example" not in open(os.path.join(pan, "config.ini"), encoding="utf-8").read())
    check("chiave giusta ma Origin di un altro sito -> 403",
          chiama("/api/azione", {"tipo": "prova"},
                 dict(buono, Origin="http://evil.example", **{"Content-Type": "application/json"}))[0] == 403)
    check("Host falso (DNS rebinding) -> 403",
          chiama("/api/messaggio", None, dict(buono, Host="evil.example"))[0] == 403)
    st, corpo = chiama("/api/messaggio", None, buono)
    d = json.loads(corpo)
    check("la password NON arriva al browser", d["smtp"]["password_app"] == "" and "SEGRETA" not in corpo)
    check("ma il pannello sa che c'e' (password_presente)", d["password_presente"] is True)
    chiama("/api/impostazioni", {"smtp": {"host": "smtp.gmail.com", "porta": "465",
           "mittente": "a@gmail.com", "password_app": ""}},
           dict(buono, **{"Content-Type": "application/json"}))
    check("salvare con password vuota non cancella quella salvata",
          "SEGRETA123456789" in open(os.path.join(pan, "config.ini"), encoding="utf-8").read())
    check("/api/ping e' pubblico (serve a riconoscere un pannello aperto)", chiama("/api/ping")[0] == 200)
    # config.ini rotto: messaggio comprensibile (prima appariva solo "1")
    ok_cfg = open(os.path.join(pan, "config.ini"), encoding="utf-8").read()
    open(os.path.join(pan, "config.ini"), "w", encoding="utf-8").write(ok_cfg + "\nriga senza uguale\n")
    st, corpo = chiama("/api/stato", None, buono)
    err = json.loads(corpo).get("errore", "")
    check("config.ini rotto: messaggio chiaro, non '1'", "config.ini" in err and len(err) > 20, err)
    st, corpo = chiama("/api/anteprima", None, buono)
    check("config.ini rotto: anteprima spiega il problema", "config.ini" in corpo, corpo[:80])
    open(os.path.join(pan, "config.ini"), "w", encoding="utf-8").write(ok_cfg)
finally:
    proc.terminate()
    try:
        proc.wait(timeout=10)
    except Exception:
        proc.kill()

print()
print("=" * 50)
print(f"RISULTATO v4.13: {ok_n} OK, {len(falliti)} FALLITI", falliti if falliti else "")
sys.exit(1 if falliti else 0)
