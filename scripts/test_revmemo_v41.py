# -*- coding: utf-8 -*-
"""Test RevMemo v4.1 — riproduce ESATTAMENTE il problema dell'utente:
"se metto la mia email ad un cliente mi fa vedere STOP".

Scenario: il centro (centro.prova.revmemo@gmail.com) si manda un promemoria
a se stesso (la sua email e' scritta su una riga cliente). Il promemoria
contiene la parola STOP nel footer. PRIMA della v4.1 il programma lo
interpretava come una risposta STOP e bloccava il cliente.

Verifiche:
 1. rimuovi_citazioni: il testo citato viene scartato, la riga solo-STOP resta
 2. caselle_del_centro: mittente + email_contatti (titolare se compilata)
 3. controlla_risposte con GMAIL FINTA:
    - email del centro stesso con footer STOP -> IGNORATA (niente Excel)
    - cliente che risponde "VA BENE" citando il footer -> NON in STOP
    - cliente che risponde "STOP" -> in STOP
    - cliente che scrive solo STOP sotto la citazione -> in STOP
    - contatore "(ignorate N email inviate dal centro stesso)" a schermo
 4. main() completo in --dry-run con la Gmail finta attiva
"""
import io
import os
import shutil
import sys
import tempfile
from contextlib import redirect_stdout

from email.message import EmailMessage

PACCHETTO = os.environ.get("REVMEMO_PKG", "/home/z/my-project/download/RevMemo")

ok_count = 0
fail_count = 0


def check(nome, condizione, extra=""):
    global ok_count, fail_count
    if condizione:
        ok_count += 1
        print(f"  [OK]   {nome}")
    else:
        fail_count += 1
        print(f"  [FALLITO] {nome} {extra}")


tmp = tempfile.mkdtemp(prefix="revmemo_v41_")
for f in ("revmemo.py", "template_email.html", "logo.png"):
    shutil.copy2(os.path.join(PACCHETTO, f), os.path.join(tmp, f))

CONFIG_TEST = """\
[smtp]
host = smtp.gmail.com
porta = 465
mittente = centro.prova.revmemo@gmail.com
password_app = aaaabbbbccccdddd
risposta_a =

[centro]
nome = Centro Revisioni Test
citta = Milano
telefono = 02 1234567
email_contatti = info@test.it
colore_tema = #1F3B66

[prodotto]
nome = RevMemo

[promemoria]
finestra_1mese = 26-30
finestra_10giorni = 8-10
finestra_1giorno = 1-2

[invii]
pausa_secondi = 0
notifica_titolare = SI
email_titolare =
file_excel = clienti_revisioni.xlsx
file_log = reminders_log.txt

[sms]
attivo = NO
"""
with open(os.path.join(tmp, "config.ini"), "w", encoding="utf-8") as f:
    f.write(CONFIG_TEST)

sys.path.insert(0, tmp)
import revmemo  # noqa: E402

FOOTER = ("Non vuoi piu' ricevere promemoria? Rispondi a questa email "
          "con oggetto STOP.")

print("== 1) rimuovi_citazioni ==")
check("senza citazione: invariato",
      revmemo.rimuovi_citazioni("VA BENE\ngrazie") == "VA BENE\ngrazie")
check("testo vuoto", revmemo.rimuovi_citazioni("") == "")
quotata_gm = ("VA BENE, GRAZIE\n\n"
              "Il giorno mar 28 ott 2026 alle ore 09:00, Centro Test "
              "<centro@test.it> ha scritto:\n"
              "> Gentile Mario,\n> le ricordiamo...\n" + "> " + FOOTER + "\n")
check("citazione Gmail it: tagliata",
      revmemo.rimuovi_citazioni(quotata_gm) == "VA BENE, GRAZIE\n")
check("footer citato NON conta come STOP",
      not revmemo.parola_stop_presente(revmemo.rimuovi_citazioni(quotata_gm)))
quotata_gt = "STOP\n\nOn 28/10/2026 Centro wrote:\n> " + FOOTER + "\n"
check("citazione '>' + On...wrote: tagliata, STOP in cima resta",
      revmemo.rimuovi_citazioni(quotata_gt) == "STOP\n")
bottom_post = ("Il giorno 28 ott 2026 Centro <c@t.it> ha scritto:\n"
               "> " + FOOTER + "\nSTOP\n")
check("solo STOP sotto la citazione: recuperato",
      revmemo.parola_stop_presente(revmemo.rimuovi_citazioni(bottom_post)))
check("'STOP!' sotto la citazione: recuperato",
      revmemo.parola_stop_presente(revmemo.rimuovi_citazioni(
          "ha scritto:\n> ciao\nSTOP!\n")))
check("riga footer citata non e' un STOP nudo",
      not revmemo.parola_stop_presente(revmemo.rimuovi_citazioni(
          "ha scritto:\n> " + FOOTER + "\n")))
separatore = ("VA BENE\n-------- Messaggio originale --------\nDa: c@t.it\n"
              + FOOTER + "\n")
check("separatore Outlook/Gmail: tagliato",
      revmemo.rimuovi_citazioni(separatore) == "VA BENE")

print("== 2) caselle_del_centro ==")
proprie = revmemo.caselle_del_centro()
check("mittente incluso", "centro.prova.revmemo@gmail.com" in proprie)
check("email_contatti incluso", "info@test.it" in proprie)
check("email_titolare vuota: niente", len(proprie) == 2, str(proprie))

print("== 3) controlla_risposte con Gmail finta ==")
percorso_excel = os.path.join(tmp, "clienti_revisioni.xlsx")
revmemo.crea_excel_se_manca(percorso_excel)

# Simulo l'errore dell'utente: la MIA email scritta su una riga cliente
# (riga 8, Carlo Neri) -> l'email del centro risulta anche "di un cliente"
from openpyxl import load_workbook  # noqa: E402
_wb = load_workbook(percorso_excel)
_wb["Clienti"].cell(row=8, column=3).value = "centro.prova.revmemo@gmail.com"
_wb.save(percorso_excel)
_wb.close()

MESSAGGI = {}


def msg(da, oggetto, corpo):
    m = EmailMessage()
    m["From"] = da
    m["To"] = "centro.prova.revmemo@gmail.com"
    m["Subject"] = oggetto
    m.set_content(corpo)
    return m.as_bytes()

MESSAGGI = {
    1: msg("Centro Revisioni Test <centro.prova.revmemo@gmail.com>",
           "Promemoria revisione AB123CD - scadenza 09/2026",
           "Gentile Cliente,\n\nla revisione scade il 30/09/2026.\n\n" + FOOTER + "\n"),
    # 2: cliente gentile che risponde citando (Gmail): NON deve andare in STOP
    2: msg("Mario Rossi <mario.rossi@email.it>",
           "Re: Promemoria revisione AB123CD - scadenza 09/2026",
           "VA BENE, GRAZIE\n\nIl giorno mar 28 ott 2026 alle ore 09:00, "
           "Centro Revisioni Test <centro.prova.revmemo@gmail.com> ha scritto:\n"
           "> Gentile Mario Rossi,\n> la revisione scade il 30/09/2026.\n"
           "> " + FOOTER + "\n"),
    # 3: cliente che chiede STOP in cima -> DEVE andare in STOP
    3: msg("Luigi Verdi <luigi.verdi@email.it>",
           "Re: Promemoria revisione EF456GH - scadenza 09/2026",
           "STOP\n\nIl giorno mar 28 ott 2026, Centro Revisioni Test "
           "<centro.prova.revmemo@gmail.com> ha scritto:\n> " + FOOTER + "\n"),
    # 4: cliente che scrive SOLO STOP sotto la citazione -> DEVE andare in STOP
    4: msg("Sofia Gialli <sofia.gialli@email.it>",
           "Re: Promemoria revisione ST654UV - scadenza 09/2026",
           "Il giorno mar 28 ott 2026, Centro Revisioni Test "
           "<centro.prova.revmemo@gmail.com> ha scritto:\n> " + FOOTER + "\nSTOP\n"),
}


class FakeIMAP:
    istanza = None

    def __init__(self, host, port, timeout=None, **kw):
        self.marked = []
        self.fetch_full = []
        FakeIMAP.istanza = self

    def login(self, u, p):
        return ("OK", [b""])

    def select(self, box):
        return ("OK", [b"4"])

    def search(self, charset, *crit):
        return ("OK", [b"1 2 3 4"])

    def fetch(self, num, spec):
        raw = MESSAGGI[int(num)].replace(b"\n", b"\r\n")
        if "HEADER" in spec:
            cut = raw.find(b"\r\n\r\n")
            hdr = raw[:cut]
            return ("OK", [(num + b" (RFC822.HEADER {%d}" % len(hdr), hdr), b")"])
        self.fetch_full.append(num)
        return ("OK", [(num + b" (RFC822 {%d}" % len(raw), raw), b")"])

    def store(self, num, flags, flag):
        self.marked.append(num)
        return ("OK", [b""])

    def logout(self):
        return ("BYE", [b""])


imaplib_originale = revmemo.imaplib.IMAP4_SSL
revmemo.imaplib.IMAP4_SSL = FakeIMAP

buf = io.StringIO()
with redirect_stdout(buf):
    eventi = revmemo.controlla_risposte(percorso_excel, dry_run=False)
uscita = buf.getvalue()

check("2 eventi (Luigi + Sofia)", len(eventi) == 2, str(eventi))
check("luigi in STOP", eventi and eventi[0]["email"] == "luigi.verdi@email.it",
      str(eventi))
check("sofia in STOP", eventi and eventi[1]["email"] == "sofia.gialli@email.it",
      str(eventi))
check("email del centro IGNORATA (mai letta intera)",
      b"1" not in FakeIMAP.istanza.fetch_full, str(FakeIMAP.istanza.fetch_full))
check("clienti veri letti fino in fondo",
      FakeIMAP.istanza.fetch_full == [b"2", b"3", b"4"],
      str(FakeIMAP.istanza.fetch_full))
check("contatore email centro a schermo",
      "(ignorate 1 email inviate dal centro stesso" in uscita, uscita)
check("mario NON in STOP", "[STOP] mario.rossi@email.it" not in uscita)

from openpyxl import load_workbook as _lw  # noqa: E402
ws = _lw(percorso_excel)["Clienti"]
check("riga 5 (mario) resta SÌ", ws.cell(row=5, column=8).value == "SÌ")
check("riga 6 (luigi) = STOP", ws.cell(row=6, column=8).value == "STOP")
check("riga 9 (sofia) = STOP", ws.cell(row=9, column=8).value == "STOP")
check("marcate lette solo le 2 STOP vere",
      FakeIMAP.istanza.marked == [b"3", b"4"], str(FakeIMAP.istanza.marked))

print("== 4) rilancio: zero doppioni ==")
buf = io.StringIO()
with redirect_stdout(buf):
    eventi2 = revmemo.controlla_risposte(percorso_excel, dry_run=False)
uscita2 = buf.getvalue()
check("rilancio: zero nuovi eventi", eventi2 == [], str(eventi2))
check("rilancio: conta ancora l'email del centro",
      "(ignorate 1 email inviate dal centro stesso" in uscita2)

print("== 5) main() completo in PROVA con Gmail finta ==")
sys.argv = ["revmemo.py", "--dry-run"]
buf = io.StringIO()
with redirect_stdout(buf):
    codice = revmemo.main()
uscita3 = buf.getvalue()
check("main ritorna 0", codice == 0)
check("dice MODALITA' PROVA", "MODALITA' PROVA" in uscita3)
check("luigi escluso come STOP", "ha chiesto STOP" in uscita3)
check("riepilogo presente", "RIEPILOGO:" in uscita3)

revmemo.imaplib.IMAP4_SSL = imaplib_originale

print("== 6) il pacchetto vero resta pulito ==")
for f in ("config.ini", "clienti_revisioni.xlsx", "reminders_log.txt", "__pycache__"):
    check(f"pacchetto senza {f}", not os.path.exists(os.path.join(PACCHETTO, f)))

print()
print("=" * 50)
print(f"RISULTATO: {ok_count} OK, {fail_count} FALLITI")
print("=" * 50)
sys.exit(1 if fail_count else 0)
