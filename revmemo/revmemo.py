#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
=====================================================================
 RevMemo - Promemoria revisioni automatici      versione 4.14 (PRO)
=====================================================================
Cosa fa:
 - Legge i clienti dal file Excel (clienti_revisioni.xlsx; se manca,
   lo crea da solo con 8 clienti di esempio)
 - PRIMA degli invii guarda la posta in arrivo della casella mittente:
   se un CLIENTE ha risposto STOP, scrive Attivo = STOP nell'Excel e
   quel cliente non riceve piu' email finche' non lo riattivi tu
   (le COPIE dei promemoria mandati a se stessi non contano; una
   risposta VERA del titolare con STOP scritto a mano viene accettata)
 - Calcola la scadenza per MESE come sul libretto (es. 09/2026 =
   ultimo giorno di settembre 2026)
 - Invia email professionali in HTML (logo, colori, contatti del
   centro) a: 30 giorni, 10 giorni e 1 giorno dalla scadenza
 - IL MESSAGGIO E' PERSONALIZZABILE AL 100% (v4.4): in config.ini,
   sezione [messaggio], il centro scrive le PAROLE delle email
   (oggetto, saluto, testo, firma, riga extra del piede e istruzioni
   di disiscrizione) usando i segnaposto {NOME} {TARGA} {SCADENZA}
   {GIORNI} {TRAFRA} e tutti i dati del centro; si possono scrivere
   anche oggetti speciali per i 3 promemoria (30 giorni / 10 giorni /
   1 giorno). Il DESIGN resta in template_email.html (se e' della
   versione precedente viene aggiornato da solo, con backup).
 - Nuovi dati del centro in config.ini: indirizzo, sito_web e orari
   (se compilati appaiono nelle email, altrimenti restano fuori)
 - Il file Excel si colora da solo: SÌ verde, NO rosso, STOP rosso
   scuro (con --colora-excel aggiorni i colori di un file gia' esistente)
 - Manda al titolare un riepilogo di cio' che e' partito (STOP inclusi)
 - Registra ogni invio in reminders_log.txt (zero doppioni: ogni
   promemoria parte UNA sola volta per ogni scadenza del cliente)
 - Backup automatico del file clienti prima di ogni esecuzione

Regole dei clienti:
 - Attivo = NO o STOP     -> nessuna email (STOP = l'ha chiesto il cliente)
 - Revisione Fatta = SI'  -> nessuna email (finche' non aggiorni la scadenza)
 - Finestre di invio (config.ini): 26-30 gg, 8-10 gg, 1-2 gg

Tutta la configurazione sta in config.ini: il codice NON si tocca.
Se un config.ini di una versione precedente non ha la sezione
[messaggio], il programma la aggiunge da solo (senza toccare i valori).

Uso:
    python revmemo.py             invio reale
    python revmemo.py --dry-run   PROVA: mostra senza inviare
    python revmemo.py --file x.xlsx   usa un file clienti alternativo
    python revmemo.py --colora-excel  mette/aggiorna i colori nell'Excel
=====================================================================
"""

import argparse
import calendar
import configparser
import datetime as dt
import email
import email.policy
import html as html_lib
import imaplib
import os
import re
import shutil
import smtplib
import ssl
import sys
import time
import zipfile
from email.message import EmailMessage
from email.utils import formataddr, formatdate, make_msgid, parseaddr

# Rende sicura la stampa a schermo su qualsiasi console Windows
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))

# LA versione del programma in UN SOLO POSTO: chi vuole mostrarla (pannello,
# riepilogo, test) la legge da qui, cosi' non puo' mai essere diversa.
VERSIONE = "4.14"

# Se c'e' il runtime incluso e NON lo stiamo usando, ripartiamo con quello:
# cosi' il Python installato sul PC non conta piu' (zero problemi di versione).
RUNTIME_PY = os.path.join(BASE, "runtime", "python.exe")
if (os.name == "nt"
        and os.path.isfile(RUNTIME_PY)
        and os.path.abspath(sys.executable).lower() != os.path.abspath(RUNTIME_PY).lower()
        and os.environ.get("REVMEMO_NEL_RUNTIME") != "1"):
    os.environ["REVMEMO_NEL_RUNTIME"] = "1"
    try:
        os.execv(RUNTIME_PY, [RUNTIME_PY, os.path.abspath(__file__)] + sys.argv[1:])
    except Exception:
        pass  # se il rilancio fallisce, prosegue con il Python corrente

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError as e:
    print("Problema con la libreria openpyxl:", e)
    print("Soluzione: fai partire il programma con la cartella runtime inclusa")
    print("nel pacchetto (doppio click su avvia_prova.bat o avvia_invio.bat).")
    print("Se il runtime non c'e', riscarica il pacchetto completo o esegui:")
    print("   pip install openpyxl")
    sys.exit(1)

# =====================================================================
# 1) CONFIGURAZIONE (da config.ini)
# =====================================================================

PERCORSO_CONFIG = os.path.join(BASE, "config.ini")

CONFIG_DEFAULT = """\
; =====================================================================
;  RevMemo - CONFIGURAZIONE
;  QUESTO e' il file da modificare (con Blocco note). Il codice non
;  si tocca mai. Qui scrivi: i dati del centro, le caselle email e
;  LE PAROLE delle email (sezione [messaggio]).
;  Dopo una modifica: salva e chiudi, poi rilancia il programma.
; =====================================================================

[smtp]
; Casella email DALLA CUI partono i promemoria: va bene GMAIL oppure
; un'altra casella (Outlook, Libero, Virgilio, Aruba...).
; Con Gmail non toccare host e porta. Con UN'ALTRA casella cambia le
; due righe qui sotto:
;   Outlook  -> host = smtp.office365.com   porta = 587
;   Libero   -> host = smtp.libero.it       porta = 465
;   Virgilio -> host = smtp.virgilio.it     porta = 465
;   Aruba    -> host = smtp.aruba.it        porta = 465
host = smtp.gmail.com
porta = 465
mittente = latuaemail@gmail.com
; Password per app (Gmail: la Password per app di Google, 16 caratteri,
; vedi LEGGIMI.txt sezione 5; altre caselle: la password dell'email
; oppure la password per app se il fornitore la chiede)
password_app =
; Risposte dei clienti: arrivano SEMPRE nella casella mittente (quella
; sopra), dove il programma legge anche le risposte STOP (automatico).
; (La vecchia opzione "risposta_a" e' stata tolta dal pannello: se
; compilata lo STOP automatico smetteva di funzionare. Lasciala vuota.)
risposta_a =

[centro]
; I DATI DEL CENTRO: appaiono nelle email che ricevono i CLIENTI
nome = Centro Revisioni Esempio
citta = Milano
; Indirizzo completo (facoltativo): es. Via Roma 12, 20121 Milano (MI)
indirizzo =
telefono = 02 1234567
email_contatti = info@centrorevisioni.it
; Sito web (facoltativo): es. www.centrorevisioni.it
sito_web =
; Orari di apertura (facoltativo): es. Lun-Ven 8:00-13:00 e 14:30-18:00, Sab 8:00-12:00
orari =
; Colore principale delle email (codice esadecimale, es. blu #1F3B66, rosso #C8102E)
colore_tema = #1F3B66

[messaggio]
; =====================================================================
;  LE PAROLE DELLE EMAIL: qui il centro scrive il MESSAGGIO come vuole.
;  (Il disegno delle email - colori, logo, disposizione - sta invece in
;   template_email.html: quello lo tocca solo chi vuole cambiare la grafica.)
;
;  SEGNAPOSTI: il programma li sostituisce da solo con i dati veri.
;   {NOME}          nome del cliente
;   {TARGA}         targa del veicolo
;   {SCADENZA}      data di scadenza (gg/mm/aaaa)
;   {GIORNI}        giorni rimanenti
;   {TRAFRA}        "tra 29 giorni" / "tra 1 giorno"
;   {NOME_CENTRO}   nome del centro (sezione [centro])
;   {CITTA}         citta' del centro
;   {INDIRIZZO}     indirizzo del centro
;   {TELEFONO}      telefono del centro
;   {EMAIL_CENTRO}  email contatti del centro
;   {SITO_WEB}      sito web del centro
;   {ORARI}         orari del centro
;   {PRODOTTO}      nome del prodotto (white label)
;
;  Nel testo: una riga VUOTA separa i paragrafi; **parola** = grassetto.
;  ATTENZIONE: le righe sotto testo = devono iniziare con QUALCHE SPAZIO
;  (come gia' sono adesso): non togliere lo spazio iniziale.
; =====================================================================

; Oggetto dell'email (puoi usare i segnaposto, es. {TARGA})
oggetto = [{NOME_CENTRO}] Promemoria revisione {TARGA} - scade {TRAFRA}
; Oggetti speciali per i 3 promemoria (LASCIA VUOTO = usa l'oggetto generale)
oggetto_1mese =
oggetto_10giorni =
oggetto_1giorno =
; Saluto iniziale (se il cliente non ha il nome esce "Cliente")
saluto = Gentile {NOME},
; Il testo del messaggio (piu' paragrafi separati da una riga vuota)
testo =
    le ricordiamo che la revisione del veicolo targato **{TARGA}**
    scade **entro il {SCADENZA}** (mancano {GIORNI} giorni).

    Per prenotare la revisione rispondi a questa email
    oppure chiama il numero **{TELEFONO}**.
; Chiusura / firma
firma = Grazie della fiducia, {NOME_CENTRO}
; Riga extra nel piede dell'email (facoltativa): es. Parcheggio gratuito
; per i clienti, Appuntamenti anche il sabato, ecc.
riga_extra_piede =
; Istruzioni di disiscrizione nel piede. IMPORTANTE: lascia la parola
; STOP dentro il testo: e' la parola che il programma cerca nelle
; risposte dei clienti per metterli in pausa.
testo_disiscrizione = Non vuoi piu' ricevere promemoria? Rispondi a questa email con oggetto STOP

[prodotto]
; Nome del prodotto mostrato in piccolo nel piede delle email.
; Cambiabile liberamente (white label).
nome = RevMemo

[promemoria]
; Finestre di invio: (da giorni-al giorni) di anticipo sulla scadenza
finestra_1mese = 26-30
finestra_10giorni = 8-10
finestra_1giorno = 1-2

[invii]
; Pausa tra un'email e l'altra (secondi) - 1.5 consigliato
pausa_secondi = 1.5
; Riepilogo giornaliero al titolare: SI / NO
notifica_titolare = SI
; Email del titolare (se vuota, usa la stessa del mittente)
email_titolare =
; Nomi dei file (non serve toccarli)
file_excel = clienti_revisioni.xlsx
file_log = reminders_log.txt

[sms]
; Modulo SMS predisposto: quando il centro decide di attivare gli SMS
; (gateway tipo Aruba/Skebby) questa sezione verra' collegata.
; Per ora lasciare NO.
attivo = NO
"""


def _si(v, predefinito=False):
    if v is None:
        return predefinito
    return str(v).strip().upper() in ("SI", "SÌ", "S", "TRUE", "VERO", "1", "YES", "Y")


# Un SOLO indirizzo email spedibile: niente spazi, virgole, punti e virgola,
# parentesi o secondi indirizzi ("a@x.it, b@y.it" manderebbe a due persone;
# un "a capo" dentro la cella farebbe crollare l'intestazione dell'email)
_EMAIL_OK = re.compile(r"^[^@\s,;<>()\[\]\"]+@[^@\s,;<>()\[\]\"]+\.[^@\s,;<>()\[\]\"]{2,}$")


def email_valida(indirizzo):
    return bool(_EMAIL_OK.match(str(indirizzo or "").strip()))


# Testi PREDEFINITI delle email (usati quando config.ini non ha la
# sezione [messaggio] o ha qualche riga vuota). Stesse parole di prima,
# con in piu' il nome del cliente nel saluto.
OGGETTO_PREDEFINITO = ("[{NOME_CENTRO}] Promemoria revisione {TARGA}"
                       " - scade {TRAFRA}")
SALUTO_PREDEFINITO = "Gentile {NOME},"
TESTO_PREDEFINITO = ("le ricordiamo che la revisione del veicolo targato"
                     " **{TARGA}**\n"
                     "scade **entro il {SCADENZA}** (mancano {GIORNI}"
                     " giorni).\n\n"
                     "Per prenotare la revisione rispondi a questa email\n"
                     "oppure chiama il numero **{TELEFONO}**.")
FIRMA_PREDEFINITA = "Grazie della fiducia, {NOME_CENTRO}"
DISISCRIZIONE_PREDEFINITA = ("Non vuoi piu' ricevere promemoria? Rispondi"
                             " a questa email con oggetto STOP")


def _multiriga(valore):
    """Pulisce un valore su piu' righe di config.ini: toglie gli spazi
    all'inizio/fine di ogni riga e le righe vuote di troppo. Una riga
    vuota interna = separazione tra paragrafi e resta."""
    righe = [r.strip() for r in str(valore).replace("\r\n", "\n").split("\n")]
    while righe and not righe[0]:
        righe.pop(0)
    while righe and not righe[-1]:
        righe.pop()
    return "\n".join(righe)


# Blocco [messaggio] che viene AGGIUNTO ai config.ini delle versioni
# precedenti (che non lo avevano): cosi' anche un centro installato da
# tempo vede le nuove opzioni nel suo config.ini, senza perdere i valori.
BLOCCO_MESSAGGIO = """\
[messaggio]
; =====================================================================
;  LE PAROLE DELLE EMAIL: qui il centro scrive il MESSAGGIO come vuole.
;  (Il disegno delle email - colori, logo, disposizione - sta invece in
;   template_email.html: quello lo tocca solo chi vuole cambiare la grafica.)
;
;  SEGNAPOSTI: il programma li sostituisce da solo con i dati veri.
;   {NOME}          nome del cliente
;   {TARGA}         targa del veicolo
;   {SCADENZA}      data di scadenza (gg/mm/aaaa)
;   {GIORNI}        giorni rimanenti
;   {TRAFRA}        "tra 29 giorni" / "tra 1 giorno"
;   {NOME_CENTRO}   nome del centro (sezione [centro])
;   {CITTA}         citta' del centro
;   {INDIRIZZO}     indirizzo del centro
;   {TELEFONO}      telefono del centro
;   {EMAIL_CENTRO}  email contatti del centro
;   {SITO_WEB}      sito web del centro
;   {ORARI}         orari del centro
;   {PRODOTTO}      nome del prodotto (white label)
;
;  Nel testo: una riga VUOTA separa i paragrafi; **parola** = grassetto.
;  ATTENZIONE: le righe sotto testo = devono iniziare con QUALCHE SPAZIO
;  (come gia' sono adesso): non togliere lo spazio iniziale.
; =====================================================================

; Oggetto dell'email (puoi usare i segnaposto, es. {TARGA})
oggetto = [{NOME_CENTRO}] Promemoria revisione {TARGA} - scade {TRAFRA}
; Oggetti speciali per i 3 promemoria (LASCIA VUOTO = usa l'oggetto generale)
oggetto_1mese =
oggetto_10giorni =
oggetto_1giorno =
; Saluto iniziale (se il cliente non ha il nome esce "Cliente")
saluto = Gentile {NOME},
; Il testo del messaggio (piu' paragrafi separati da una riga vuota)
testo =
    le ricordiamo che la revisione del veicolo targato **{TARGA}**
    scade **entro il {SCADENZA}** (mancano {GIORNI} giorni).

    Per prenotare la revisione rispondi a questa email
    oppure chiama il numero **{TELEFONO}**.
; Chiusura / firma
firma = Grazie della fiducia, {NOME_CENTRO}
; Riga extra nel piede dell'email (facoltativa): es. Parcheggio gratuito
; per i clienti, Appuntamenti anche il sabato, ecc.
riga_extra_piede =
; Istruzioni di disiscrizione nel piede. IMPORTANTE: lascia la parola
; STOP dentro il testo: e' la parola che il programma cerca nelle
; risposte dei clienti per metterli in pausa.
testo_disiscrizione = Non vuoi piu' ricevere promemoria? Rispondi a questa email con oggetto STOP

"""

# Nuove righe del [centro] aggiunte ai config.ini vecchi (prima della
# sezione successiva, cosi' restano dentro [centro])
RIGHE_CENTRO_NUOVE = [
    "; Indirizzo completo (facoltativo): es. Via Roma 12, 20121 Milano (MI)",
    "indirizzo =",
    "; Sito web (facoltativo): es. www.centrorevisioni.it",
    "sito_web =",
    "; Orari di apertura (facoltativo): es. Lun-Ven 8:00-13:00 e 14:30-18:00",
    "orari =",
    "",
]


def aggiorna_config():
    """config.ini nato con una versione precedente: gli AGGIUNGE (senza
    toccare nessun valore gia' scritto) la sezione [messaggio] e i nuovi
    campi indirizzo/sito_web/orari del [centro]. Cosi' basta sostituire
    revmemo.py e il file di configurazione resta auto-documentante."""
    try:
        with open(PERCORSO_CONFIG, "r", encoding="utf-8", newline="") as f:
            testo = f.read()
    except Exception:
        return
    if "[messaggio]" in testo and "indirizzo" in testo:
        return
    nl = "\r\n" if "\r\n" in testo else "\n"
    fatte = []

    if "[messaggio]" not in testo:
        if not testo.endswith(nl):
            testo += nl
        testo += nl + BLOCCO_MESSAGGIO.replace("\n", nl)
        fatte.append("aggiunta la sezione [messaggio] (le parole delle email)")

    if "[centro]" in testo and "\nindirizzo" not in testo.replace("\r\n", "\n"):
        righe = testo.split(nl)
        inizio = next((i for i, r in enumerate(righe)
                       if r.strip().lower() == "[centro]"), None)
        if inizio is not None:
            fine = len(righe)
            for j in range(inizio + 1, len(righe)):
                nuda = righe[j].strip()
                if nuda.startswith("[") and nuda.endswith("]"):
                    fine = j
                    break
            righe[fine:fine] = [r.replace("\n", "") for r in RIGHE_CENTRO_NUOVE]
            testo = nl.join(righe)
            fatte.append("aggiunti i campi indirizzo, sito_web e orari al [centro]")

    if not fatte:
        return
    try:
        with open(PERCORSO_CONFIG, "w", encoding="utf-8", newline="") as f:
            f.write(testo)
        for m in fatte:
            print(f"config.ini aggiornato: {m} (i tuoi valori non sono toccati).")
    except Exception as e:
        print(f"AVVISO: non riesco ad aggiornare config.ini ({e}):")
        print("        le nuove opzioni restano comunque disponibili.")


def carica_config():
    if not os.path.exists(PERCORSO_CONFIG):
        try:
            with open(PERCORSO_CONFIG, "w", encoding="utf-8", newline="") as f:
                f.write(CONFIG_DEFAULT)
            print("config.ini non c'era: l'ho ricreato adesso (valori di partenza).")
            print("ATTENZIONE: se avevi gia' messo mittente o password, reinseriscili.\n")
        except Exception as e:
            print(f"ERRORE: manca config.ini e non riesco a crearlo: {e}")
            sys.exit(1)
    else:
        aggiorna_config()
    cfg = configparser.ConfigParser(interpolation=None)
    try:
        cfg.read(PERCORSO_CONFIG, encoding="utf-8")
    except configparser.ParsingError as e:
        print("ERRORE: config.ini non si riesce a leggere (riga con formato sbagliato).")
        print("Causa piu' comune: le righe dopo 'testo =' devono iniziare con")
        print("QUALCHE SPAZIO; controlla di non aver tolto lo spazio iniziale.")
        print(f"Dettagli: {e}")
        sys.exit(1)

    c = {}
    smtp = cfg["smtp"] if cfg.has_section("smtp") else {}
    c["smtp_host"] = smtp.get("host", "smtp.gmail.com").strip() or "smtp.gmail.com"
    try:
        c["smtp_porta"] = int(str(smtp.get("porta", "465")).strip() or "465")
    except (TypeError, ValueError):
        c["smtp_porta"] = 465
    c["mittente"] = smtp.get("mittente", "").strip()
    # Tolgo TUTTI gli spazi: Google mostra la chiave a gruppi di 4
    # ("abcd efgh ijkl mnop") e se viene incollata con gli spazi il login fallirebbe.
    c["password"] = "".join(smtp.get("password_app", "").split())
    c["risposta_a"] = smtp.get("risposta_a", "").strip()

    off = cfg["centro"] if cfg.has_section("centro") else {}
    c["nome_centro"] = off.get("nome", "Centro Revisioni")
    c["citta"] = off.get("citta", "").strip()
    c["indirizzo"] = off.get("indirizzo", "").strip()
    c["telefono"] = off.get("telefono", "").strip()
    c["email_contatti"] = off.get("email_contatti", "").strip()
    c["sito_web"] = off.get("sito_web", "").strip()
    c["orari"] = off.get("orari", "").strip()
    colore = off.get("colore_tema", "#1F3B66").strip()
    if not colore.startswith("#"):
        colore = "#" + colore
    c["colore"] = colore

    # LE PAROLE DELLE EMAIL (sezione [messaggio]; se manca, valgono i
    # testi predefiniti = quelli di sempre)
    msg = cfg["messaggio"] if cfg.has_section("messaggio") else {}
    oggetto = msg.get("oggetto", OGGETTO_PREDEFINITO).strip()
    c["msg_oggetto"] = oggetto or OGGETTO_PREDEFINITO
    c["msg_oggetti_tipo"] = {}
    for tipo, chiave in ((30, "oggetto_1mese"), (10, "oggetto_10giorni"),
                         (1, "oggetto_1giorno")):
        v = msg.get(chiave, "").strip()
        c["msg_oggetti_tipo"][tipo] = v or None
    c["msg_saluto"] = _multiriga(msg.get("saluto", SALUTO_PREDEFINITO)) \
        or SALUTO_PREDEFINITO
    c["msg_testo"] = _multiriga(msg.get("testo", TESTO_PREDEFINITO)) \
        or TESTO_PREDEFINITO
    c["msg_firma"] = _multiriga(msg.get("firma", FIRMA_PREDEFINITA)) \
        or FIRMA_PREDEFINITA
    c["msg_extra"] = msg.get("riga_extra_piede", "").strip()
    disiscrizione = msg.get("testo_disiscrizione", DISISCRIZIONE_PREDEFINITA).strip()
    c["msg_disiscrizione"] = disiscrizione or DISISCRIZIONE_PREDEFINITA

    prod = cfg["prodotto"] if cfg.has_section("prodotto") else {}
    c["prodotto"] = prod.get("nome", "RevMemo")

    prom = cfg["promemoria"] if cfg.has_section("promemoria") else {}

    def finestra(val, predef):
        try:
            a, b = str(val).split("-")
            a, b = int(a), int(b)
            return (min(a, b), max(a, b))     # "30-26" vale come "26-30"
        except Exception:
            try:
                n = int(val)
                return (n, n)
            except Exception:
                return predef

    c["finestre"] = {
        30: finestra(prom.get("finestra_1mese", "26-30"), (26, 30)),
        10: finestra(prom.get("finestra_10giorni", "8-10"), (8, 10)),
        1: finestra(prom.get("finestra_1giorno", "1-2"), (1, 2)),
    }
    c["etichette"] = {30: "30 giorni prima", 10: "10 giorni prima", 1: "1 giorno prima"}

    inv = cfg["invii"] if cfg.has_section("invii") else {}
    # "1,5" (virgola all'italiana) vale come 1.5; un testo strano non
    # deve far crollare il programma: si usa il valore consigliato
    try:
        c["pausa"] = max(0.0, float(str(inv.get("pausa_secondi", "1.5"))
                                    .strip().replace(",", ".")))
    except (TypeError, ValueError):
        c["pausa"] = 1.5
    c["notifica_titolare"] = _si(inv.get("notifica_titolare", "SI"), True)
    c["email_titolare"] = inv.get("email_titolare", "").strip()

    c["file_excel"] = inv.get("file_excel", "clienti_revisioni.xlsx")
    c["file_log"] = inv.get("file_log", "reminders_log.txt")
    return c


CFG = carica_config()
SMTP_HOST = CFG["smtp_host"]
SMTP_PORT = CFG["smtp_porta"]
MITTENTE_EMAIL = CFG["mittente"]
MITTENTE_PASSWORD = CFG["password"]
RISPOSTA_A = CFG["risposta_a"]
NOME_CENTRO = CFG["nome_centro"]
PRODOTTO = CFG["prodotto"]
FILE_EXCEL = CFG["file_excel"]
FILE_LOG = CFG["file_log"]

# =====================================================================
# 2) LETTURA DEL FILE EXCEL
# =====================================================================

COLONNE = {
    "nome":     "Nome Cliente",
    "email":    "Email",
    "telefono": "Telefono",          # opzionale (per SMS futuri)
    "targa":    "Targa",
    "scadenza": "Scadenza Revisione",
    "fatta":    "Revisione Fatta",
    "attivo":   "Attivo",
}
RIGA_INTESTAZIONI = 4
FOGLIO = "Clienti"


def ultimo_giorno_mese(anno, mese):
    return dt.date(anno, mese, calendar.monthrange(anno, mese)[1])


def converti_scadenza(val):
    """(data_scadenza, da_mese) — accetta 09/2026, 9/2026, date Excel
    auto-convertite (giorno 1 = mese del libretto) e date complete."""
    if isinstance(val, dt.datetime):
        val = val.date()
    if isinstance(val, dt.date):
        if val.day == 1:
            return ultimo_giorno_mese(val.year, val.month), True
        return val, False
    if isinstance(val, str):
        v = val.strip()
        for f in ("%m/%Y", "%m-%Y"):
            try:
                d = dt.datetime.strptime(v, f)
                return ultimo_giorno_mese(d.year, d.month), True
            except ValueError:
                continue
        for f in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
            try:
                return dt.datetime.strptime(v, f).date(), False
            except ValueError:
                continue
    return None, False


def individua_colonne(ws):
    """Trova le colonne dal nome delle intestazioni (riga 4).
    Ritorna (colonna_di, mancanti)."""
    indice = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=RIGA_INTESTAZIONI, column=col).value
        if val is not None:
            indice[str(val).strip().lower()] = col
    colonna_di, mancanti = {}, []
    for chiave, nome_col in COLONNE.items():
        c = indice.get(nome_col.strip().lower())
        if c is None:
            if chiave == "telefono":      # unica colonna facoltativa
                continue
            mancanti.append(nome_col)
        else:
            colonna_di[chiave] = c
    return colonna_di, mancanti


def leggi_clienti(percorso_excel):
    if not os.path.exists(percorso_excel):
        print(f"ERRORE: non trovo il file Excel: {percorso_excel}")
        sys.exit(1)
    try:
        wb = load_workbook(percorso_excel, data_only=True)
    except PermissionError:
        print("ERRORE: il file Excel non si puo' aprire perche' risulta APERTO")
        print("in Excel (o bloccato da OneDrive mentre sincronizza).")
        print("Soluzione: salva (Ctrl+S) e chiudi Excel, aspetta qualche secondo")
        print("e rilancia il programma. Se continua: riavvia il PC.")
        sys.exit(1)
    except zipfile.BadZipFile:
        print("ERRORE: il file Excel non e' leggibile (non e' un file Excel valido).")
        print("Aprilo con Excel: se si apre, salvalo di nuovo (Ctrl+S) e riprova.")
        sys.exit(1)
    ws = wb[FOGLIO] if FOGLIO in wb.sheetnames else wb.active

    colonna_di, mancanti = individua_colonne(ws)
    if mancanti:
        print("ERRORE: nel foglio Excel mancano le colonne: " + ", ".join(mancanti))
        print("Le intestazioni alla riga 4 devono essere: "
              + " | ".join(COLONNE.values()))
        sys.exit(1)

    clienti = []
    for row in range(RIGA_INTESTAZIONI + 1, ws.max_row + 1):
        def cella(chiave):
            if chiave not in colonna_di:
                return None
            return ws.cell(row=row, column=colonna_di[chiave]).value

        nome, email, targa = cella("nome"), cella("email"), cella("targa")
        scadenza, fatta, attivo = cella("scadenza"), cella("fatta"), cella("attivo")
        telefono = cella("telefono")

        # riga vuota o nota a fondo tabella -> salta
        if all(v is None or str(v).strip() == "" for v in (email, targa, scadenza)):
            continue

        data_scad, da_mese = converti_scadenza(scadenza)
        attivo_testo = str(attivo).strip().upper() if attivo is not None else ""
        clienti.append({
            "riga": row,
            "nome": str(nome).strip() if nome else "",
            "email": str(email).strip() if email else "",
            "telefono": str(telefono).strip() if telefono else "",
            "targa": str(targa).strip().upper() if targa else "",
            "scadenza": data_scad,
            "da_mese": da_mese,
            "fatta": _si(fatta),
            # Cella Attivo VUOTA = attivo: le guide dicono che solo NO e STOP
            # escludono un cliente. Prima un cliente con la cella lasciata
            # vuota veniva saltato in silenzio e non riceveva mai nulla.
            "attivo": True if attivo_testo == "" else _si(attivo),
            "attivo_testo": attivo_testo,
        })
    return clienti

# =====================================================================
# 3) REGISTRO INVII (anti-doppioni)
# =====================================================================

def percorso_log():
    return os.path.join(BASE, FILE_LOG)


_DATA_ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def gia_inviato(email, targa, tipo, scadenza):
    """True se quel promemoria e' GIA' partito per questa scadenza.
    La chiave e' email;targa;tipo;scadenza: ogni promemoria parte UNA
    sola volta per ciclo (se il cliente rinnova la revisione, la scadenza
    cambia e i promemoria ripartono da soli). Le righe del vecchio
    formato (senza scadenza) valgono comunque come 'gia' inviato':
    nessun rischio di doppioni dopo l'aggiornamento."""
    percorso = percorso_log()
    if not os.path.exists(percorso):
        return False
    email_l, targa_u = email.lower().strip(), targa.upper().strip()
    with open(percorso, "r", encoding="utf-8") as f:
        for line in f:
            parti = line.split(";")
            if len(parti) < 5:
                continue
            if (parti[1].strip().lower() == email_l
                    and parti[2].strip().upper() == targa_u
                    and parti[3].strip() == str(tipo)):
                quarto = parti[4].strip()
                # formato nuovo: campo 4 = scadenza AAAA-MM-GG
                # formato vecchio: campo 4 = descrizione -> comunque inviato
                if quarto == scadenza or not _DATA_ISO.match(quarto):
                    return True
    return False


def registra_invio(email, targa, tipo, giorni, descrizione, scadenza):
    riga = (f"{dt.date.today().isoformat()};{email.lower()};{targa.upper()};"
            f"{tipo};{scadenza};{descrizione} ({giorni} gg)")
    with open(percorso_log(), "a", encoding="utf-8") as f:
        f.write(riga + "\n")

# =====================================================================
# 3-bis) RISPOSTE DEI CLIENTI: STOP AUTOMATICO (lettura della casella)
# =====================================================================
# Come funziona: prima di inviare, il programma si collega alla casella
# mittente (Gmail/IMAP, porta 993) con la STESSA email e la STESSA
# password_app di config.ini, guarda le email degli ultimi 7 giorni e,
# se un cliente ha risposto con la parola STOP, scrive STOP nella
# colonna Attivo dell'Excel. Le email degli altri mittenti non vengono
# toccate (restano lette/non lette come sono).
#
# PROTEZIONI v4.1/v4.2 (imparate dai test veri):
# 1) le COPIE dei promemoria che il centro si manda a se stesso vengono
#    ignorate: contengono la parola "STOP" nel footer e non sono
#    risposte (nessun Re:/In-Reply-To, oggetto "Promemoria...");
# 2) MA una risposta VERA del titolare dalla sua casella (Re:/In-Reply-To)
#    o un'email con oggetto STOP vengono accettate: cosi' lo STOP si puo'
#    provare anche dalla propria casella (basta premere RISPONDI);
# 3) nelle risposte viene letta solo la parte scritta a mano: il testo
#    citato sotto ("Il giorno ... ha scritto:") viene scartato, cosi' un
#    "VA BENE" che cita il footer con STOP non e' uno STOP.

IMAP_PORTA = 993
_MESI_EN = ("Jan", "Feb", "Mar", "Apr", "May", "Jun",
            "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
GIORNI_DI_RICERCA = 7          # guarda le email degli ultimi N giorni
MASSIMO_EMAIL_DA_VEDERE = 150  # oltre: considera solo le piu' recenti (e lo dice)


def contesto_tls():
    """Contesto TLS che VERIFICA il certificato del server (nome compreso).
    Se non ricevono un contesto, smtplib e imaplib ne creano uno che NON
    controlla nulla: chiunque sulla rete potrebbe spacciarsi per il server
    di posta e leggere la password. Qui si usa sempre quello verificato
    (certificati del sistema: su Windows quelli di Windows)."""
    return ssl.create_default_context()


def causa_tls(e):
    """Spiegazione in italiano se l'errore e' un certificato non valido."""
    if isinstance(e, ssl.SSLCertVerificationError):
        return ("Il certificato del server di posta non risulta valido: "
                "controlla che data e ora del PC siano giuste e che "
                "l'antivirus non intercetti le connessioni sicure.")
    return ""


# Memoria degli STOP gia' registrati. Senza, una risposta STOP che resta
# nella posta (ultimi 7 giorni) rimetterebbe in pausa un cliente che il
# titolare ha appena RIATTIVATO a mano. Una riga per email: data;chiave.
def percorso_stop_gestiti():
    return os.path.join(BASE, "stop_gestiti.txt")


def chiave_messaggio(msg_hdr, mittente, oggetto):
    """Identifica una email: il suo Message-ID; se manca, mittente+data+oggetto."""
    mid = str(msg_hdr.get("Message-ID", "") or "").strip()
    if mid:
        return mid
    data = str(msg_hdr.get("Date", "") or "").strip()
    return f"{mittente}|{data}|{oggetto.strip()}"[:300]


def carica_stop_gestiti():
    gestiti = set()
    try:
        with open(percorso_stop_gestiti(), "r", encoding="utf-8") as f:
            for riga in f:
                parti = riga.rstrip("\r\n").split(";", 1)
                if len(parti) == 2 and parti[1]:
                    gestiti.add(parti[1])
    except OSError:
        pass
    return gestiti


def segna_stop_gestito(chiave):
    try:
        with open(percorso_stop_gestiti(), "a", encoding="utf-8") as f:
            f.write(f"{dt.date.today().isoformat()};"
                    f"{' '.join(str(chiave).split())}\n")
    except OSError:
        pass


def host_imap_da(host_smtp):
    """smtp.gmail.com -> imap.gmail.com (stesso fornitore).
    Outlook/Hotmail/Live non seguono lo schema smtp.->imap.: il loro
    server di posta in arrivo e' outlook.office365.com."""
    h = (host_smtp or "").strip().lower()
    if any(x in h for x in ("office365", "outlook", "hotmail", "live")):
        return "outlook.office365.com"
    if "smtp." in h:
        return h.replace("smtp.", "imap.", 1)
    return h or "imap.gmail.com"


def parola_stop_presente(testo):
    """True se c'e' la parola STOP da sola (stop, Stop, STOP, anche
    dentro una frase o con '!' accanto). 'Stoppa'/'stoppare' NON contano."""
    return re.search(r"\bstop\b", testo or "", re.IGNORECASE) is not None


def caselle_del_centro():
    """Gli indirizzi del CENTRO STESSO (mittente, titolare, contatti):
    le email che partono da qui non sono mai risposte dei clienti.
    Il promemoria contiene la parola STOP nel footer: senza questa
    protezione un invio di prova a se stessi diventerebbe uno STOP falso
    (e il cliente 'se stesso' non riceverebbe mai piu' email)."""
    caselle = {(MITTENTE_EMAIL or "").strip().lower()}
    for chiave in ("email_titolare", "email_contatti"):
        v = (CFG.get(chiave) or "").strip().lower()
        if "@" in v:
            caselle.add(v)
    return caselle


_MARCATORI_CITAZIONE = ("ha scritto:", "wrote:",
                        "messaggio originale", "original message",
                        "messaggio inoltrato", "forwarded message")
_INIZI_INTESTAZIONE = ("da:", "from:", "de:", "von:")
_INIZI_CAMPI = ("inviato:", "sent:", "data:", "date:", "oggetto:",
                "subject:", "a:", "to:", "cc:")


def _inizia_citazione_outlook(righe, i):
    """Outlook (e Outlook web) non usano '>' ne' 'ha scritto:': sotto la
    risposta mettono un blocco 'Da: ... Inviato: ... A: ... Oggetto: ...'
    seguito dal promemoria originale, che contiene la scritta STOP del piede."""
    if not righe[i].strip().lower().startswith(_INIZI_INTESTAZIONE):
        return False
    for k in range(i + 1, min(i + 6, len(righe))):
        if righe[k].strip().lower().startswith(_INIZI_CAMPI):
            return True
    return False


def _riga_piede_promemoria(riga):
    """Riga del piede dei NOSTRI promemoria (quella con 'oggetto STOP'):
    non e' mai una richiesta del cliente, ovunque sia finita nel messaggio."""
    s = riga.lower()
    frasi = ("con oggetto stop", str(CFG.get("msg_disiscrizione") or "").strip().lower())
    return any(f and f in s for f in frasi)


def rimuovi_citazioni(testo):
    """Ritorna la risposta SENZA il testo citato (il promemoria originale
    che il programma di posta riporta sotto quando si risponde): la parola
    STOP dentro il footer del promemoria NON deve contare se il cliente ha
    scritto altro. Resta sempre valida una riga che contiene SOLO la parola
    STOP (anche se il cliente la scrive sotto la citazione). Riconosce le
    citazioni di Gmail/Apple Mail ('ha scritto:', '>') e quelle di Outlook
    ('Da: ... Inviato: ...'); in piu' scarta ovunque la riga del piede."""
    righe = (testo or "").splitlines()
    risultato = testo or ""
    for i, riga in enumerate(righe):
        s = riga.strip().lower()
        if (s.startswith(">") or any(m in s for m in _MARCATORI_CITAZIONE)
                or _inizia_citazione_outlook(righe, i)):
            # riga con SOLO la parola STOP dopo la citazione = scritta dal cliente
            coda = [r for r in righe[i + 1:]
                    if r.strip().strip("!.").upper() == "STOP"]
            risultato = "\n".join(righe[:i] + coda)
            break
    rimaste = risultato.splitlines()
    if any(_riga_piede_promemoria(r) for r in rimaste):
        return "\n".join(r for r in rimaste if not _riga_piede_promemoria(r))
    return risultato


def _primo_blocco(dati):
    """Estrae i byte del messaggio dalla risposta di imaplib.fetch.
    La risposta e' tipo [(b'1 (RFC822 {123}', b'<contenuto>'), b')']:
    nel primo elemento della tupla c'e' il prefisso con la dimensione,
    il contenuto vero e' l'altro pezzo."""
    for elemento in dati or []:
        if isinstance(elemento, tuple):
            candidati = [p for p in elemento if isinstance(p, bytes)]
            # scarta il prefisso tipo b'1 (RFC822 {123}': il contenuto e' l'altro
            buoni = [p for p in candidati
                     if not re.search(rb"\{\d+\}\s*$", p)
                     and not re.match(rb"^\d+ \(", p)]
            if buoni:
                return buoni[0]
            if candidati:
                return candidati[-1]
        elif isinstance(elemento, bytes) and len(elemento) > 5:
            return elemento
    return None


def estrai_da_e_testo(msg_raw):
    """Da un'email grezza: (indirizzo_mittente, oggetto, corpo_in_testo)."""
    try:
        msg = email.message_from_bytes(msg_raw, policy=email.policy.default)
        mitt = parseaddr(str(msg.get("From", "")))[1].strip().lower()
        oggetto = str(msg.get("Subject", "") or "")
        corpo = ""
        if msg.is_multipart():
            for parte in msg.walk():
                if parte.get_content_type() == "text/plain":
                    corpo = parte.get_content()
                    break
            if not corpo.strip():
                for parte in msg.walk():
                    if parte.get_content_type() == "text/html":
                        corpo = parte.get_content()
                        break
        else:
            corpo = msg.get_content()
    except Exception:
        try:
            msg = email.message_from_bytes(msg_raw)
            mitt = parseaddr(str(msg.get("From", "")))[1].strip().lower()
            oggetto = str(msg.get("Subject", "") or "")
            corpo = (msg.get_payload(decode=True) or b"").decode("utf-8", "replace")
        except Exception:
            return "", "", ""
    # se il corpo e' HTML tolgo i tag (cosi' <b>STOP</b> vale come STOP)
    corpo = html_lib.unescape(re.sub(r"<[^>]+>", " ", corpo))
    return mitt, oggetto, corpo


def email_clienti_nel_excel(percorso_excel):
    """Tutte le email dei clienti nell'Excel (in minuscolo)."""
    try:
        wb = load_workbook(percorso_excel, data_only=True)
    except Exception:
        return set()
    try:
        ws = wb[FOGLIO] if FOGLIO in wb.sheetnames else wb.active
        colonna_di, _ = individua_colonne(ws)
        if "email" not in colonna_di:
            return set()
        col = colonna_di["email"]
        emails = set()
        for riga in range(RIGA_INTESTAZIONI + 1, ws.max_row + 1):
            v = ws.cell(row=riga, column=col).value
            if v is not None and "@" in str(v):
                emails.add(str(v).strip().lower())
        return emails
    finally:
        wb.close()


def segna_stop_nel_excel(percorso_excel, email_mittente, esito=None):
    """Scrive STOP nella colonna Attivo di tutte le righe di quel cliente.
    Ritorna [(riga, nome, targa), ...] delle righe aggiornate ([] = niente).
    Se si passa un dizionario 'esito', in caso di ERRORE (Excel aperto,
    file illeggibile, file occupato) ci scrive esito["errore"] = True:
    cosi' chi chiama distingue 'era gia' in pausa' da 'non sono riuscito
    a scrivere'. La scrittura avviene sotto la serratura condivisa col
    pannello: niente salvataggi nello stesso istante."""
    serratura = SerraturaExcel()
    if not serratura.prendi():
        if esito is None:
            esito = {}
        esito["errore"] = True
        print("  ATTENZIONE: il file clienti e' occupato da un'altra scrittura")
        print("  (pannello?): lo STOP verra' registrato al prossimo giro.")
        return []
    try:
        return _segna_stop_serrato(percorso_excel, email_mittente, esito)
    finally:
        serratura.rilascia()


def _segna_stop_serrato(percorso_excel, email_mittente, esito=None):
    """(la scrittura STOP vera, sotto serratura)"""
    if esito is None:
        esito = {}
    try:
        wb = load_workbook(percorso_excel)
    except PermissionError:
        esito["errore"] = True
        print("  ATTENZIONE: l'Excel risulta APERTO: non posso registrare lo STOP.")
        print("  Chiudi Excel e rilancia il programma: lo STOP verra' registrato allora.")
        return []
    except Exception as e:
        esito["errore"] = True
        print(f"  ATTENZIONE: non riesco ad aggiornare l'Excel ({e})")
        return []
    try:
        ws = wb[FOGLIO] if FOGLIO in wb.sheetnames else wb.active
        colonna_di, _ = individua_colonne(ws)
        if "email" not in colonna_di or "attivo" not in colonna_di:
            return []
        obiettivo = email_mittente.strip().lower()
        col_nome = colonna_di.get("nome", colonna_di["email"])
        col_targa = colonna_di.get("targa", colonna_di["email"])
        toccate = []
        for riga in range(RIGA_INTESTAZIONI + 1, ws.max_row + 1):
            v = ws.cell(row=riga, column=colonna_di["email"]).value
            if v is None or str(v).strip().lower() != obiettivo:
                continue
            cella_attivo = ws.cell(row=riga, column=colonna_di["attivo"])
            if str(cella_attivo.value).strip().upper() == "STOP":
                continue          # era gia' in pausa: niente da fare
            nome = ws.cell(row=riga, column=col_nome).value
            targa = ws.cell(row=riga, column=col_targa).value
            cella_attivo.value = "STOP"
            toccate.append((riga,
                            str(nome).strip() if nome else "",
                            str(targa).strip().upper() if targa else ""))
        if toccate:
            salva_excel_sicuro(wb, percorso_excel)
            for _, _, targa in toccate:
                registra_invio(email_mittente, targa, "STOP", 0,
                               "risposta STOP registrata", "-")
        return toccate
    except PermissionError:
        esito["errore"] = True
        print("  ATTENZIONE: l'Excel risulta APERTO: lo STOP NON e' stato salvato.")
        print("  Chiudi Excel e rilancia il programma: verra' registrato al prossimo avvio.")
        return []
    except Exception as e:
        esito["errore"] = True
        print(f"  ATTENZIONE: aggiornamento Excel non riuscito ({e})")
        return []


def controlla_risposte(percorso_excel, dry_run=False):
    """Guarda la posta in arrivo della casella mittente: se un cliente ha
    risposto STOP, scrive Attivo = STOP nell'Excel. Ritorna la lista degli
    eventi ([] = nessuno, None = controllo non eseguito)."""
    if not MITTENTE_EMAIL:
        return None
    if not MITTENTE_PASSWORD:
        if dry_run:
            print("Risposte STOP: controllo saltato (manca la password per app in config.ini)")
        return None
    print(f"Controllo le risposte degli ultimi {GIORNI_DI_RICERCA} giorni "
          f"nella casella {MITTENTE_EMAIL} ...")
    try:
        posta = imaplib.IMAP4_SSL(host_imap_da(SMTP_HOST), IMAP_PORTA,
                                  ssl_context=contesto_tls(), timeout=30)
        posta.login(MITTENTE_EMAIL, MITTENTE_PASSWORD)
    except Exception as e:
        print(f"AVVISO: non riesco a controllare le risposte STOP "
              f"({e.__class__.__name__}: {e}).")
        if causa_tls(e):
            print("        " + causa_tls(e))
        print("         Proseguo comunque con gli invii.")
        return None
    eventi = []
    try:
        posta.select("INBOX")
        limite = dt.date.today() - dt.timedelta(days=GIORNI_DI_RICERCA)
        dal = f"{limite.day:02d}-{_MESI_EN[limite.month - 1]}-{limite.year}"
        tip, dati = posta.search(None, "SINCE", dal)
        numeri = (dati[0] or b"").split() if tip == "OK" else []
        if len(numeri) > MASSIMO_EMAIL_DA_VEDERE:
            print(f"  AVVISO: nella posta ci sono {len(numeri)} email degli ultimi "
                  f"{GIORNI_DI_RICERCA} giorni: guardo solo le "
                  f"{MASSIMO_EMAIL_DA_VEDERE} piu' recenti.")
            numeri = numeri[-MASSIMO_EMAIL_DA_VEDERE:]
        if not numeri:
            print("  nessun nuovo STOP.")
            return eventi
        indirizzi_clienti = email_clienti_nel_excel(percorso_excel)
        caselle_proprie = caselle_del_centro()
        gia_gestiti = carica_stop_gestiti()
        ignorate_proprie = 0
        for num in numeri:
            tip, dati = posta.fetch(num, "(RFC822.HEADER)")
            intestazioni = _primo_blocco(dati)
            if not intestazioni:
                continue
            msg_hdr = email.message_from_bytes(intestazioni)
            mitt_rapido = parseaddr(str(msg_hdr.get("From", "")))[1].strip().lower()
            if not mitt_rapido:
                continue
            oggetto_rapido = str(msg_hdr.get("Subject", "") or "")
            chiave_msg = chiave_messaggio(msg_hdr, mitt_rapido, oggetto_rapido)
            if chiave_msg in gia_gestiti:
                continue          # STOP gia' registrato (e magari riattivato dopo)
            if mitt_rapido in caselle_proprie:
                # Email che PARTE dal centro stesso. Quasi sempre e' la copia
                # di un promemoria (o del riepilogo) mandato a se stessi:
                # quelle si IGNORANO, perche' contengono la parola STOP nel
                # footer e non sono risposte vere. Ma se il titolare RISPONDE
                # davvero (Re:/In-Reply-To) o scrive STOP nell'oggetto, e' un
                # gesto umano: si prosegue e il corpo decide se e' uno STOP.
                e_risposta = bool(msg_hdr.get("In-Reply-To")
                                  or msg_hdr.get("References")) or \
                    oggetto_rapido.strip().lower().startswith(("re:", "r:"))
                if not e_risposta and not parola_stop_presente(oggetto_rapido):
                    ignorate_proprie += 1
                    continue
            elif mitt_rapido not in indirizzi_clienti:
                continue          # non e' uno dei clienti: non ci interessa
            # BODY.PEEK[] legge il messaggio SENZA segnarlo come letto: con
            # (RFC822) ogni risposta dei clienti (anche "vorrei prenotare")
            # risultava gia' letta nella casella del titolare, che rischiava
            # di non accorgersene mai.
            tip, dati = posta.fetch(num, "(BODY.PEEK[])")
            msg_raw = _primo_blocco(dati)
            if not msg_raw:
                continue
            mitt, oggetto, corpo = estrai_da_e_testo(msg_raw)
            if not parola_stop_presente(oggetto + "\n" + rimuovi_citazioni(corpo)):
                continue
            if mitt_rapido in caselle_proprie and mitt_rapido not in indirizzi_clienti:
                print(f"  [INFO] STOP letto da {mitt_rapido} ma questa email "
                      f"non e' su nessuna riga cliente dell'Excel:")
                print("         niente da fermare (aggiungi la riga o usa "
                      "l'email di un cliente di prova).")
                continue
            if dry_run:
                print(f"  [PROVA] {mitt} ha chiesto STOP: nell'Excel "
                      f"verrebbe messo Attivo = STOP")
                eventi.append({"email": mitt, "righe": [], "prova": True})
                continue
            esito = {}
            toccate = segna_stop_nel_excel(percorso_excel, mitt, esito)
            if not toccate:
                if not esito.get("errore"):
                    # era gia' in pausa: ricordo comunque questa email, cosi'
                    # una riattivazione futura non viene annullata da lei
                    segna_stop_gestito(chiave_msg)
                    gia_gestiti.add(chiave_msg)
                continue          # Excel non salvato: riprovo al prossimo avvio
            segna_stop_gestito(chiave_msg)
            gia_gestiti.add(chiave_msg)
            try:
                posta.store(num, "+FLAGS", "\\Seen")
            except Exception:
                pass
            nomi = ", ".join(n or "(senza nome)" for _, n, _ in toccate)
            targhe = ", ".join(t or "(targa?)" for _, _, t in toccate)
            print(f"  [STOP] {mitt} -> Attivo = STOP nell'Excel "
                  f"({nomi} - targhe {targhe})")
            eventi.append({"email": mitt, "righe": toccate, "prova": False})
        if ignorate_proprie:
            print(f"  (ignorate {ignorate_proprie} email inviate dal centro stesso: "
                  f"non sono risposte dei clienti)")
        if not eventi:
            print("  nessun nuovo STOP.")
    except Exception as e:
        print(f"AVVISO: controllo risposte interrotto ({e.__class__.__name__}: {e}).")
        print("         Proseguo comunque con gli invii.")
    finally:
        try:
            posta.logout()
        except Exception:
            pass
    return eventi

# =====================================================================
# 4) TEMPLATE EMAIL (HTML + testo) e MESSAGGIO PERSONALIZZABILE
# =====================================================================
# Come funziona dalla v4.4 (zero toccate al codice per i centri):
#  - LE PAROLE: config.ini sezione [messaggio] -> oggetto, saluto, testo,
#    firma, riga extra del piede, istruzioni di disiscrizione. Dentro si
#    usano i segnaposto ({NOME}, {TARGA}, {SCADENZA}, {GIORNI}, {TRAFRA}
#    e tutti i dati del centro); nel testo una riga vuota = paragrafo
#    nuovo e **parola** = grassetto.
#  - IL DISEGNO: template_email.html (colori, logo, disposizione). Il
#    template riceve le parole gia' pronte: {SALUTO_HTML}, {TESTO_HTML},
#    {FIRMA_HTML}. Se il file e' della versione precedente viene
#    aggiornato da solo all'avvio (copia di sicurezza salvata in
#    template_email_vecchio.html); se manca, viene ricreato.

PERCORSO_TEMPLATE = os.path.join(BASE, "template_email.html")
PERCORSO_TEMPLATE_VECCHIO = os.path.join(BASE, "template_email_vecchio.html")
PERCORSO_LOGO = os.path.join(BASE, "logo.png")

TEMPLATE_DEFAULT_HTML = """\
<!DOCTYPE html>
<html>
<body style="margin:0;padding:0;background-color:#f4f5f7;font-family:Arial,Helvetica,sans-serif;">
<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background-color:#f4f5f7;padding:24px 0;">
<tr><td align="center">
<table role="presentation" width="600" cellpadding="0" cellspacing="0" style="background-color:#ffffff;border-radius:8px;overflow:hidden;">

<!-- INTESTAZIONE con colore e logo del centro -->
<tr><td style="background-color:{COLORE};padding:18px 32px;">
<table role="presentation" width="100%"><tr>
<td style="color:#ffffff;font-size:20px;font-weight:bold;background-color:{COLORE};">{NOME_CENTRO}</td>
<td align="right" style="background-color:{COLORE};">{LOGO_IMG}</td>
</tr></table>
</td></tr>

<!-- CORPO: le PAROLE arrivano da config.ini sezione [messaggio].
     Qui si cambia solo il DESIGN (colori, misure, disposizione). -->
<tr><td style="padding:32px;">
{SALUTO_HTML}
{TESTO_HTML}
{FIRMA_HTML}

<!-- CARTA GRANDE con giorni rimanenti -->
<table role="presentation" width="100%" cellpadding="0" cellspacing="0"
       style="background-color:#f8f9fa;border-left:4px solid {COLORE};border-radius:4px;">
<tr><td style="padding:18px 24px;">
<div style="font-size:12px;letter-spacing:1px;color:#888888;">GIORNI RIMANENTI</div>
<div style="font-size:38px;font-weight:bold;color:{COLORE};line-height:1.1;">{GIORNI}</div>
<div style="font-size:13px;color:#888888;">scadenza: {SCADENZA} &nbsp;|&nbsp; {TRAFRA}</div>
</td></tr></table>

</td></tr>

<!-- PIEDE -->
<tr><td style="padding:16px 32px;background-color:#fafafa;border-top:1px solid #eeeeee;">
<p style="margin:0;font-size:12px;color:#888888;line-height:1.6;">
<b>{NOME_CENTRO}</b>{CITTA_FOOTER} &nbsp;|&nbsp; Tel. {TELEFONO}{EMAIL_FOOTER}{INDIRIZZO_FOOTER}<br>
{ORARI_FOOTER}{RIGA_EXTRA_FOOTER}Hai ricevuto questa email perch&eacute; cliente di {NOME_CENTRO}.<br>
{TESTO_DISISCRIZIONE}<br>
<span style="font-size:10px;color:#bbbbbb;">Promemoria automatico inviato con {PRODOTTO} &copy; {ANNO}</span>
</p>
</td></tr>

</table>
</td></tr>
</table>
</body>
</html>
"""


def leggi_template():
    if os.path.exists(PERCORSO_TEMPLATE):
        try:
            with open(PERCORSO_TEMPLATE, "r", encoding="utf-8") as f:
                return f.read()
        except Exception as e:
            print(f"AVVISO: non riesco a leggere template_email.html ({e}):")
            print("        le email partiranno in versione testo semplice.")
            return None
    return None


def aggiorna_template():
    """Tiene template_email.html allineato: se manca lo crea dal modello
    interno; se e' della versione vecchia (senza {TESTO_HTML}, quindi non
    sa leggere il messaggio scritto in config.ini) lo aggiorna mettendo
    al sicuro la copia precedente in template_email_vecchio.html."""
    attuale = None
    if os.path.exists(PERCORSO_TEMPLATE):
        try:
            with open(PERCORSO_TEMPLATE, "r", encoding="utf-8") as f:
                attuale = f.read()
        except Exception:
            return
        if attuale is not None and "{TESTO_HTML}" in attuale:
            return                      # gia' aggiornato: niente da fare
    try:
        if attuale is not None:
            shutil.copy2(PERCORSO_TEMPLATE, PERCORSO_TEMPLATE_VECCHIO)
            with open(PERCORSO_TEMPLATE, "w", encoding="utf-8", newline="") as f:
                f.write(TEMPLATE_DEFAULT_HTML)
            print("template_email.html portato alla nuova versione: le PAROLE")
            print("delle email ora si scrivono in config.ini, sezione [messaggio].")
            print("La copia precedente e' salva come template_email_vecchio.html")
        else:
            with open(PERCORSO_TEMPLATE, "w", encoding="utf-8", newline="") as f:
                f.write(TEMPLATE_DEFAULT_HTML)
            print("template_email.html non c'era: l'ho ricreato (design standard).")
    except Exception as e:
        print(f"AVVISO: non riesco ad aggiornare template_email.html ({e}).")
        print("        Se esiste gia' un template vecchio, il messaggio scritto")
        print("        in config.ini appare solo nella versione testo dell'email.")


_LOGO_CACHE = None


def leggi_logo():
    """Legge logo.png una volta sola; avvisa se e' illeggibile o vuoto
    (cosi' mai email mute senza spiegazione, es. file OneDrive non sincronizzato)."""
    global _LOGO_CACHE
    if _LOGO_CACHE is None and os.path.exists(PERCORSO_LOGO):
        try:
            with open(PERCORSO_LOGO, "rb") as f:
                dati = f.read()
            if len(dati) < 100:
                print("AVVISO: logo.png e' vuoto o danneggiato: le email usciranno SENZA logo.")
                print("        Ricopialo dal pacchetto o mettici il logo del centro.")
            else:
                _LOGO_CACHE = dati
        except Exception as e:
            print(f"AVVISO: non riesco a leggere logo.png ({e}): email senza logo.")
    return _LOGO_CACHE


def sostituisci(testo, valori):
    for k, v in valori.items():
        testo = testo.replace(k, v)
    return testo


def togli_grassetto(testo):
    """Versione testo semplice: **parola** diventa solo parola."""
    return re.sub(r"\*\*(.+?)\*\*", r"\1", testo)


def testo_a_html(testo):
    """Il testo scritto in config.ini in HTML pronto per l'email:
    riga vuota = paragrafo nuovo, riga nuova = andare a capo,
    **parola** = grassetto. I caratteri speciali (&, <, >) vengono
    resi sicuri da soli (mai email rotte per colpa di una &)."""
    paragrafi = []
    for blocco in re.split(r"\n\s*\n", (testo or "").strip()):
        righe = []
        for r in blocco.split("\n"):
            r = html_lib.escape(r)
            r = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", r)
            righe.append(r)
        paragrafi.append("<br>".join(righe))
    stile = "margin:0 0 16px 0;font-size:15px;color:#333333;line-height:1.6;"
    return "".join(f'<p style="{stile}">{p}</p>' for p in paragrafi)


def componi_email(cliente, giorni, tipo=None):
    """Costruisce (oggetto, testo, html) dell'email per un cliente.
    Le parole vengono da config.ini [messaggio]; 'tipo' (30/10/1) fa
    usare l'oggetto speciale del singolo promemoria se il centro l'ha
    scritto (altrimenti si usa l'oggetto generale)."""
    data_scad = cliente["scadenza"].strftime("%d/%m/%Y")
    targa = cliente["targa"]
    if giorni > 1:
        tra = f"tra {giorni} giorni"
    else:
        tra = "tra 1 giorno"

    base = {
        "NOME": (cliente["nome"] or "").strip() or "Cliente",
        "TARGA": targa,
        "SCADENZA": data_scad,
        "GIORNI": str(giorni),
        "TRAFRA": tra,
        "NOME_CENTRO": NOME_CENTRO,
        "CITTA": CFG["citta"],
        "INDIRIZZO": CFG["indirizzo"],
        "TELEFONO": CFG["telefono"],
        "EMAIL_CENTRO": CFG["email_contatti"],
        "SITO_WEB": CFG["sito_web"],
        "ORARI": CFG["orari"],
        "COLORE": CFG["colore"],
        "PRODOTTO": PRODOTTO,
        "ANNO": str(dt.date.today().year),
    }
    valori_testo = {f"{{{k}}}": v for k, v in base.items()}
    valori_html = {k: html_lib.escape(v) for k, v in valori_testo.items()}

    modello_oggetto = (CFG["msg_oggetti_tipo"].get(tipo) if tipo else None) \
        or CFG["msg_oggetto"]
    oggetto = sostituisci(modello_oggetto, valori_testo)

    # Le parti del messaggio con i segnaposto gia' sostituiti: la versione
    # RAW serve per l'HTML (che trasforma **grassetto** in <b>), quella
    # "pulita" per il testo semplice (dove il grassetto non esiste)
    saluto_raw = sostituisci(CFG["msg_saluto"], valori_testo)
    corpo_raw = sostituisci(CFG["msg_testo"], valori_testo)
    firma_raw = sostituisci(CFG["msg_firma"], valori_testo)
    saluto = togli_grassetto(saluto_raw)
    corpo = togli_grassetto(corpo_raw)
    firma = togli_grassetto(firma_raw)

    testo = f"{saluto}\n\n{corpo}\n\n{firma}"
    if CFG["indirizzo"]:
        testo += f"\n{CFG['indirizzo']}"
    contatti = [v for v in (CFG["telefono"], CFG["email_contatti"],
                            CFG["sito_web"]) if v]
    if contatti:
        testo += "\nTel. " + " | ".join(contatti)
    if CFG["orari"]:
        testo += f"\nOrari: {CFG['orari']}"
    if CFG["msg_extra"]:
        testo += f"\n{CFG['msg_extra']}"
    testo += (f"\n---\n{CFG['msg_disiscrizione']}\n"
              f"Promemoria inviato con {PRODOTTO}\n")

    template = leggi_template()
    html = None
    if template:
        logo_presente = leggi_logo() is not None
        logo_img = (f'<img src="cid:logo" height="38" alt="logo {NOME_CENTRO}" '
                    'style="display:block;border:0;">') if logo_presente else ""
        citta_footer = (f" &nbsp;|&nbsp; {html_lib.escape(CFG['citta'])}"
                        if CFG["citta"] else "")
        email_footer = (f" &nbsp;|&nbsp; {html_lib.escape(CFG['email_contatti'])}"
                        if CFG["email_contatti"] else "")
        indirizzo_footer = (f" &nbsp;|&nbsp; {html_lib.escape(CFG['indirizzo'])}"
                            if CFG["indirizzo"] else "")
        orari_footer = (f"Orari: {html_lib.escape(CFG['orari'])}<br>"
                        if CFG["orari"] else "")
        extra_footer = (f"{html_lib.escape(CFG['msg_extra'])}<br>"
                        if CFG["msg_extra"] else "")

        if "{TESTO_HTML}" in template:
            # Template nuovo: le parole arrivano da config.ini [messaggio]
            valori = dict(valori_html)
            valori["{SALUTO_HTML}"] = testo_a_html(saluto_raw)
            valori["{TESTO_HTML}"] = testo_a_html(corpo_raw)
            valori["{FIRMA_HTML}"] = testo_a_html(firma_raw)
            valori["{LOGO_IMG}"] = logo_img
            valori["{CITTA_FOOTER}"] = citta_footer
            valori["{EMAIL_FOOTER}"] = email_footer
            valori["{INDIRIZZO_FOOTER}"] = indirizzo_footer
            valori["{ORARI_FOOTER}"] = orari_footer
            valori["{RIGA_EXTRA_FOOTER}"] = extra_footer
            valori["{TESTO_DISISCRIZIONE}"] = \
                html_lib.escape(CFG["msg_disiscrizione"])
            html = sostituisci(template, valori)
        else:
            # Template della versione precedente: sostituzione classica
            # (il testo di config.ini resta nella versione testo; il
            # template viene comunque aggiornato da solo all'avvio)
            giorni_label = f"{giorni} GIORNI" if giorni > 1 else "1 GIORNO"
            valori = {
                "{NOME_CENTRO}": NOME_CENTRO,
                "{CITTA}": CFG["citta"],
                "{TELEFONO}": CFG["telefono"],
                "{EMAIL_CENTRO}": CFG["email_contatti"],
                "{COLORE}": CFG["colore"],
                "{TARGA}": targa,
                "{SCADENZA}": data_scad,
                "{GIORNI}": str(giorni),
                "{GIORNI_LABEL}": giorni_label,
                "{TRAFRA}": tra,
                "{PRODOTTO}": PRODOTTO,
                "{ANNO}": str(dt.date.today().year),
                "{LOGO_IMG}": logo_img,
                "{CITTA_FOOTER}": citta_footer,
                "{EMAIL_FOOTER}": email_footer,
            }
            html = sostituisci(template, valori)
    return oggetto, testo, html

# =====================================================================
# 5) INVIO (una sola connessione, con retry e pausa)
# =====================================================================

class Sender:
    """Gestisce la connessione SMTP: si connette al primo invio e resta
    aperta; se cade si riconnette; ogni email ha 3 tentativi."""

    def __init__(self):
        self.server = None
        # Errore per cui riprovare e' inutile (password sbagliata, certificato
        # non valido, server senza cifratura): gli invii si fermano subito.
        self.errore_fatale = None

    def _connetti(self):
        if self.server is None:
            if SMTP_PORT == 587:
                # porta 587 = STARTTLS (Outlook e molti altri fornitori):
                # prima ci si collega in chiaro, poi si cripta la linea.
                # Se la cifratura non parte NON si prosegue: la password
                # non deve mai viaggiare in chiaro.
                srv = smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=30)
                try:
                    srv.ehlo()
                    srv.starttls(context=contesto_tls())
                    srv.ehlo()
                except BaseException:
                    try:
                        srv.close()
                    except Exception:
                        pass
                    raise
                self.server = srv
            else:
                # porta 465 (e altre) = SSL dalla prima riga (Gmail ecc.)
                self.server = smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT,
                                               timeout=30,
                                               context=contesto_tls())
            self.server.login(MITTENTE_EMAIL, MITTENTE_PASSWORD)

    def _chiudi(self):
        try:
            if self.server is not None:
                self.server.quit()
        except Exception:
            pass
        self.server = None

    def invia(self, destinatario, oggetto, testo, html=None, logo=None,
              con_stop=False):
        """Invia una email. Ritorna None se e' partita, altrimenti l'errore.
        con_stop=True (email ai CLIENTI) aggiunge l'intestazione
        List-Unsubscribe: Gmail e Outlook mostrano il bottone 'annulla
        iscrizione', che manda a questa casella un'email con oggetto STOP."""
        if self.errore_fatale is not None:
            return self.errore_fatale
        ultimo_errore = None
        for tentativo in range(1, 4):
            try:
                self._connetti()
                msg = EmailMessage()
                msg["From"] = formataddr((NOME_CENTRO, MITTENTE_EMAIL))
                msg["To"] = destinatario
                msg["Subject"] = oggetto
                msg["Date"] = formatdate(localtime=True)
                dominio = MITTENTE_EMAIL.split("@")[-1] if "@" in MITTENTE_EMAIL else None
                msg["Message-ID"] = make_msgid(domain=dominio)
                if con_stop and MITTENTE_EMAIL:
                    msg["List-Unsubscribe"] = f"<mailto:{MITTENTE_EMAIL}?subject=STOP>"
                if RISPOSTA_A:
                    msg["Reply-To"] = RISPOSTA_A
                msg.set_content(testo)
                if html:
                    msg.add_alternative(html, subtype="html")
                    if logo:
                        # Content-ID DEVE avere le parentesi angolari (<logo>):
                        # e' lo standard RFC 2392 per src="cid:logo". Senza
                        # parentesi Gmail e altri client NON mostrano il logo.
                        msg.get_payload()[1].add_related(
                            logo, maintype="image", subtype="png", cid="<logo>")
                self.server.send_message(msg)
                return None
            except smtplib.SMTPAuthenticationError as e:
                # password per app / utente sbagliati: riprovare 3 volte per
                # ogni cliente non serve e rischia di far bloccare la casella
                self._chiudi()
                self.errore_fatale = e
                return e
            except smtplib.SMTPNotSupportedError as e:
                self._chiudi()          # niente STARTTLS: mai password in chiaro
                self.errore_fatale = e
                return e
            except ssl.SSLCertVerificationError as e:
                self._chiudi()
                self.errore_fatale = e
                return e
            except smtplib.SMTPRecipientsRefused as e:
                self._chiudi()          # indirizzo rifiutato dal server
                return e
            except smtplib.SMTPResponseException as e:
                ultimo_errore = e
                self._chiudi()
                if 500 <= e.smtp_code < 600:
                    return e            # errore definitivo: riprovare non serve
                if tentativo < 3:
                    time.sleep(3 * tentativo)
            except (smtplib.SMTPException, OSError) as e:
                ultimo_errore = e
                self._chiudi()
                if tentativo < 3:
                    time.sleep(3 * tentativo)
            except ValueError as e:
                # es. "a capo" dentro un campo (cella Excel con Alt+Invio):
                # l'intestazione non e' valida. Questo cliente salta, gli
                # altri partono lo stesso (prima il programma si fermava).
                self._chiudi()
                return e
        return ultimo_errore

    def chiudi(self):
        self._chiudi()

# =====================================================================
# 6) BACKUP AUTOMATICO DEL FILE CLIENTI
# =====================================================================

BACKUP_ULTIMI = 20          # gli ultimi N backup, qualunque giorno
BACKUP_GIORNI = 30          # + il piu' recente di ognuno degli ultimi N giorni


def pulisci_backup(cartella):
    """Ritenzione dei backup. Prima si tenevano solo i 14 piu' RECENTI in
    assoluto: ogni salvataggio del pannello fa un backup, quindi dopo 14
    modifiche nella stessa giornata sparivano TUTTI i backup dei giorni
    prima. Ora: gli ultimi 20 + l'ultimo di ogni giorno degli ultimi 30."""
    files = [os.path.join(cartella, f) for f in os.listdir(cartella)
             if f.startswith("clienti_") and f.endswith(".xlsx")]
    files.sort(key=os.path.getmtime, reverse=True)
    tieni = set(files[:BACKUP_ULTIMI])
    giorni_visti = set()
    limite = time.time() - BACKUP_GIORNI * 86400
    for p in files:                      # dal piu' recente al piu' vecchio
        m = os.path.getmtime(p)
        if m < limite:
            continue
        giorno = dt.date.fromtimestamp(m)
        if giorno not in giorni_visti:
            giorni_visti.add(giorno)
            tieni.add(p)
    for p in files:
        if p not in tieni:
            try:
                os.remove(p)
            except OSError:
                pass


def backup_excel(percorso_excel):
    try:
        cartella = os.path.join(BASE, "backup")
        os.makedirs(cartella, exist_ok=True)
        nome = f"clienti_{dt.date.today().isoformat()}_{dt.datetime.now().strftime('%H%M')}.xlsx"
        shutil.copy2(percorso_excel, os.path.join(cartella, nome))
        pulisci_backup(cartella)
        return True
    except Exception as e:
        print(f"  (backup non riuscito: {e} - continuo comunque)")
        return False

# =====================================================================
# 7) RIEPILOGO GIORNALIERO AL TITOLARE
# =====================================================================

def invia_riepilogo(sender, inviati, errori, esclusi, scadute, dry_run,
                    eventi_stop=None):
    eventi_stop = eventi_stop or []
    if not CFG["notifica_titolare"]:
        return
    if not inviati and not errori and not eventi_stop:
        return                      # niente da segnalare: zero email al titolare
    oggi = dt.date.today().strftime("%d/%m/%Y")
    destinatario = CFG["email_titolare"] or MITTENTE_EMAIL

    # Tutto cio' che viene dall'Excel o dagli errori va "reso sicuro": un
    # nome con & o < (o un messaggio d'errore con <...>) non deve rompere
    # la tabella del riepilogo.
    esc = html_lib.escape
    righe_html = ""
    for i in inviati:
        righe_html += (f"<tr><td style='padding:6px 10px;border-bottom:1px solid #eee;'>{esc(str(i['targa']))}</td>"
                       f"<td style='padding:6px 10px;border-bottom:1px solid #eee;'>{esc(str(i['email']))}</td>"
                       f"<td style='padding:6px 10px;border-bottom:1px solid #eee;'>{esc(str(i['desc']))}</td></tr>")
    righe_err = ""
    for e in errori:
        righe_err += (f"<tr><td style='padding:6px 10px;border-bottom:1px solid #eee;'>{esc(str(e['targa']))}</td>"
                      f"<td style='padding:6px 10px;border-bottom:1px solid #eee;color:#C0392B;'>{esc(str(e['errore']))}</td></tr>")
    righe_stop = ""
    for s in eventi_stop:
        dettaglio = ", ".join(
            f"{n or 'cliente'} (riga {r}, targa {t or '?'})" for r, n, t in s.get("righe", []))
        righe_stop += (f"<tr><td style='padding:6px 10px;border-bottom:1px solid #eee;'>{esc(str(s['email']))}</td>"
                       f"<td style='padding:6px 10px;border-bottom:1px solid #eee;'>{esc(dettaglio or 'richiesta STOP')}</td></tr>")

    html = f"""<div style="font-family:Arial,sans-serif;color:#333;max-width:600px;">
  <h2 style="color:{CFG['colore']};margin-bottom:4px;">Riepilogo promemoria - {oggi}</h2>
  <p style="margin-top:0;">Generato automaticamente da {PRODOTTO}</p>
  <p><b>{len(inviati)}</b> promemoria inviati &nbsp;|&nbsp;
     <b>{len(eventi_stop)}</b> richieste STOP &nbsp;|&nbsp;
     <b>{len(errori)}</b> errori &nbsp;|&nbsp;
     <b>{esclusi}</b> clienti esclusi &nbsp;|&nbsp;
     <b>{scadute}</b> revisioni gia' scadute</p>
  {f"<h3>Invii di oggi</h3><table style='border-collapse:collapse;width:100%;'><tr style='background:{CFG['colore']};color:#fff;'><th style='padding:6px 10px;text-align:left;'>Targa</th><th style='padding:6px 10px;text-align:left;'>Cliente</th><th style='padding:6px 10px;text-align:left;'>Promemoria</th></tr>{righe_html}</table>" if righe_html else ""}
  {f"<h3>Errori da controllare</h3><table style='border-collapse:collapse;width:100%;'><tr style='background:#C0392B;color:#fff;'><th style='padding:6px 10px;text-align:left;'>Targa</th><th style='padding:6px 10px;text-align:left;'>Errore</th></tr>{righe_err}</table>" if righe_err else ""}
  {f"<h3>Clienti che hanno risposto STOP (ora in pausa)</h3><table style='border-collapse:collapse;width:100%;'><tr style='background:#7f8c8d;color:#fff;'><th style='padding:6px 10px;text-align:left;'>Email</th><th style='padding:6px 10px;text-align:left;'>Dettaglio</th></tr>{righe_stop}</table>" if righe_stop else ""}
  <p style="font-size:12px;color:#888;">Chi risponde STOP viene messo in pausa da solo (Attivo = STOP nell'Excel). Per riattivare un cliente: rimetti Attivo = SÌ.</p>
</div>"""
    testo = (f"Riepilogo promemoria del {oggi}\n"
             f"Inviati: {len(inviati)} | Errori: {len(errori)} | "
             f"Esclusi: {esclusi} | Scadute: {scadute}")
    oggetto = (f"[{PRODOTTO}] Riepilogo del {oggi} - "
               f"{len(inviati)} promemoria inviati")
    if dry_run:
        print("\n  [RIEPILOGO TITOLARE - solo PROVA]")
        print(f"  Oggetto: {oggetto}")
        print(f"  A: {destinatario}")
        print(f"  Inviati: {len(inviati)}, Errori: {len(errori)}, STOP: {len(eventi_stop)}, "
              f"Esclusi: {esclusi}, Scadute: {scadute}")
        return
    errore = sender.invia(destinatario, oggetto, testo, html)
    if errore:
        print(f"  ATTENZIONE: riepilogo al titolare non inviato ({errore})")

# =====================================================================
# 8) PROGRAMMA PRINCIPALE
# =====================================================================

def trova_tipo(giorni):
    for tipo, (da, a) in CFG["finestre"].items():
        if da <= giorni <= a:
            return tipo
    return None


# --- colori automatici del file Excel --------------------------------
# Solo le CELLE di Revisione Fatta e Attivo si colorano: SÌ = verde,
# NO = rosso, STOP = rosso scuro. STOP NON e' nel menu a tendina: lo
# scrive solo il programma quando un cliente lo chiede. Sono regole di
# FORMATTAZIONE CONDIZIONALE: stanno nel file e cambiano da sole quando
# scrivi in Excel (il programma non serve).
COLORI_EXCEL = {
    "si":   ("C6EFCE", "006100"),    # sfondo, testo
    "no":   ("FFC7CE", "9C0006"),
    "stop": ("C00000", "FFFFFF"),
}


def applica_colori(ws, colonna_di, ultima_riga=300):
    """Regole di colore SOLO sulle celle delle colonne Revisione Fatta
    (SÌ verde, NO rosso) e Attivo (SÌ verde, NO rosso, STOP rosso
    scuro). Usa OR(UPPER(...)) cosi' contano anche SI senza accento e
    le minuscole."""
    prima = RIGA_INTESTAZIONI + 1
    col_fatta = get_column_letter(colonna_di["fatta"])
    col_attivo = get_column_letter(colonna_di["attivo"])
    si_sf, si_tx = COLORI_EXCEL["si"]
    no_sf, no_tx = COLORI_EXCEL["no"]
    stop_sf, stop_tx = COLORI_EXCEL["stop"]

    def aggiungi(intervallo, formula, sfondo, testo=None, grassetto=False):
        riempimento = PatternFill(start_color=sfondo, end_color=sfondo,
                                  fill_type="solid")
        ws.conditional_formatting.add(intervallo, FormulaRule(
            formula=[formula],
            fill=riempimento,
            font=Font(color=testo, bold=grassetto) if testo else None))

    rng_fatta = f"{col_fatta}{prima}:{col_fatta}{ultima_riga}"
    aggiungi(rng_fatta,
             f'OR(UPPER({col_fatta}{prima})="SÌ",UPPER({col_fatta}{prima})="SI")',
             si_sf, si_tx)
    aggiungi(rng_fatta, f'UPPER({col_fatta}{prima})="NO"', no_sf, no_tx)
    rng_attivo = f"{col_attivo}{prima}:{col_attivo}{ultima_riga}"
    aggiungi(rng_attivo,
             f'OR(UPPER({col_attivo}{prima})="SÌ",UPPER({col_attivo}{prima})="SI")',
             si_sf, si_tx)
    aggiungi(rng_attivo, f'UPPER({col_attivo}{prima})="NO"', no_sf, no_tx)
    aggiungi(rng_attivo, f'UPPER({col_attivo}{prima})="STOP"',
             stop_sf, stop_tx, grassetto=True)


def crea_excel_se_manca(percorso_excel):
    """Se il file clienti non c'e', lo crea qui: intestazioni alla riga 4,
    8 clienti di esempio, menu SÌ/NO e colori. Cosi' nel pacchetto non ci
    sono file doppi e aggiornare il programma non cancella mai i clienti."""
    if os.path.exists(percorso_excel):
        return
    try:
        def mese_libretto(giorni_desiderati):
            """'MM/YYYY' del mese il cui ULTIMO giorno cade a ~N giorni da oggi."""
            d = dt.date.today() + dt.timedelta(days=giorni_desiderati)
            fine = ultimo_giorno_mese(d.year, d.month)
            delta = (fine - dt.date.today()).days
            if delta < giorni_desiderati - 6:
                y, m = (d.year + 1, 1) if d.month == 12 else (d.year, d.month + 1)
                fine = ultimo_giorno_mese(y, m)
            elif delta > giorni_desiderati + 6:
                y, m = (d.year - 1, 12) if d.month == 1 else (d.year, d.month - 1)
                fine = ultimo_giorno_mese(y, m)
            return f"{fine.month:02d}/{fine.year}"

        esempi = [
            # (nome, email, telefono, targa, scadenza, revisione fatta, attivo)
            ("Mario Rossi",   "mario.rossi@email.it",   "347 1122334", "AB123CD", mese_libretto(29),  "NO", "SÌ"),
            ("Luigi Verdi",   "luigi.verdi@email.it",   "339 5566778", "EF456GH", mese_libretto(9),   "NO", "SÌ"),
            ("Anna Bianchi",  "anna.bianchi@email.it",  "331 8899001", "JK789LM", mese_libretto(2),   "NO", "SÌ"),
            ("Carlo Neri",    "carlo.neri@email.it",    "320 2233445", "OP321QR", mese_libretto(55),  "NO", "SÌ"),
            ("Sofia Gialli",  "sofia.gialli@email.it",  "349 6677889", "ST654UV", mese_libretto(-5),  "NO", "SÌ"),
            ("Paolo Marroni", "paolo.marroni@email.it", "333 4455667", "WX987YZ", mese_libretto(29),  "SÌ", "SÌ"),
            ("Elena Azzurri", "elena.azzurri@email.it", "346 7788990", "ZT147HP", mese_libretto(40),  "NO", "NO"),
            ("Franco Viola",  "franco.viola@email.it",  "328 2244668", "QR258WX", mese_libretto(100), "NO", "SÌ"),
        ]

        wb = Workbook()
        ws = wb.active
        ws.title = FOGLIO

        ws.cell(row=1, column=2, value="PROMEMORIA REVISIONI - ELENCO CLIENTI").font = \
            Font(bold=True, size=14)
        nota = ("Una riga per VEICOLO (stesso cliente con 2 macchine = 2 righe). Scrivi dalla riga 5 "
                "e cancella gli esempi. Scadenza come sul libretto: MESE/ANNO (es. 09/2026 = fino al 30/09/2026). "
                "Revisione Fatta = SÌ ferma le email finche' non aggiorni il mese. Attivo = NO esclude il cliente; "
                "STOP significa che il cliente l'ha chiesto rispondendo STOP all'email. "
                "I colori si mettono da soli nelle celle: verde = SÌ, rosso = NO, "
                "rosso scuro = STOP (lo scrive il programma).")
        ws.cell(row=2, column=2, value=nota).alignment = \
            Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=2, end_row=3, end_column=8)

        intestazioni = (COLONNE["nome"], COLONNE["email"], COLONNE["telefono"],
                        COLONNE["targa"], COLONNE["scadenza"], COLONNE["fatta"],
                        COLONNE["attivo"])
        for i, titolo in enumerate(intestazioni, start=2):
            cella = ws.cell(row=RIGA_INTESTAZIONI, column=i, value=titolo)
            cella.font = Font(bold=True, color="FFFFFF")
            cella.fill = PatternFill("solid", fgColor="1F3B66")
            cella.alignment = Alignment(horizontal="center", vertical="center")

        for r, (nome, mitt, tel, targa, scad, fatta, attivo) in enumerate(esempi, start=RIGA_INTESTAZIONI + 1):
            for c, v in enumerate((nome, mitt, tel, targa, scad, fatta, attivo), start=2):
                ws.cell(row=r, column=c, value=v)

        for colonna, larghezza in zip("BCDEFGH", (22, 28, 14, 11, 18, 16, 10)):
            ws.column_dimensions[colonna].width = larghezza

        tendina = DataValidation(type="list", formula1='"SÌ,NO"', allow_blank=True)
        ws.add_data_validation(tendina)
        tendina.add(f"G{RIGA_INTESTAZIONI + 1}:G300")
        tendina.add(f"H{RIGA_INTESTAZIONI + 1}:H300")

        colonna_di = {"nome": 2, "email": 3, "telefono": 4, "targa": 5,
                      "scadenza": 6, "fatta": 7, "attivo": 8}
        applica_colori(ws, colonna_di, ultima_riga=300)

        # Colonna Scadenza in formato TESTO: scrivendo 01/2026 resta
        # 01/2026 (Excel non la trasforma piu' in una data tipo "gen-26")
        for r in range(RIGA_INTESTAZIONI + 1, 301):
            ws.cell(row=r, column=6).number_format = "@"

        ws.freeze_panes = "A5"
        salva_excel_sicuro(wb, percorso_excel)
        print(f"File clienti non trovato: l'ho creato con 8 clienti di esempio "
              f"({os.path.basename(percorso_excel)}).")
        print("Aprilolo, cancella gli esempi e scrivi i tuoi clienti dalla riga 5.\n")
    except Exception as e:
        print(f"AVVISO: non riesco a creare il file clienti ({e}).")
        print("Chiudi Excel se e' aperto e rilancia il programma.\n")


def migliora_excel(percorso_excel):
    """Mette (o aggiorna) i colori in un file Excel GIA' esistente senza
    toccare i dati: celle SÌ verde, NO rosso, STOP rosso scuro, tendina
    SÌ/NO (STOP non e' un'opzione del menu) e colonna Scadenza in formato
    testo. Serve per i file creati con le versioni vecchie del programma.
    Sotto la serratura condivisa col pannello: niente colori a meta'."""
    serratura = SerraturaExcel()
    if not serratura.prendi():
        print("ERRORE: il file clienti e' occupato da un'altra scrittura (pannello?).")
        print("Riprova colora_excel.bat tra qualche secondo.")
        return False
    try:
        return _migliora_excel_serrato(percorso_excel)
    finally:
        serratura.rilascia()


def _migliora_excel_serrato(percorso_excel):
    """(i colori veri, sotto serratura)"""
    if not os.path.exists(percorso_excel):
        print(f"ERRORE: non trovo il file Excel: {percorso_excel}")
        return False
    try:
        wb = load_workbook(percorso_excel)
    except PermissionError:
        print("ERRORE: il file Excel risulta APERTO in Excel (o bloccato da OneDrive).")
        print("Chiudi Excel, aspetta qualche secondo e rilancia colora_excel.bat.")
        return False
    except Exception as e:
        print(f"ERRORE: il file Excel non e' leggibile ({e}).")
        return False
    ws = wb[FOGLIO] if FOGLIO in wb.sheetnames else wb.active
    colonna_di, mancanti = individua_colonne(ws)
    if mancanti:
        print("ERRORE: nel foglio mancano le colonne: " + ", ".join(mancanti))
        print("Le intestazioni devono stare alla riga 4: "
              + " | ".join(COLONNE.values()))
        return False

    prima = RIGA_INTESTAZIONI + 1
    ultima = max(ws.max_row + 50, 300)
    colonna_scad = colonna_di["scadenza"]

    # 1) scadenze trasformate in data da Excel (es. "gen-26"): le
    #    riscrivo come MESE/ANNO e metto la colonna in formato testo
    convertite = 0
    for r in range(prima, ultima + 1):
        cella = ws.cell(row=r, column=colonna_scad)
        v = cella.value
        if isinstance(v, dt.datetime):
            v = v.date()
        if isinstance(v, dt.date):
            if v.day == 1:                      # es. "01/2026" -> 1 gen
                cella.value = f"{v.month:02d}/{v.year}"
                cella.number_format = "@"
                convertite += 1
            # date complete (giorno != 1): le lascio, il programma
            # le accetta comunque
        else:
            cella.number_format = "@"

    # 2) tendina SÌ / NO su Revisione Fatta e Attivo (STOP non e' una
    #    opzione: lo scrive solo il programma quando il cliente lo chiede;
    #    qui sistema anche i file toccati dalla versione sbagliata di oggi)
    giusta = '"SÌ,NO"'
    trovata = False
    for dv in ws.data_validations.dataValidation:
        copre = False
        for intervallo in dv.sqref.ranges:
            if intervallo.max_row >= prima and (
                    intervallo.min_col <= colonna_di["fatta"] <= intervallo.max_col
                    or intervallo.min_col <= colonna_di["attivo"] <= intervallo.max_col):
                copre = True
                break
        if copre:
            dv.formula1 = giusta
            dv.showErrorMessage = False   # il menu e' solo un aiuto
            trovata = True
    if not trovata:
        tendina = DataValidation(type="list", formula1=giusta, allow_blank=True)
        ws.add_data_validation(tendina)
        for chiave in ("fatta", "attivo"):
            lettera = get_column_letter(colonna_di[chiave])
            tendina.add(f"{lettera}{prima}:{lettera}{ultima}")

    # 3) colori (prima tolgo le regole vecchie, per non accumularne)
    ws.conditional_formatting = ConditionalFormattingList()
    applica_colori(ws, colonna_di, ultima_riga=ultima)

    # 4) intestazioni bloccate + larghezze colonne se mai impostate
    if not ws.freeze_panes:
        ws.freeze_panes = "A5"
    for chiave, larghezza in (("nome", 22), ("email", 28), ("telefono", 14),
                              ("targa", 11), ("scadenza", 18),
                              ("fatta", 16), ("attivo", 10)):
        lettera = get_column_letter(colonna_di[chiave])
        if lettera not in ws.column_dimensions \
                or ws.column_dimensions[lettera].width is None:
            ws.column_dimensions[lettera].width = larghezza

    try:
        salva_excel_sicuro(wb, percorso_excel)
    except PermissionError:
        print("ERRORE: il file Excel risulta APERTO: non posso salvarlo.")
        print("Chiudi Excel e rilancia colora_excel.bat.")
        return False
    print("Fatto! Colori aggiornati nel file Excel:")
    print(f"  - {os.path.basename(percorso_excel)}")
    print("  - celle: SÌ = verde, NO = rosso, STOP = rosso scuro")
    if convertite:
        print(f"  - riscritte {convertite} scadenze da data Excel a MESE/ANNO (es. 01/2026)")
    print("Chiudi e riapri il file per vedere i colori.")
    return True


LOCK_SCADUTO_SECONDI = 15 * 60     # blocco illeggibile da cosi' tanto = abbandonato


def percorso_lock():
    return os.path.join(BASE, "revmemo.lock")


def _pid_vivo(pid):
    """True se esiste un processo con quel numero ADESSO (Windows e Linux).
    Stesse regole del pannello (pannello.py): su Windows si controlla anche
    che il processo non sia gia' terminato (i numeri vengono riusati)."""
    try:
        pid = int(pid)
    except (TypeError, ValueError):
        return False
    if pid <= 0:
        return False
    if os.name == "nt":
        try:
            import ctypes
            k32 = ctypes.windll.kernel32
            handle = k32.OpenProcess(0x1000, False, pid)
            if not handle:
                return False       # non esiste piu' (o non si puo' guardare)
            try:
                codice_uscita = ctypes.c_ulong()
                if not k32.GetExitCodeProcess(handle, ctypes.byref(codice_uscita)):
                    return False
                return codice_uscita.value == 259      # STILL_ACTIVE
            finally:
                k32.CloseHandle(handle)
        except Exception:
            return False
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    except PermissionError:
        return True        # esiste ma e' di un altro utente: e' vivo
    except OSError:
        return False
    return True


def _leggi_pid_lock(percorso):
    """Il numero di processo scritto dentro un file di blocco (int).
    None se illeggibile o senza numero."""
    try:
        with open(percorso, "r", encoding="utf-8") as f:
            testo = (f.read() or "").strip()
    except OSError:
        return None
    return int(testo) if testo.isdigit() else None


class BloccoInvio:
    """Impedisce DUE invii contemporanei (es. il pannello e l'attivita'
    pianificata di Windows partiti insieme): senza, entrambi vedrebbero
    "non ancora inviato" e il cliente riceverebbe l'email doppia.
    Il blocco e' un file che contiene il numero del programma che sta
    inviando e che viene "toccato" a ogni cliente (battito).
    Chi lo trova lo lascia STARE se il proprietario e' VIVO (anche se
    vecchio: un invio lento non viene mai rubato); lo riprende da solo
    se il proprietario e' morto (subito, senza attese) o se il blocco e'
    illeggibile da troppo tempo (15 minuti)."""

    def __init__(self):
        self.preso = False

    def prendi(self):
        percorso = percorso_lock()
        for _ in range(3):
            try:
                fd = os.open(percorso, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            except FileExistsError:
                try:
                    eta = time.time() - os.path.getmtime(percorso)
                except OSError:
                    continue                  # sparito proprio ora: riprovo
                pid = _leggi_pid_lock(percorso)
                if pid is not None and _pid_vivo(pid):
                    return False     # l'altro invio e' VIVO: non si tocca
                if pid is None and eta < LOCK_SCADUTO_SECONDI:
                    return False     # blocco recente ma illeggibile: rispetto
                try:
                    os.remove(percorso)   # orfano (morto) o abbandonato
                except OSError:
                    return False
                continue
            except OSError:
                return True    # cartella non scrivibile: meglio lavorare che bloccarsi
            with os.fdopen(fd, "w") as f:
                f.write(str(os.getpid()))
            self.preso = True
            return True
        return False

    def batti(self):
        if self.preso:
            try:
                os.utime(percorso_lock(), None)
            except OSError:
                pass

    def rilascia(self):
        if self.preso:
            try:
                os.remove(percorso_lock())
            except OSError:
                pass
            self.preso = False


# --- Serratura del file CLIENTI (condivisa fra pannello e motore) -------
# Il pannello e il motore sono DUE programmi separati: la serratura interna
# del pannello tiene in fila i suoi thread, questa tiene in fila anche
# l'invio automatico. Vale solo per le SCRITTURE (breve: si prende prima
# di caricare il file e si lascia dopo il salvataggio); le letture non la
# chiedono perche' il salvataggio e' atomico (salva_excel_sicuro: chi
# legge vede il file vecchio o quello nuovo, mai uno scritto a meta').

SERRATURA_EXCEL_ATTESA = 15     # secondi che chi scrive aspetta il suo turno
SERRATURA_EXCEL_SCADUTA = 120   # blocco illeggibile e piu' vecchio di cosi' = abbandonato


def percorso_serratura_excel():
    return os.path.join(BASE, "excel.lock")


class SerraturaExcel:
    """Fa salvare il file clienti UNO ALLA VOLTA anche a programmi diversi
    (pannello contro motore). Regola come per BloccoInvio: a un
    proprietario VIVO non si ruba mai; un proprietario morto si riprende
    subito, senza far aspettare nessuno."""

    def __init__(self, attesa=SERRATURA_EXCEL_ATTESA):
        self.preso = False
        try:
            self._attesa = max(1, int(attesa))
        except (TypeError, ValueError):
            self._attesa = SERRATURA_EXCEL_ATTESA

    def prendi(self):
        percorso = percorso_serratura_excel()
        inizio = time.time()
        while True:
            try:
                fd = os.open(percorso, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            except FileExistsError:
                self._se_orfano_riprendi(percorso)
                if time.time() - inizio >= self._attesa:
                    return False
                time.sleep(0.25)
                continue
            except OSError:
                return True    # cartella non scrivibile: meglio lavorare
            with os.fdopen(fd, "w") as f:
                f.write(str(os.getpid()))
            self.preso = True
            return True

    def _se_orfano_riprendi(self, percorso):
        try:
            eta = time.time() - os.path.getmtime(percorso)
        except OSError:
            return                     # sparito proprio ora: nulla da fare
        pid = _leggi_pid_lock(percorso)
        if pid is not None and _pid_vivo(pid):
            return                     # proprietario vivo: aspetta il turno
        if pid is None and eta < SERRATURA_EXCEL_SCADUTA:
            return                     # illeggibile ma recente: aspetta
        try:
            os.remove(percorso)        # orfano (morto) o abbandonato
        except OSError:
            pass

    def rilascia(self):
        if not self.preso:
            return
        try:
            if _leggi_pid_lock(percorso_serratura_excel()) in (os.getpid(), None):
                os.remove(percorso_serratura_excel())
        except OSError:
            pass
        self.preso = False


def salva_excel_sicuro(wb, percorso_excel):
    """Salvataggio ATOMICO del file clienti: prima su un file provvisorio,
    poi lo scambia con quello vero. Chi legge nel frattempo (pannello,
    backup) vede il file vecchio o quello nuovo, MAI uno scritto a meta'.
    Se il file e' aperto in Excel (o e' protetto) fallisce con
    PermissionError, esattamente come il salvataggio di prima."""
    if os.path.exists(percorso_excel):
        # posso scrivere il file vero? Se e' aperto in Excel o protetto
        # esce qui con PermissionError, come ci si aspetta da sempre
        with open(percorso_excel, "ab"):
            pass
    provvisorio = percorso_excel + ".nuovo"
    try:
        wb.save(provvisorio)
        os.replace(provvisorio, percorso_excel)
    finally:
        try:
            if os.path.exists(provvisorio):
                os.remove(provvisorio)
        except OSError:
            pass


def main():
    parser = argparse.ArgumentParser(description="RevMemo - promemoria revisioni")
    parser.add_argument("--dry-run", action="store_true",
                        help="modalita' prova: mostra senza inviare")
    parser.add_argument("--file", default=None,
                        help="file clienti alternativo")
    parser.add_argument("--colora-excel", action="store_true",
                        help="mette/aggiorna i colori nel file clienti ed esce")
    args = parser.parse_args()

    percorso_excel = args.file or os.path.join(BASE, FILE_EXCEL)
    if args.colora_excel:
        sys.exit(0 if migliora_excel(percorso_excel) else 1)
    crea_excel_se_manca(percorso_excel)
    aggiorna_template()      # crea/aggiorna il design email (con backup)

    print("=" * 62)
    print(f"  {PRODOTTO} v{VERSIONE} - Promemoria revisioni",
          "(MODALITA' PROVA - nessuna email partira')" if args.dry_run else "")
    print("=" * 62)
    print(f"Oggi: {dt.date.today().strftime('%d/%m/%Y')}")
    print(f"Centro: {NOME_CENTRO}   |   Mittente: {MITTENTE_EMAIL or '(non impostato)'}")
    print(f"File clienti: {percorso_excel}")
    if args.dry_run and not MITTENTE_PASSWORD:
        print("Nota: password non ancora inserita in config.ini - la prova funziona lo stesso")
    print()

    if not args.dry_run:
        if not MITTENTE_EMAIL or not MITTENTE_PASSWORD:
            print("ERRORE: compila config.ini sezione [smtp] prima dell'invio vero:")
            print("  mittente = la casella da cui inviare")
            print("  password_app = la Password per app di Google (16 caratteri)")
            print("Guida: LEGGIMI.txt, sezione 5")
            sys.exit(1)
        blocco = BloccoInvio()
        if not blocco.prendi():
            print("Un altro invio e' gia' in corso (pannello o attivita' pianificata):")
            print("aspetta che finisca, cosi' nessun cliente riceve l'email doppia.")
            print("Se invece non sta girando niente, riprova tra qualche minuto.")
            return 2
        try:
            backup_excel(percorso_excel)
            return _giro(args, percorso_excel, blocco)
        finally:
            blocco.rilascia()
    return _giro(args, percorso_excel)


def _giro(args, percorso_excel, blocco=None):
    """Il lavoro vero: STOP dei clienti, invii, riepilogo al titolare."""
    # Prima di inviare: leggo la posta in arrivo per gli STOP dei clienti
    eventi_stop = controlla_risposte(percorso_excel, args.dry_run) or []
    print()

    clienti = leggi_clienti(percorso_excel)
    print(f"Trovati {len(clienti)} clienti nel file.\n")

    sender = Sender()
    inviati, errori = [], []
    esclusi = scadute = saltati_log = 0
    oggi = dt.date.today()

    for c in clienti:
        if blocco:
            blocco.batti()
        targa = c["targa"] or "(targa mancante)"
        riga_lbl = f"riga {c['riga']:>2} | {(c['nome'] or '(senza nome)')[:18]:<18} | {targa:<8} | "

        if not c["attivo"]:
            motivo = ("ha chiesto STOP -> nessuna email"
                      if c["attivo_testo"] == "STOP"
                      else "cliente NON attivo -> saltato")
            print(f"  - {riga_lbl}{motivo}")
            esclusi += 1
            continue
        if c["fatta"]:
            print(f"  - {riga_lbl}revisione GIA' FATTA -> nessuna email")
            esclusi += 1
            continue
        if not email_valida(c["email"]):
            print(f"  - {riga_lbl}email mancante o non valida -> saltato")
            esclusi += 1
            continue
        if c["scadenza"] is None:
            print(f"  - {riga_lbl}scadenza non leggibile (scrivi es. 09/2026) -> saltato")
            esclusi += 1
            continue

        giorni = (c["scadenza"] - oggi).days
        if giorni < 0:
            print(f"  - {riga_lbl}revisione SCADUTA da {-giorni} gg -> nessun invio")
            scadute += 1
            esclusi += 1
            continue
        if giorni == 0:
            print(f"  - {riga_lbl}scade OGGI: ultimo giorno utile, nessun promemoria")
            esclusi += 1
            continue

        tipo = trova_tipo(giorni)
        if tipo is None:
            print(f"  - {riga_lbl}mancano {giorni:>3} gg -> non e' momento di inviare")
            esclusi += 1
            continue

        descrizione = CFG["etichette"][tipo]
        oggetto, testo, html = componi_email(c, giorni, tipo)

        scadenza_iso = c["scadenza"].isoformat()
        if gia_inviato(c["email"], c["targa"], tipo, scadenza_iso):
            print(f"  - {riga_lbl}promemoria '{descrizione}' GIA' INVIATO -> saltato")
            saltati_log += 1
            continue

        if args.dry_run:
            print(f"  > {riga_lbl}PROVA invio '{descrizione}' a {c['email']}")
            print(f"      Oggetto: {oggetto}")
            print(f"      Email HTML con logo: {'SI' if html else 'no (template mancante)'}")
            print(f"      Messaggio scritto in config.ini sezione [messaggio]")
            inviati.append({"targa": targa, "email": c["email"], "desc": descrizione})
            continue

        errore = sender.invia(c["email"], oggetto, testo, html, leggi_logo(),
                              con_stop=True)
        if errore is None:
            try:
                registra_invio(c["email"], c["targa"], tipo, giorni, descrizione, scadenza_iso)
            except OSError as e:
                # l'email E' partita ma il registro non si scrive: domani
                # potrebbe ripartire uguale. Meglio dirlo forte.
                print(f"  ATTENZIONE: email inviata ma NON scritta nel registro ({e}).")
                print(f"              Rischio doppione: controlla {percorso_log()}")
            print(f"  [OK] {riga_lbl}email inviata a {c['email']} ({descrizione})")
            inviati.append({"targa": targa, "email": c["email"], "desc": descrizione})
        else:
            print(f"  [X] {riga_lbl}ERRORE: {errore}")
            errori.append({"targa": targa, "errore": str(errore)})
            if sender.errore_fatale is not None:
                print()
                print("  Gli invii si FERMANO qui: la casella non accetta il collegamento")
                print("  (password per app o server sbagliati, oppure collegamento non sicuro).")
                if causa_tls(errore):
                    print("  " + causa_tls(errore))
                print("  Controlla Impostazioni nel pannello e rilancia.")
                break
        time.sleep(CFG["pausa"])

    if not args.dry_run:
        sender.chiudi()

    print()
    print("-" * 62)
    print(f"RIEPILOGO: {len(inviati)} inviate, {saltati_log} doppioni saltati, "
          f"{esclusi} esclusi, {len(eventi_stop)} STOP, {len(errori)} errori.")
    if args.dry_run:
        invia_riepilogo(sender, inviati, errori, esclusi, scadute, dry_run=True,
                        eventi_stop=eventi_stop)
        print("\nTutto corretto? Per l'invio vero: doppio click su avvia_invio.bat")
    else:
        invia_riepilogo(sender, inviati, errori, esclusi, scadute, dry_run=False,
                        eventi_stop=eventi_stop)
        if inviati:
            print(f"Registro aggiornato: {percorso_log()}")
    print("-" * 62)
    return 0


if __name__ == "__main__":
    sys.exit(main())
